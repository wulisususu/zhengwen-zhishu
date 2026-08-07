import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from pptx import Presentation
from pptx.util import Inches

SCRIPTS = Path(__file__).parents[1] / "skill" / "ppt-master" / "scripts"
sys.path.insert(0, str(SCRIPTS))
from pptx_animations import create_grouped_sequence_timing_xml, create_sequence_timing_xml
from svg_to_pptx.pptx_package.builder import create_pptx_with_native_svg
from pptx_animations import read_slide_animation_sequence

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS = {"p": P}

def test_explicit_targets_have_timing_and_no_orphans(tmp_path):
    path = tmp_path / "p4-source.pptx"
    prs = Presentation()
    prs.slides._sldIdLst.clear()
    target_pages = []
    for page, count in enumerate((2, 2, 2, 1, 0), 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        target_pages.append([slide.shapes.add_textbox(Inches(1 + i), Inches(1), Inches(2), Inches(1)).shape_id for i in range(count)])
    prs.save(path)
    effects = (("fade", "wipe"), ("zoom", "appear"), ("fade", "appear"), ("wipe",))
    triggers = ("on-click", "with-previous", "after-previous", "on-click")
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    for page, shape_ids in enumerate(target_pages[:4], 1):
        name = f"ppt/slides/slide{page}.xml"
        timing = create_sequence_timing_xml(list(zip(shape_ids, [0] * len(shape_ids), effects[page - 1])), trigger=triggers[page - 1])
        parts[name] = parts[name].replace(b"</p:sld>", timing.encode() + b"</p:sld>")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items(): archive.writestr(name, data)
    exported = set()
    for page, has_timing in enumerate((True, True, True, True, False), 1):
        with zipfile.ZipFile(path) as archive: root = ET.fromstring(archive.read(f"ppt/slides/slide{page}.xml"))
        assert bool(root.findall("p:timing", NS)) is has_timing
        exported.update((page, int(target.attrib["spid"])) for target in root.findall(".//p:spTgt", NS))
    assert exported == {(page, sid) for page, ids in enumerate(target_pages, 1) for sid in ids}
    assert len(exported) == 7


def test_grouped_targets_keep_each_trigger_structure():
    root = ET.fromstring(create_grouped_sequence_timing_xml([
        (2, 0, "fade", 0.5, "on-click"),
        (3, 120, "wipe", 0.7, "with-previous"),
        (4, 240, "zoom", 0.9, "after-previous"),
    ]))
    assert {target.attrib["spid"] for target in root.findall(".//p:spTgt", NS)} == {"2", "3", "4"}
    node_types = {node.attrib.get("nodeType") for node in root.findall(".//p:cTn", NS)}
    assert {"clickEffect", "withEffect", "afterEffect"} <= node_types
    ids = [node.attrib["id"] for node in root.findall(".//p:cTn", NS)]
    assert len(ids) == len(set(ids))


def test_converter_readback_honors_explicit_on_click_trigger(tmp_path):
    """A group-level trigger overrides the slide default during read-back."""
    svg = tmp_path / "page_01.svg"
    output = tmp_path / "explicit-on-click.pptx"
    svg.write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720">'
        '<g id="box"><rect x="10" y="10" width="100" height="100" fill="#ff0000"/></g>'
        "</svg>",
        encoding="utf-8",
    )
    config = {
        "slides": {
            "page_01": {
                "animation": {"effect": "auto", "trigger": "after-previous", "duration": 0.4},
                "groups": {
                    "box": {
                        "effect": "fade", "duration": 0.5, "delay": 0,
                        "order": 1, "trigger": "on-click",
                    }
                },
            }
        }
    }
    assert create_pptx_with_native_svg(
        [svg], output, animation="auto", animation_config=config,
        animation_cli_overrides={}, verbose=False, pptx_structure="flat",
    )
    with zipfile.ZipFile(output) as archive:
        summary = read_slide_animation_sequence(archive.read("ppt/slides/slide1.xml"))
    assert summary.trigger == "on-click"
