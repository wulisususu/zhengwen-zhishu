"""Internal draft export; this module does not grant formal issuing rights."""
from io import BytesIO
from docx import Document
from docx.shared import Mm, Pt


def internal_docx(payload: dict) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.bottom_margin = Mm(37), Mm(35)
    section.left_margin, section.right_margin = Mm(28), Mm(26)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "仿宋", Pt(16)
    section.header.paragraphs[0].text = "未审核 · 内部草稿"
    section.footer.paragraphs[0].text = "仅供内部流转，不作正式发文"
    doc.add_heading(payload["title"], level=0)
    fields = payload.get("fields", {})
    if fields.get("recipient"):
        doc.add_paragraph(fields["recipient"] + "：")
    for line in payload.get("body", "").splitlines():
        if line.strip():
            doc.add_paragraph(line)
    if fields.get("issuer"):
        doc.add_paragraph(fields["issuer"])
    if payload.get("signDate"):
        doc.add_paragraph(payload["signDate"])
    if fields.get("undertaker"):
        doc.add_paragraph("承办：" + fields["undertaker"])
    if payload.get("policyRefs"):
        doc.add_heading("政策依据（引用记录，须人工核验）", level=1)
        labels = {"current": "现行", "abolished": "已废止", "unverified": "未核验"}
        for ref in payload["policyRefs"]:
            doc.add_paragraph(f"《{ref['title']}》 {ref['clause']} · 第 {ref.get('version') or '未知'} 版 · {labels[ref['status']]}")
            if ref.get("statusNote"):
                doc.add_paragraph("核验提示：" + ref["statusNote"])
            if ref.get("excerpt"):
                doc.add_paragraph(ref["excerpt"])
            doc.add_paragraph("来源：" + ref["sourceUrl"])
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def notice_docx(draft, payload: dict) -> bytes:
    """Render the saved notice with the fixed kit bridge and provenance appendix."""
    import tempfile
    from pathlib import Path
    from fastapi import HTTPException
    from app.services.gongwen.review import render_notice
    from app.services.gongwen.settings import load_settings
    from app.services.gongwen.runner import GongwenToolError
    with tempfile.TemporaryDirectory(prefix="gov-notice-export-") as tmp:
        target = Path(tmp) / "draft.docx"
        try:
            render_notice(draft, target, load_settings(), internal_draft=True)
        except GongwenToolError as exc:
            raise HTTPException(503, "通知生成失败：" + exc.message) from None
        doc = Document(target)
    if payload.get("policyRefs"):
        doc.add_page_break()
        doc.add_paragraph("政策依据与审校记录（内部附件）")
        labels = {"current": "现行", "abolished": "已废止", "unverified": "未核验"}
        for ref in payload["policyRefs"]:
            doc.add_paragraph(f"《{ref['title']}》 {ref['clause']} · 第 {ref.get('version') or '未知'} 版 · {labels[ref['status']]}")
            if ref.get("statusNote"):
                doc.add_paragraph("核验提示：" + ref["statusNote"])
            if ref.get("excerpt"):
                doc.add_paragraph(ref["excerpt"])
            doc.add_paragraph("来源：" + ref.get("sourceUrl", ""))
    report = payload.get("_review", {})
    doc.add_paragraph("内部草稿；本次审校发现严重问题 %s 项；格式机检：%s。须人工核验。" % (report.get("summary", {}).get("critical", "未知"), report.get("formatCheck", {}).get("status", "未完成")))
    out = BytesIO()
    doc.save(out)
    return out.getvalue()
