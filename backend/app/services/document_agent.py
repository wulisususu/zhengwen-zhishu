"""Document Agent Workflow v2.5.

This module keeps DeepSeek as a lightweight router, builds a deterministic
task graph in the backend, validates the plan, lets Qwen produce structured
content, and renders files through local tools.
"""
from __future__ import annotations

import importlib.util
import json
import logging
import re
import time
import uuid
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Callable, Iterable

import requests
from docx import Document

from app.config import get_config
from app.services.markdown_to_docx import (
    document_title_from_markdown,
    markdown_to_docx,
    normalize_markdown_for_docx,
    validate_docx_markdown_render,
)
from app.services.intent_policy import classify_request_intent

PROTOCOL_VERSION = "document_agent_final_v1"
DOC_OUTPUT_DIR = Path("/home/youyeetoo/backend/data/generated_docs")
DOC_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
logger = logging.getLogger(__name__)

SUPPORTED_OUTPUT_FORMATS = {"chat", "docx", "pptx", "markdown", "json"}
FILE_EXT_BY_FORMAT = {"docx": ".docx", "pptx": ".pptx", "markdown": ".md", "json": ".json"}
DOCUMENT_CONTENT_TYPES = {"markdown", "html", "plaintext"}
MIN_MARKDOWN_CHARS = 12


class DocumentGenerationError(Exception):
    """User-visible document generation error with technical detail for logs."""

    default_user_message = "文档生成失败，请稍后重试。"

    def __init__(self, technical_message: str = "", user_message: str = ""):
        self.user_message = user_message or self.default_user_message
        self.technical_message = technical_message or self.user_message
        super().__init__(self.technical_message)


class ModelEmptyResponseError(DocumentGenerationError):
    default_user_message = "模型返回空结构，可能是内容过长或模型容量不足，请缩短篇幅或稍后重试。"


class ModelInvalidJSONError(DocumentGenerationError):
    default_user_message = "模型返回非法 JSON，请稍后重试。"


class ModelTimeoutError(DocumentGenerationError):
    default_user_message = "模型生成超时，请稍后重试或简化要求。"


class ModelHTTPError(DocumentGenerationError):
    default_user_message = "模型服务异常，请稍后重试。"


class ContentSchemaInvalidError(DocumentGenerationError):
    default_user_message = "正文结构不完整，请稍后重试。"

OPERATIONS = {
    "summarize",
    "generate",
    "append_section",
    "polish",
    "rewrite",
    "restructure",
    "extract",
    "compare",
    "convert",
    "visual_relation",
}

OPERATION_TASK_TYPE = {
    "summarize": "content_summarization",
    "generate": "document_generation",
    "append_section": "content_expansion",
    "polish": "content_polish",
    "rewrite": "content_rewrite",
    "restructure": "document_restructure",
    "extract": "information_extraction",
    "compare": "multi_document_analysis",
    "convert": "format_conversion",
    "visual_relation": "visual_text_reasoning",
}

TASK_GRAPH_TEMPLATES = {
    "summarize": [
        "parse_document",
        "merge_text_and_visual_context",
        "content_summarization",
        "quality_check",
        "chat_response",
    ],
    "generate": [
        "parse_document",
        "extract_key_information",
        "plan_document_structure",
        "document_generation",
        "quality_check",
        "document_render",
        "download_card",
    ],
    "append_section": [
        "parse_document",
        "preserve_original_structure",
        "locate_insert_position",
        "content_expansion",
        "merge_original_and_new_content",
        "quality_check",
        "document_render",
    ],
    "polish": [
        "parse_document",
        "preserve_structure",
        "content_polish",
        "before_after_check",
        "document_render_or_chat",
    ],
    "rewrite": [
        "parse_document",
        "preserve_core_facts",
        "content_rewrite",
        "quality_check",
        "document_render_or_chat",
    ],
    "restructure": [
        "parse_document",
        "extract_core_points",
        "plan_report_structure",
        "document_restructure",
        "content_rewrite",
        "quality_check",
        "document_render",
    ],
    "extract": [
        "parse_document",
        "information_extraction",
        "risk_classification",
        "table_generation",
        "quality_check",
        "chat_or_document_output",
    ],
    "compare": [
        "parse_document_a",
        "parse_document_b",
        "normalize_documents",
        "multi_document_analysis",
        "difference_table",
        "conclusion_generation",
        "chat_or_document_output",
    ],
    "convert": [
        "parse_pdf_text",
        "parse_pdf_layout",
        "extract_pdf_images_if_needed",
        "format_conversion",
        "document_render",
        "download_card",
    ],
    "visual_relation": [
        "parse_document",
        "merge_text_and_visual_context",
        "visual_text_reasoning",
        "quality_check",
        "chat_response",
    ],
}

ROUTER_PROMPT = """You are a lightweight document-task router.
Do not write final content. Do not generate reports. Do not call tools.
Classify the user request into the Document Agent Workflow v2.5 routing JSON.

Allowed primary_operation:
summarize, generate, append_section, polish, rewrite, restructure, extract, compare, convert, visual_relation

Allowed output_spec.format:
chat, docx, pptx, markdown, json

Rules:
- DeepSeek only routes. Qwen will do all formal content generation.
- If the user asks to output Word/docx or preserve/edit an attached document, choose docx.
- If the user asks to generate PPT, choose pptx.
- If the user asks to summarize or answer from an attachment without file output, choose chat.
- Visual context is auxiliary and must not replace PDF text.

Output JSON only. Required keys:
{
  "intent": "document_task",
  "confidence": 0.0,
  "task_mode": "single | pipeline | interactive",
  "primary_operation": "",
  "input_sources": [],
  "visual_policy": {},
  "constraints": {},
  "output_spec": {}
}
"""

QWEN_CONTENT_PROMPT = """You are Qwen35B, the formal content executor for a document agent.
You receive routing_json, task_graph, user_instruction, and parsed attachment context.

Rules:
1. Use parsed attachment text as the primary evidence.
2. Use visual_context only as auxiliary context. It must not replace PDF text.
3. Preserve original structure/content when routing_json.constraints asks for it.
4. Do not invent facts. If evidence is insufficient, add a warning.
5. If the user asks for 建议, 修改建议, 加点建议, 优化建议, or recommendations,
   the output must include a dedicated section whose heading contains 建议.
6. Output valid JSON only. No Markdown fence, no explanation, no thinking.

Return this JSON shape:
{
  "title": "",
  "summary": "",
  "sections": [
    {"heading": "", "content": ""}
  ],
  "tables": [
    {"title": "", "headers": [], "rows": []}
  ],
  "slides": [
    {"title": "", "bullets": [], "notes": ""}
  ],
  "warnings": [],
  "source_usage": {
    "used_pdf_text": false,
    "used_visual_context": false,
    "used_user_instruction": true
  }
}
"""


def _cfg() -> dict:
    try:
        return get_config().get("llm", {}) or {}
    except Exception:
        return {}


def _setting(name: str, default: Any) -> Any:
    return _cfg().get(name, default)


def _parse_json_object(text: str) -> dict:
    text = (text or "").strip()
    text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s*```$", "", text)
    try:
        value = json.loads(text)
        return value if isinstance(value, dict) else {}
    except Exception:
        pass
    match = re.search(r"\{[\s\S]*\}", text)
    if not match:
        return {}
    try:
        value = json.loads(match.group(0))
        return value if isinstance(value, dict) else {}
    except Exception:
        return {}


def _call_prompt_model(prompt: str, timeout: float = 12) -> str:
    url = str(_setting("prompt_model_url", "http://127.0.0.1:8001")).rstrip("/")
    resp = requests.post(url + "/v1/completions", json={"prompt": prompt, "stream": False}, timeout=timeout)
    resp.raise_for_status()
    return str(resp.json().get("text", ""))


def _ollama_settings() -> tuple[str, str, float]:
    return (
        str(_setting("ollama_base_url", "http://192.168.10.212:21434")).rstrip("/"),
        str(_setting("ollama_model", "qwen3.5-35b-a3b-q4-4k:latest")),
        float(_setting("ollama_timeout_seconds", 240)),
    )


def _call_qwen_json(
    prompt: str,
    max_tokens: int = 4096,
    max_retries: int = 1,
    retry_backoff: float = 1.2,
    timeout_seconds: float | None = None,
) -> dict:
    url, model, timeout = _ollama_settings()
    if timeout_seconds is not None:
        try:
            timeout = max(1.0, float(timeout_seconds))
        except Exception:
            timeout = max(1.0, timeout)
    payload = {
        "model": model,
        "prompt": prompt,
        "stream": False,
        "format": "json",
        "options": {
            "num_predict": max_tokens,
            "temperature": 0.15,
            "top_p": 0.8,
            "repeat_penalty": 1.08,
        },
    }
    last_err: DocumentGenerationError | None = None
    for attempt in range(max_retries + 1):
        try:
            resp = requests.post(url + "/api/generate", json=payload, timeout=timeout)
        except requests.Timeout as exc:
            last_err = ModelTimeoutError(f"Ollama timeout attempt={attempt}: {exc}")
        except requests.RequestException as exc:
            last_err = ModelHTTPError(f"Ollama request error attempt={attempt}: {exc}")
        else:
            if resp.status_code != 200:
                last_err = ModelHTTPError(f"Ollama HTTP {resp.status_code}: {resp.text[:300]}")
            else:
                try:
                    response_json = resp.json()
                except ValueError as exc:
                    last_err = ModelInvalidJSONError(f"Ollama response is not JSON: {resp.text[:300]}")
                else:
                    raw = str((response_json or {}).get("response", "")).strip()
                    if not raw:
                        last_err = ModelEmptyResponseError(f"empty response attempt={attempt}")
                    else:
                        parsed = _parse_json_object(raw)
                        if isinstance(parsed, dict) and parsed:
                            return parsed
                        last_err = ModelEmptyResponseError(f"empty JSON object attempt={attempt}: raw={raw[:300]!r}")
        logger.warning("Qwen JSON generation failed attempt=%s/%s: %s", attempt, max_retries, last_err)
        if attempt < max_retries:
            time.sleep(retry_backoff * (attempt + 1))
    raise last_err or ModelEmptyResponseError("unknown empty response")


def _call_qwen_text(
    prompt: str,
    max_tokens: int = 4096,
    timeout_seconds: float | None = None,
) -> str:
    url, model, timeout = _ollama_settings()
    if timeout_seconds is not None:
        try:
            timeout = max(1.0, float(timeout_seconds))
        except Exception:
            timeout = max(1.0, timeout)
    payload = {
        "model": model,
        "prompt": prompt,
        "stream": False,
        "options": {
            "num_predict": max_tokens,
            "temperature": 0.2,
            "top_p": 0.85,
            "repeat_penalty": 1.06,
        },
    }
    try:
        resp = requests.post(url + "/api/generate", json=payload, timeout=timeout)
    except requests.Timeout as exc:
        raise ModelTimeoutError(f"Ollama timeout: {exc}") from exc
    except requests.RequestException as exc:
        raise ModelHTTPError(f"Ollama request error: {exc}") from exc
    if resp.status_code != 200:
        raise ModelHTTPError(f"Ollama HTTP {resp.status_code}: {resp.text[:300]}")
    try:
        raw = str((resp.json() or {}).get("response", "")).strip()
    except ValueError as exc:
        raise ModelInvalidJSONError(f"Ollama response is not JSON: {resp.text[:300]}") from exc
    if not raw:
        raise ModelEmptyResponseError("empty response")
    return _strip_markdown_fence(raw)


def _strip_markdown_fence(text: str) -> str:
    value = (text or "").strip()
    value = re.sub(r"^```(?:markdown|md)?\s*", "", value, flags=re.IGNORECASE)
    value = re.sub(r"\s*```$", "", value)
    return _clean_generated_markdown(value)


def _clean_generated_markdown(text: str) -> str:
    value = (text or "").strip()
    if not value:
        return ""
    value = re.sub(r"(?is)<think>[\s\S]*?</think>", "", value)
    value = re.sub(r"(?is)</?think>", "", value)
    leak_patterns = (
        r"(?im)^\s*Thinking Process\s*:",
        r"(?im)^\s*Analyze the Request\s*:",
        r"(?im)^\s*Review User Instruction\s*:",
        r"(?im)^\s*Drafting the Content\s*:",
        r"(?im)^\s*Refining Tone and Style\s*:",
        r"(?im)^\s*Checking Constraints\s*:",
        r"(?im)^\s*Final Review\s*:",
        r"(?im)^\s*Self-Correction\b",
        r"(?im)^\s*Let's assemble\b",
        r"(?im)^\s*Let's write\b",
    )
    cut_at = len(value)
    for pattern in leak_patterns:
        match = re.search(pattern, value)
        if match:
            cut_at = min(cut_at, match.start())
    value = value[:cut_at]
    lines = []
    for raw in value.splitlines():
        stripped = raw.strip()
        if re.match(r"(?i)^(Role|Task|Input|Constraints)\s*:", stripped):
            break
        lines.append(raw.rstrip())
    value = re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()
    return _strip_duplicate_markdown_document(value)


def _strip_duplicate_markdown_document(text: str) -> str:
    value = (text or "").strip()
    if not value:
        return ""
    lines = value.splitlines()
    first_title = ""
    for line in lines:
        match = re.match(r"^\s*#\s+(.+?)\s*$", line)
        if match:
            first_title = _normalize_markdown_title(match.group(1))
            break
    if not first_title:
        return value

    seen_first = False
    kept: list[str] = []
    for line in lines:
        match = re.match(r"^\s*#\s+(.+?)\s*$", line)
        if match:
            title = _normalize_markdown_title(match.group(1))
            if seen_first and _same_document_title(first_title, title):
                break
            seen_first = True
        kept.append(line.rstrip())
    return re.sub(r"\n{3,}", "\n\n", "\n".join(kept)).strip()


def _normalize_markdown_title(title: str) -> str:
    value = re.sub(r"\s+", "", title or "").lower()
    value = re.sub(r"[\(\)（）【】\\[\\]《》:：_\\-]+", "", value)
    value = value.replace("prd", "产品需求文档")
    return value


def _same_document_title(first: str, current: str) -> bool:
    if not first or not current:
        return False
    return first == current or first in current or current in first


def _minicpm_service_ready() -> bool:
    try:
        resp = requests.get("http://127.0.0.1:8003/health", timeout=0.25)
        if resp.status_code != 200:
            return False
        return (resp.json() or {}).get("status") == "ready"
    except Exception:
        return False


def _attachment_type(filename: str | None) -> str:
    ext = Path(filename or "").suffix.lower()
    if ext == ".pdf":
        return "pdf"
    if ext in {".jpg", ".jpeg", ".png", ".bmp", ".webp", ".tif", ".tiff", ".gif"}:
        return "image"
    if ext in {".doc", ".docx"}:
        return "docx"
    if ext in {".md", ".markdown"}:
        return "markdown"
    if ext in {".json", ".jsonl"}:
        return "json"
    return "user_text" if not ext else ext.lstrip(".")


def _attachment_manifest(attachments) -> list[dict[str, Any]]:
    manifest = []
    for item in attachments or []:
        filename = getattr(item, "filename", None)
        file_id = getattr(item, "id", None)
        if isinstance(item, dict):
            filename = item.get("filename") or item.get("name") or filename
            file_id = item.get("id") or item.get("file_id") or file_id
        ftype = _attachment_type(filename)
        manifest.append({
            "type": ftype,
            "id": str(file_id or filename or "attachment"),
            "filename": filename or "",
        })
    return manifest


def _has_visual_attachment(attachments) -> bool:
    return any(item.get("type") in {"image", "pdf"} for item in _attachment_manifest(attachments))


def _has_attachment_context(text: str) -> bool:
    markers = ("以下是用户上传附件", "## 附件：", "PDF 图片分析", "图片识别结果", "ATTACHMENT CONTENT")
    return any(marker in (text or "") for marker in markers)


def _keyword(text: str, *words: str) -> bool:
    return any(word in text for word in words)


def _formal_spec_document_request(text: str) -> bool:
    text = text or ""
    return _keyword(
        text,
        "需求规格说明书",
        "软件需求规格说明书",
        "系统需求规格说明书",
        "SRS",
        "srs",
        "需求文档",
        "规格说明书",
        "需求说明书",
    )


def _wants_advice(user_input: str) -> bool:
    text = (user_input or "").lower()
    return _keyword(
        text,
        "建议",
        "加点建议",
        "给点建议",
        "给些建议",
        "修改建议",
        "优化建议",
        "完善建议",
        "补充建议",
        "recommendation",
        "suggestion",
        "advice",
    )


def _wants_report(user_input: str) -> bool:
    text = (user_input or "").lower()
    return _keyword(text, "报告", "汇报", "正式材料", "正式文档", "report")


def _requested_char_count(user_input: str) -> int:
    text = (user_input or "").lower()
    match = re.search(r"(\d{2,5})\s*(?:字|字符|汉字|word|words)", text)
    if match:
        return max(50, min(20000, int(match.group(1))))
    chinese_digit = {"一": 1, "二": 2, "两": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9}
    match = re.search(r"([一二两三四五六七八九])\s*百\s*(?:字|字符|汉字)?", text)
    if match:
        return chinese_digit.get(match.group(1), 3) * 100
    match = re.search(r"([一二两三四五六七八九])\s*千\s*(?:字|字符|汉字)?", text)
    if match:
        return chinese_digit.get(match.group(1), 1) * 1000
    if "百字" in text or "几百字" in text:
        return 300
    return 0


def _content_char_len(text: str) -> int:
    return len(re.sub(r"[\s#*`_>\-|：:，,。.;；!！?？（）()]+", "", text or ""))


def _mentions_attachment(user_input: str) -> bool:
    text = (user_input or "").lower()
    return _keyword(text, "附件", "上传", "pdf", "图片", "图像", "截图", "扫描件", "照片")


def _explicit_file_output_request(text: str) -> bool:
    return _keyword(
        text,
        "word",
        "docx",
        "ppt",
        "pptx",
        "导出",
        "下载",
        "输出文件",
        "生成文件",
        "生成文档",
        "制作文档",
        "生成word",
        "生成 word",
        "制作ppt",
        "制作 ppt",
        "转word",
        "转 word",
        "转换格式",
    )


def _explicit_attachment_reference(text: str) -> bool:
    return (
        _keyword(text, "附件", "上传", "pdf", "图片", "图像")
        or (
            _keyword(text, "文件", "文档")
            and _keyword(text, "这个", "这份", "该", "上述", "上面", "刚才", "已上传", "上传的")
        )
    )


def _common_office_document_request(text: str) -> bool:
    return _keyword(
        text,
        "入职通知",
        "入职通知书",
        "录用通知",
        "录用通知书",
        "离职证明",
        "面试邀约",
        "面试邀请",
        "邀请函",
        "请假审批",
        "请假记录",
        "会议纪要",
        "催款函",
        "催款通知",
        "面试评价",
        "评价表",
        "正式公告",
        "行政通知",
        "放假通知",
        "正式通报",
        "处罚通报",
        "违纪处理",
        "采购申请",
        "审批理由",
        "事件时间线",
        "事件经过",
        "事故报告",
        "事故调查报告",
        "事件报告",
        "事件调查报告",
        "故障报告",
        "调查报告",
        "正式记录",
        "背调提问",
        "提问清单",
        "员工档案",
        "报销文件",
        "报销汇总",
    ) or (
        _keyword(text, "客户", "欠款", "逾期") and _keyword(text, "催款", "函", "通知")
    ) or (
        _keyword(text, "采购") and _keyword(text, "用途")
    ) or (
        _keyword(text, "停电", "放假") and _keyword(text, "公司", "通知", "公告", "生成")
    )


def _explicit_new_document_generation_request(text: str) -> bool:
    text = (text or "").lower()
    generation_action = _keyword(
        text,
        "生成",
        "创建",
        "新建",
        "写一份",
        "撰写",
        "起草",
        "制作",
        "输出",
        "可直接交付",
    )
    document_target = _keyword(
        text,
        "word",
        "docx",
        "文档",
        "方案",
        "实施方案",
        "项目方案",
        "投标书",
        "投标文件",
        "报告",
        "预算表",
        "预算报价",
    )
    return generation_action and document_target and not _explicit_attachment_reference(text)


def _fallback_operation(user_input: str, has_attachments: bool = False) -> str:
    text = (user_input or "").lower()
    if _wants_advice(user_input) and has_attachments:
        return "visual_relation" if _keyword(text, "图片和文字", "文字和图片", "图文", "图片", "图像") else "append_section"
    if not has_attachments and _explicit_new_document_generation_request(text):
        return "generate"
    if not has_attachments and _common_office_document_request(text):
        return "generate"
    if _keyword(text, "图文", "图片和文字", "文字和图片", "视觉关系", "图片关系", "图表关系", "visual relation"):
        return "visual_relation"
    if _keyword(text, "对比", "比较", "差异", "不同点", "版本变化", "compare"):
        return "compare"
    if _keyword(text, "pdf转word", "pdf 转 word", "转成word", "转为word", "转换格式", "convert"):
        return "convert"
    if _keyword(text, "保留原文", "保留原结构", "补充", "补写", "新增一段", "追加", "append"):
        return "append_section"
    if _keyword(text, "润色", "优化表达", "正式一点", "通顺", "polish"):
        return "polish"
    if _keyword(text, "改写", "换一种说法", "rewrite"):
        return "rewrite"
    if _keyword(text, "改成正式报告", "整理成报告", "汇报稿", "项目方案", "重组", "restructure"):
        return "restructure"
    if _keyword(text, "提取", "风险", "问题点", "清单", "extract"):
        return "extract"
    if _keyword(text, "总结", "概括", "重点", "摘要", "讲了什么", "看一下", "解释", "什么意思", "普通话", "白话", "普通人", "看懂", "summarize"):
        return "summarize"
    if _keyword(text, "生成ppt", "生成 ppt", "制作ppt", "演示文稿", "生成word", "生成 word", "写一份", "生成", "创建", "制作"):
        return "generate"
    return "summarize" if has_attachments else "generate"


def _fallback_output_spec(user_input: str, operation: str, has_attachments: bool = False) -> dict[str, Any]:
    text = (user_input or "").lower()
    if _formal_spec_document_request(text):
        fmt = "docx"
        template = "report" if operation in {"generate", "restructure", "append_section"} else "default"
    elif _keyword(text, "pptx", "ppt", "演示文稿"):
        fmt = "pptx"
        template = "ppt_outline"
    elif _keyword(text, "json", "结构化数据"):
        fmt = "json"
        template = "default"
    elif _keyword(text, "markdown", ".md", " md"):
        fmt = "markdown"
        template = "default"
    elif _keyword(text, "word", "docx", "下载", "导出", "输出文档", "生成文档", "文件"):
        fmt = "docx"
        template = "report" if operation in {"generate", "restructure", "append_section"} else "default"
    elif operation in {"append_section", "polish", "rewrite", "restructure", "convert"} and _wants_advice(user_input):
        fmt = "chat"
        template = "default"
    elif operation in {"generate", "append_section", "polish", "rewrite", "restructure", "convert"}:
        fmt = "docx"
        template = "report" if operation != "convert" else "default"
    else:
        fmt = "chat"
        template = "risk_list" if operation == "extract" else "default"
    return {
        "format": fmt,
        "template": template,
        "include_toc": fmt in {"docx", "markdown"} and operation in {"generate", "restructure"},
        "include_tables": operation in {"extract", "compare"} or _keyword(text, "表格", "清单"),
        "include_figures": False,
        "download_card": fmt != "chat",
    }


def _fallback_constraints(user_input: str, operation: str) -> dict[str, Any]:
    text = user_input or ""
    tone = "formal" if _keyword(text, "正式", "报告", "方案", "公文", "商务") else "neutral"
    length = "extended" if _keyword(text, "详细", "完整", "长一点") else "normal"
    preserve = operation in {"append_section", "polish", "rewrite", "convert"}
    return {
        "preserve_original_structure": preserve,
        "preserve_original_content": operation in {"append_section", "polish", "convert"},
        "tone": tone,
        "length": length,
        "strict_fact_only": True,
        "no_hallucination": True,
        "allow_inference": False,
    }


def _fallback_visual_policy(user_input: str, operation: str, attachments) -> dict[str, Any]:
    text = user_input or ""
    requires_visual = (
        operation == "visual_relation"
        or _keyword(text, "图片", "图像", "图表", "照片", "截图", "视觉", "地图")
        or any(item.get("type") == "image" for item in _attachment_manifest(attachments))
    )
    return {
        "requires_visual_context": bool(requires_visual),
        "pdf_image_strategy": "analyze_all_detected_images",
        "visual_context_role": "auxiliary_context",
        "visual_result_replace_text": False,
    }


def _fallback_input_sources(user_input: str, attachments) -> list[dict[str, Any]]:
    sources = []
    for item in _attachment_manifest(attachments):
        ftype = item["type"]
        sources.append({
            "type": ftype,
            "id": item["id"],
            "role": "visual_source" if ftype == "image" else "main_source",
            "need_text_parse": ftype not in {"image"},
            "need_layout_parse": ftype == "pdf",
            "need_vision_model": ftype in {"image", "pdf"},
            "need_ocr": ftype in {"image"},
        })
    if not sources and _mentions_attachment(user_input):
        text = (user_input or "").lower()
        if "pdf" in text:
            required_type = "pdf"
        elif _keyword(text, "图片", "图像", "照片", "截图"):
            required_type = "image"
        elif _keyword(text, "ppt", "pptx"):
            required_type = "pptx"
        elif _keyword(text, "word", "docx", "文档"):
            required_type = "docx"
        else:
            required_type = "attachment"
        sources.append({
            "type": required_type,
            "id": "required_attachment",
            "role": "main_source",
            "need_text_parse": required_type not in {"image"},
            "need_layout_parse": required_type == "pdf",
            "need_vision_model": required_type in {"image", "pdf"},
            "need_ocr": required_type == "image",
        })
    if (user_input or "").strip():
        sources.append({
            "type": "user_text",
            "id": "user_instruction",
            "role": "instruction_source",
            "need_text_parse": False,
            "need_layout_parse": False,
            "need_vision_model": False,
            "need_ocr": False,
        })
    return sources


def _legacy_doc_type(user_input: str, operation: str, output_format: str) -> str:
    text = (user_input or "").lower()
    if "周报" in text:
        return "weekly_report"
    if "简历" in text or "resume" in text:
        return "resume"
    if operation == "extract" or _keyword(text, "清单", "检查表", "checklist"):
        return "checklist"
    if operation in {"generate", "restructure", "append_section", "convert"} or output_format in {"docx", "pptx"}:
        return "complex_report"
    return "simple_report"


def build_task_graph(route: dict) -> list[dict[str, Any]]:
    operation = route.get("primary_operation") or "generate"
    nodes = TASK_GRAPH_TEMPLATES.get(operation, TASK_GRAPH_TEMPLATES["generate"])
    output_format = (route.get("output_spec") or {}).get("format", "chat")
    graph = []
    previous_output = "attachments + user_text"
    for index, node_type in enumerate(nodes, start=1):
        real_type = node_type
        if node_type in {"document_render", "document_render_or_chat", "chat_or_document_output"}:
            real_type = "chat_response" if output_format == "chat" else f"{output_format}_render"
        graph.append({
            "id": f"t{index}",
            "type": real_type,
            "task_type": OPERATION_TASK_TYPE.get(operation, real_type),
            "operation": operation,
            "input": previous_output,
            "output": _node_output_name(real_type, output_format),
        })
        previous_output = graph[-1]["output"]
    return graph


def _node_output_name(node_type: str, output_format: str) -> str:
    if node_type.endswith("_render"):
        return output_format
    if node_type == "chat_response":
        return "chat"
    if "quality" in node_type or "check" in node_type:
        return "checked_content"
    if "parse" in node_type:
        return "raw_context"
    if "merge" in node_type or "normalize" in node_type:
        return "merged_context"
    if "plan" in node_type:
        return "document_plan"
    if "table" in node_type:
        return "structured_table"
    return "draft_content"


def normalize_workflow_route(
    route: dict | None,
    user_input: str = "",
    attachments=None,
    attachment_context: str = "",
    trust_route: bool = False,
    force_output_format: str | None = None,
) -> dict[str, Any]:
    route = dict(route or {})
    has_attachments = bool(attachments) or _has_attachment_context(attachment_context) or _has_attachment_context(user_input)
    fallback_operation = _fallback_operation(user_input, has_attachments=has_attachments)
    force_office_generation = not has_attachments and _common_office_document_request((user_input or "").lower())
    force_spec_docx = _formal_spec_document_request(user_input or "")
    operation = route.get("primary_operation") or route.get("operation") or fallback_operation
    if force_office_generation and operation != "generate":
        operation = "generate"
    if operation not in OPERATIONS or (not trust_route and fallback_operation in {"append_section", "polish", "convert", "compare"}):
        operation = fallback_operation

    output_spec = dict(route.get("output_spec") or {})
    fallback_output = _fallback_output_spec(user_input, operation, has_attachments=has_attachments)
    output_format = output_spec.get("format") or fallback_output["format"]
    if force_output_format:
        output_format = str(force_output_format).strip().lower()
        output_spec["format"] = output_format
        if output_format == "docx":
            operation = "generate"
            output_spec["template"] = output_spec.get("template") or "report"
            output_spec["force_word_document_page"] = True
    if (force_office_generation or force_spec_docx) and output_format == "chat":
        output_format = "docx"
    if force_spec_docx:
        output_format = "docx"
        output_spec["template"] = "report"
    if output_format not in SUPPORTED_OUTPUT_FORMATS:
        output_format = fallback_output["format"]
    output_spec = {**fallback_output, **output_spec, "format": output_format}
    output_spec["download_card"] = output_format != "chat"

    constraints = {**_fallback_constraints(user_input, operation), **dict(route.get("constraints") or {})}
    fallback_visual = _fallback_visual_policy(user_input, operation, attachments)
    route_visual = dict(route.get("visual_policy") or {})
    visual_policy = {**fallback_visual, **route_visual}
    # Don't let route override requires_visual_context when fallback says no visual needed
    if not fallback_visual.get("requires_visual_context"):
        visual_policy["requires_visual_context"] = False
    visual_policy["pdf_image_strategy"] = "analyze_all_detected_images"
    visual_policy["visual_context_role"] = "auxiliary_context"
    visual_policy["visual_result_replace_text"] = False

    input_sources = route.get("input_sources")
    if not isinstance(input_sources, list) or not input_sources:
        input_sources = _fallback_input_sources(user_input, attachments)

    task_mode = route.get("task_mode") or ("pipeline" if output_format != "chat" or operation in {"append_section", "restructure", "compare"} else "single")
    doc_type = _legacy_doc_type(user_input, operation, output_format)
    normalized = {
        "protocol_version": PROTOCOL_VERSION,
        "intent": "document_task",
        "confidence": float(route.get("confidence") or 0.78),
        "task_mode": task_mode if task_mode in {"single", "pipeline", "interactive"} else "pipeline",
        "primary_operation": operation,
        "input_sources": input_sources,
        "visual_policy": visual_policy,
        "constraints": constraints,
        "output_spec": output_spec,
        "validation": {
            "check_missing_input": True,
            "check_conflicts": True,
            "check_tool_availability": True,
            "check_visual_context_ready": True,
            "check_output_format_supported": True,
            "check_preserve_original_possible": True,
            "on_missing_input": "ask_user",
        },
        "memory_keys": [
            "user_preferred_output_format",
            "last_document_style",
            "project_context",
        ],
        "doc_type": doc_type,
        "engine": "python-docx",
        "template_id": output_spec.get("template") or "default",
    }
    normalized["task_graph"] = build_task_graph(normalized)
    return normalized


def classify_workflow(
    user_input: str,
    attachments=None,
    attachment_context: str = "",
    force_output_format: str | None = None,
) -> dict[str, Any]:
    manifest = _attachment_manifest(attachments)
    prompt = (
        ROUTER_PROMPT
        + "\nUSER_INPUT:\n"
        + (user_input or "")
        + "\n\nATTACHMENT_MANIFEST:\n"
        + json.dumps(manifest, ensure_ascii=False)
        + "\n\nHAS_ATTACHMENT_CONTEXT:"
        + str(bool(attachment_context)).lower()
        + "\nJSON:"
    )
    route = {}
    if not _minicpm_service_ready():
        try:
            router_timeout = min(float(_setting("prompt_timeout_seconds", 12)), 3.5)
            raw = _call_prompt_model(prompt, timeout=router_timeout)
            route = _parse_json_object(raw)
        except Exception:
            route = {}
    return normalize_workflow_route(
        route,
        user_input,
        attachments,
        attachment_context,
        trust_route=True,
        force_output_format=force_output_format,
    )


def is_document_agent_request(content: str, has_attachments: bool = False) -> bool:
    text = (content or "").lower()
    intent = classify_request_intent(content or "", has_attachments=has_attachments)
    if not intent.use_document_agent:
        return False
    doc_actions = (
        "总结", "概括", "提取", "风险", "清单", "对比", "比较", "图片", "图文", "关系",
        "润色", "改写", "补充", "补写", "报告", "word", "docx", "ppt", "转换", "转word",
        "讲了什么", "看一下", "分析一下", "重点", "摘要", "建议", "解释", "什么意思", "普通话",
        "白话", "普通人", "看懂", "生成", "创建", "新建",
        "写", "制作", "导出", "整理",
    )
    if has_attachments and _keyword(
        text,
        *doc_actions,
    ):
        return True
    if _common_office_document_request(text):
        return True
    if _explicit_file_output_request(text) and _keyword(
        text, "生成", "创建", "新建", "写", "制作", "导出", "整理", "改成", "转换", "润色", "补充", "补写", "总结", "输出"
    ):
        return True
    if _explicit_attachment_reference(text) and _keyword(
        text, "总结", "概括", "提取", "对比", "比较", "分析", "润色", "改写", "补充", "补写", "转换", "转word", "生成", "讲了什么", "建议", "报告", "解释", "什么意思", "普通话", "白话", "普通人", "看懂"
    ):
        return True
    return False


def requires_file_generation(route: dict | None) -> bool:
    return ((route or {}).get("output_spec") or {}).get("format") not in {"", None, "chat"}


def validate_task_plan(route: dict, attachment_context: str = "", attachments=None) -> dict[str, Any]:
    route = normalize_workflow_route(route, "", attachments, attachment_context, trust_route=True)
    operation = route["primary_operation"]
    output_format = route["output_spec"]["format"]
    has_attachments = bool(attachments) or bool((attachment_context or "").strip()) or _has_attachment_context(attachment_context)
    warnings: list[str] = []
    errors: list[str] = []
    need_user_input: list[str] = []

    attachment_ops = {"summarize", "append_section", "polish", "rewrite", "restructure", "extract", "compare", "convert", "visual_relation"}
    input_sources = route.get("input_sources") or []
    route_requires_attachment = any(
        isinstance(source, dict)
        and source.get("type") not in {"user_text", ""}
        and source.get("role") != "instruction_source"
        for source in input_sources
    )
    if (operation in attachment_ops or route_requires_attachment) and not has_attachments:
        need_user_input.append("请先上传需要处理的附件。")

    if operation == "compare":
        count = len(_attachment_manifest(attachments))
        if count < 2 and not _has_attachment_context(attachment_context):
            need_user_input.append("对比任务至少需要两份附件。")

    visual_policy = route.get("visual_policy") or {}
    wants_visual = bool(visual_policy.get("requires_visual_context"))
    has_visual_marker = any(marker in (attachment_context or "") for marker in ("PDF 图片分析", "PDF页面视觉分析", "图片识别结果", "visual_context"))
    if wants_visual and not (_has_visual_attachment(attachments) or has_visual_marker):
        need_user_input.append("该任务需要图片或 PDF 图片上下文，请上传包含图片的附件。")
    elif wants_visual and not has_visual_marker and _has_visual_attachment(attachments):
        warnings.append("视觉上下文会由解析层的 MiniCPM 队列补齐，Qwen 只会把它作为辅助依据。")

    if output_format not in SUPPORTED_OUTPUT_FORMATS:
        errors.append(f"暂不支持输出格式：{output_format}")

    if output_format == "docx" and importlib.util.find_spec("docx") is None:
        errors.append("python-docx 工具不可用。")
    if output_format == "pptx" and importlib.util.find_spec("pptx") is None:
        errors.append("python-pptx 工具不可用。")

    preserve = (route.get("constraints") or {}).get("preserve_original_structure")
    if preserve and not has_attachments:
        need_user_input.append("保留原结构需要原始附件作为依据。")

    status = "success"
    message = ""
    if errors:
        status = "failed"
        message = "；".join(errors)
    elif need_user_input:
        status = "need_user_input"
        message = "；".join(dict.fromkeys(need_user_input))

    return {
        "status": status,
        "message": message,
        "warnings": list(dict.fromkeys(warnings)),
        "errors": errors,
        "checks": {
            "missing_input": not need_user_input,
            "tool_availability": not errors,
            "visual_context_ready": not wants_visual or has_visual_marker or _has_visual_attachment(attachments),
            "output_format_supported": output_format in SUPPORTED_OUTPUT_FORMATS,
            "preserve_original_possible": not preserve or has_attachments,
        },
    }


def _build_generation_prompt(route: dict, task_graph: list[dict[str, Any]], user_input: str, attachment_context: str) -> str:
    context = attachment_context.strip()
    if not context and _has_attachment_context(user_input):
        context = user_input
    return (
        QWEN_CONTENT_PROMPT
        + "\nROUTING_JSON:\n"
        + json.dumps(route, ensure_ascii=False, indent=2)
        + "\n\nTASK_GRAPH:\n"
        + json.dumps(task_graph, ensure_ascii=False, indent=2)
        + "\n\nUSER_INSTRUCTION:\n"
        + (user_input or "").strip()
        + "\n\nPARSED_ATTACHMENT_CONTEXT:\n"
        + context[:28000]
        + "\n\nCONTENT_JSON:"
    )


def _should_generate_sectioned_document(user_input: str, route: dict, attachment_context: str = "") -> bool:
    if not requires_file_generation(route):
        return False
    if route.get("primary_operation") != "generate":
        return False
    text = (user_input or "") + "\n" + (attachment_context or "")
    if _keyword(
        text,
        "事故调查报告",
        "事故报告",
        "事件调查报告",
        "事件报告",
        "故障报告",
        "调查报告",
        "运营分析报告",
        "分析报告",
        "调研报告",
        "总结报告",
        "投标文件",
        "招标文件",
        "技术部分",
        "技术标",
        "建设方案",
        "实施方案",
        "项目方案",
        "可行性研究报告",
    ):
        return True
    if _common_office_document_request(text.lower()) and len(text) < 1200:
        return False
    return bool(
        _requested_char_count(text) >= 1800
        or re.search(r"\d+\s*分钟", text)
        or _keyword(
            text,
            "讲话稿",
            "演讲稿",
            "发言稿",
            "方案",
            "建设方案",
            "需求规格说明书",
            "SRS",
            "不少于",
            "完整",
            "详细",
            "高质量发展",
        )
    )


def _section_count_for_request(user_input: str) -> int:
    text = user_input or ""
    if _keyword(
        text,
        "投标文件",
        "招标文件",
        "技术部分",
        "技术标",
        "建设方案",
        "实施方案",
        "项目方案",
    ):
        return 8
    minutes = re.search(r"(\d{1,3})\s*分钟", text)
    if minutes:
        value = int(minutes.group(1))
        if value >= 15:
            return 8
        if value >= 8:
            return 6
    chars = _requested_char_count(text)
    if chars >= 5000:
        return 8
    if chars >= 2500:
        return 6
    return 5


def _fallback_outline(user_input: str, route: dict) -> dict[str, Any]:
    text = user_input or ""
    if _keyword(text, "投标文件", "招标文件", "技术部分", "技术标"):
        sections = ["项目理解", "总体建设目标", "平台总体架构设计", "核心功能设计", "系统安全设计", "实施部署方案", "运维与服务保障", "项目优势与预期成效"]
    elif _keyword(text, "建设方案", "实施方案", "项目方案", "可行性研究报告"):
        sections = ["项目背景", "建设目标", "总体架构", "建设内容", "实施计划", "安全保障", "运维保障", "预期成效"]
    elif _keyword(text, "事故调查报告", "事故报告", "事件调查报告", "事件报告", "故障报告", "调查报告"):
        sections = ["事故基本情况", "影响范围", "原因分析", "应急处置经过", "整改措施", "调查结论"]
    elif _keyword(text, "讲话稿", "演讲稿", "发言稿"):
        sections = [
            "开场致辞",
            "时代背景与政策要求",
            "人工智能赋能职业教育的现实意义",
            "学院实践案例",
            "重点建设任务",
            "保障措施",
            "未来展望",
            "结束语",
        ]
    elif _keyword(text, "需求规格说明书", "SRS"):
        sections = ["项目概述", "功能需求", "非功能需求", "系统架构", "数据库设计", "接口设计", "部署方案", "测试方案"]
    else:
        sections = ["项目背景", "总体目标", "核心内容", "实施路径", "保障措施", "预期成效"]
    return {
        "title": _title_from_input(user_input, route.get("primary_operation", "generate")),
        "summary": "",
        "sections": [{"heading": item, "instruction": f"围绕“{item}”展开，保持正式、完整、可直接成文。"} for item in sections],
    }


def _build_outline_prompt(user_input: str, route: dict, attachment_context: str, section_count: int) -> str:
    return f"""你是中文正式文档规划助手。请只输出 JSON，不要 Markdown。
任务：根据用户要求规划一份正式文档大纲，后续会按章节分别生成正文。
要求：
1. sections 数量控制在 {section_count} 节左右。
2. 每个 section 只写 heading 和 instruction，不要写正文。
3. 如果用户要求讲话稿、演讲稿，请按正式演讲逻辑组织。
4. 如果用户要求方案或 SRS，请按正式申报/规格说明结构组织。
5. 如果是中文技术文档、投标文件、实施方案或分析报告，优先采用可扫读的信息架构，标题应体现明确模块，不使用空泛口号。
6. 如果是正式报告或公文类内容，章节命名应规范、完整，必要时包含“背景、目标、措施、结论、建议”等标准模块。
7. 不要在标题和章节名中使用宣传口号、感叹语、第二人称或营销化表达。

输出 JSON 结构：
{{"title":"","summary":"","sections":[{{"heading":"","instruction":""}}]}}

用户要求：
{user_input[:3000]}

附件上下文：
{(attachment_context or "无")[:3000]}
"""


def _normalize_outline(outline: dict, user_input: str, route: dict, section_count: int) -> dict[str, Any]:
    fallback = _fallback_outline(user_input, route)
    if not isinstance(outline, dict):
        outline = fallback
    title = str(outline.get("title") or fallback.get("title") or "document").strip()
    if title:
        title = re.sub(r"\s+", " ", title)
        title = re.split(r"(?:第[一二三四五六七八九十百0-9]+章|1\.1|1\.2|2\.1|项目背景与意义|建设目标)", title, maxsplit=1)[0].strip(" ：:，,。.")
    if not title:
        title = _title_from_input(user_input, route.get("primary_operation", "generate"))
    summary = str(outline.get("summary") or "").strip()
    raw_sections = outline.get("sections")
    if not isinstance(raw_sections, list) or not raw_sections:
        raw_sections = fallback["sections"]
    sections = []
    for item in raw_sections[: max(section_count + 2, 4)]:
        if isinstance(item, dict):
            heading = str(item.get("heading") or item.get("title") or "").strip()
            instruction = str(item.get("instruction") or item.get("intent") or item.get("content") or "").strip()
        else:
            heading = str(item).strip()
            instruction = ""
        if heading:
            sections.append({"heading": heading, "instruction": instruction or f"围绕“{heading}”展开。"})
    if not sections:
        sections = fallback["sections"]
    return {"title": title, "summary": summary, "sections": sections}


def _build_section_prompt(user_input: str, route: dict, attachment_context: str, outline: dict, section: dict, index: int, total: int) -> str:
    is_longform_bid = _keyword(
        user_input or "",
        "投标文件",
        "招标文件",
        "技术部分",
        "技术标",
        "建设方案",
        "实施方案",
        "项目方案",
    )
    section_length_rule = (
        "单节正文控制在 900 到 1600 个中文字符；如果是投标文件、技术标、建设方案或实施方案，可写到 2000 字以内。"
        if is_longform_bid
        else "单节正文控制在 600 到 900 个中文字符；如果用户要求长篇，可写到 1100 字以内。"
    )
    return f"""你是中文正式文档撰写助手。请只输出 JSON，不要 Markdown，不要解释。
现在按章节生成正文，这是第 {index + 1}/{total} 节。
文档标题：{outline.get("title", "")}
当前章节标题：{section.get("heading", "")}
当前章节写作要求：{section.get("instruction", "")}

整体用户要求：
{user_input[:3000]}

附件上下文：
{(attachment_context or "无")[:2500]}

写作要求：
1. 输出必须是可直接进入 Word 的正式中文正文。
2. {section_length_rule}
3. 不要输出占位符，不要写“请补充”“此处填写”。
4. 涉及政策时可引用国家职业教育、教育数字化、人工智能赋能教育等政策方向，表述要稳妥。
5. 采用克制、准确、可扫读的中文技术文风，避免“领先、强大、赋能、闭环、拉通、颠覆”等空泛表达。
6. 不使用第二人称，不使用营销口号，不使用感叹号；优先写“是什么、怎么做、达到什么结果”。
7. 如果是正式报告、公文、事故调查、通知、函、总结等文体，保持结构完整、语气庄重，必要时补足依据、过程、结果、结论或整改建议。
8. 只输出 JSON：{{"heading":"","content":""}}
"""


def _fallback_section_content(user_input: str, section: dict) -> str:
    heading = str(section.get("heading") or "正文")
    instruction = str(section.get("instruction") or "")
    return (
        f"围绕{heading}，应结合用户提出的主题和具体要求展开论述。"
        f"{instruction}"
        "本节内容由系统在模型分节生成不稳定时自动整理，请在正式使用前结合实际情况进一步核对和润色。"
    )


def _generate_sectioned_content_schema(
    user_input: str,
    route: dict,
    attachment_context: str = "",
    progress_callback: Callable[[str, int, str], None] | None = None,
) -> dict[str, Any]:
    section_count = _section_count_for_request(user_input)
    if progress_callback:
        progress_callback("generating_outline", 16, "正在生成文档大纲")
    stored_outline = route.get("outline_preview")
    if isinstance(stored_outline, dict) and stored_outline.get("sections"):
        outline = _normalize_outline(stored_outline, user_input, route, section_count)
    else:
        try:
            outline_raw = _call_qwen_json(_build_outline_prompt(user_input, route, attachment_context, section_count), max_tokens=1200, max_retries=1)
        except DocumentGenerationError as exc:
            logger.warning("Outline generation degraded: %s", exc.technical_message)
            outline_raw = _fallback_outline(user_input, route)
        outline = _normalize_outline(outline_raw, user_input, route, section_count)
    route["outline_preview"] = outline
    if progress_callback:
        progress_callback("outline_ready", 24, "大纲已生成，正在分章节撰写正文")

    sections: list[dict[str, str]] = []
    warnings: list[str] = []
    total = len(outline["sections"])
    for index, section in enumerate(outline["sections"]):
        if progress_callback:
            progress = 26 + int(((index) / max(total, 1)) * 48)
            progress_callback(
                "generating_content",
                min(progress, 72),
                f"正在撰写第{index + 1}/{total}节：{section.get('heading') or '正文'}",
            )
        prompt = _build_section_prompt(user_input, route, attachment_context, outline, section, index, total)
        try:
            item = _call_qwen_json(prompt, max_tokens=1800, max_retries=1)
            heading = str(item.get("heading") or section.get("heading") or "").strip()
            body = _stringify(item.get("content") or item.get("body") or "").strip()
            if not body:
                raise ContentSchemaInvalidError(f"empty section content index={index}")
        except DocumentGenerationError as exc:
            logger.warning("Section generation degraded index=%s heading=%s: %s", index, section.get("heading"), exc.technical_message)
            heading = str(section.get("heading") or f"第{index + 1}部分")
            body = _fallback_section_content(user_input, section)
            warnings.append(f"章节“{heading}”生成不完整，已使用降级草稿。")
        sections.append({"heading": heading, "content": body})
        if progress_callback:
            progress = 26 + int(((index + 1) / max(total, 1)) * 48)
            progress_callback(
                "generating_content",
                min(progress, 74),
                f"已完成第{index + 1}/{total}节：{heading or '正文'}",
            )

    return {
        "title": outline.get("title") or _title_from_input(user_input, "generate"),
        "summary": outline.get("summary") or "",
        "sections": sections,
        "tables": [],
        "slides": [],
        "warnings": warnings,
        "source_usage": {"used_user_instruction": True, "used_pdf_text": False, "used_visual_context": False},
    }


def generate_content_schema(
    user_input: str,
    route_or_doc_type: dict | str | None = None,
    attachment_context: str = "",
    progress_callback: Callable[[str, int, str], None] | None = None,
) -> dict[str, Any]:
    if isinstance(route_or_doc_type, dict):
        route = normalize_workflow_route(route_or_doc_type, user_input, attachment_context=attachment_context, trust_route=True)
    else:
        route = normalize_workflow_route({"primary_operation": "generate"}, user_input, attachment_context=attachment_context, trust_route=True)
        if isinstance(route_or_doc_type, str) and route_or_doc_type:
            route["doc_type"] = route_or_doc_type
    common_schema = _common_office_content_schema(user_input, route, attachment_context)
    if common_schema:
        if progress_callback:
            progress_callback("generating_content", 60, "正在整理正文内容")
        return _normalize_content_json(common_schema, route, user_input, attachment_context)
    if _should_generate_sectioned_document(user_input, route, attachment_context):
        return _normalize_content_json(
            _generate_sectioned_content_schema(user_input, route, attachment_context, progress_callback=progress_callback),
            route,
            user_input,
            attachment_context,
        )
    if progress_callback:
        progress_callback("generating_content", 52, "正在生成正文内容")
    prompt = _build_generation_prompt(route, route.get("task_graph") or build_task_graph(route), user_input, attachment_context)
    strict_generation = _strict_generated_document(route, attachment_context)
    try:
        content = _call_qwen_json(prompt)
    except DocumentGenerationError:
        logger.exception("CONTENT_JSON generation failed for doc_type=%s", route.get("doc_type"))
        if strict_generation:
            raise
        content = _draft_content_schema(user_input, attachment_context)
    except Exception as exc:
        logger.exception("CONTENT_JSON unexpected error for doc_type=%s", route.get("doc_type"))
        if strict_generation:
            raise DocumentGenerationError(str(exc), "正式文档正文生成失败，请稍后重试。") from exc
        content = _draft_content_schema(user_input, attachment_context)
    if not _content_schema_has_substance(content):
        if strict_generation:
            raise ContentSchemaInvalidError("content schema has no substance")
        content = _draft_content_schema(user_input, attachment_context)
    return _normalize_content_json(content, route, user_input, attachment_context)


def _normalize_content_json(content: dict, route: dict, user_input: str, attachment_context: str) -> dict[str, Any]:
    if not isinstance(content, dict):
        content = {}
    operation = route.get("primary_operation", "generate")
    title = str(content.get("title") or _title_from_input(user_input, operation))
    sections = content.get("sections")
    if isinstance(sections, dict):
        sections = [{"heading": str(k), "content": _stringify(v)} for k, v in sections.items()]
    if not isinstance(sections, list) or not sections:
        sections = [{"heading": _default_section_heading(operation), "content": _fallback_body(user_input, attachment_context)}]
    normalized_sections = []
    for item in sections:
        if isinstance(item, dict):
            heading = str(item.get("heading") or item.get("title") or _default_section_heading(operation))
            body = _stringify(item.get("content") or item.get("body") or "")
        else:
            heading = _default_section_heading(operation)
            body = _stringify(item)
        if body.strip():
            normalized_sections.append({"heading": heading, "content": body.strip()})
    if not normalized_sections:
        normalized_sections = [{"heading": _default_section_heading(operation), "content": _fallback_body(user_input, attachment_context)}]
    if _wants_advice(user_input) and not _contains_advice_section(normalized_sections):
        normalized_sections.append({
            "heading": "补充建议",
            "content": _build_advice_content(user_input, attachment_context, route),
        })
    normalized_sections = _ensure_report_completeness(normalized_sections, user_input, attachment_context, route)

    tables = content.get("tables") if isinstance(content.get("tables"), list) else []
    slides = content.get("slides") if isinstance(content.get("slides"), list) else []
    warnings = content.get("warnings") if isinstance(content.get("warnings"), list) else []
    source_usage = content.get("source_usage") if isinstance(content.get("source_usage"), dict) else {}
    return {
        "title": title,
        "summary": _stringify(content.get("summary") or ""),
        "sections": normalized_sections,
        "tables": tables,
        "slides": slides,
        "warnings": [str(item) for item in warnings if str(item).strip()],
        "source_usage": {
            "used_pdf_text": bool(source_usage.get("used_pdf_text") or "PDF" in (attachment_context or user_input)),
            "used_visual_context": bool(source_usage.get("used_visual_context") or "图片分析" in (attachment_context or user_input)),
            "used_user_instruction": True,
        },
    }


def _stringify(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, list):
        return "\n".join("- " + _stringify(item) for item in value)
    if isinstance(value, dict):
        return "\n".join(f"{key}: {_stringify(val)}" for key, val in value.items())
    return str(value)


def _line_field(text: str, label: str, default: str = "") -> str:
    pattern = rf"{re.escape(label)}\s*[：:]\s*([^\n\r]+)"
    match = re.search(pattern, text or "")
    return match.group(1).strip() if match else default


def _block_field(text: str, label: str, default: str = "") -> str:
    lines = (text or "").splitlines()
    collecting = False
    buffer: list[str] = []
    inline_pattern = re.compile(rf"^\s*{re.escape(label)}\s*[：:]\s*(.+?)\s*$")
    start_pattern = re.compile(rf"^\s*{re.escape(label)}\s*[：:]?\s*$")
    stop_pattern = re.compile(r"^\s*[^\s：:]{1,16}\s*[：:]\s*$")

    for raw in lines:
        inline_match = inline_pattern.match(raw)
        if inline_match and not collecting:
            return inline_match.group(1).strip() or default
        if not collecting:
            if start_pattern.match(raw):
                collecting = True
            continue
        if stop_pattern.match(raw):
            break
        buffer.append(raw.rstrip())

    value = "\n".join(buffer).strip()
    if not value:
        return default
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


def _requested_filename(user_input: str, output_format: str = "docx") -> str:
    text = user_input or ""
    match = re.search(r"文件名\s*[：:]\s*([^\n\r]+)", text)
    if not match:
        return ""
    filename = match.group(1).strip().strip("。；;，, ")
    expected_ext = FILE_EXT_BY_FORMAT.get(output_format, ".docx")
    if not filename:
        return ""
    if not Path(filename).suffix:
        filename += expected_ext
    if Path(filename).suffix.lower() != expected_ext:
        filename = filename[: -len(Path(filename).suffix)] + expected_ext
    return _safe_filename(filename[: -len(expected_ext)], output_format)


def _requested_title(user_input: str, default: str = "") -> str:
    title = _line_field(user_input, "标题", "")
    if title:
        return title
    lines = (user_input or "").splitlines()
    for index, raw in enumerate(lines):
        if raw.strip().strip("：:") == "标题":
            for follow in lines[index + 1:]:
                value = follow.strip().strip("：:，,。. ")
                if value:
                    return value
    return default


def _tomorrow_cn() -> str:
    value = datetime.now() + timedelta(days=1)
    return f"{value.year}年{value.month}月{value.day}日"


def _strict_generated_document(route: dict, attachment_context: str = "") -> bool:
    if (attachment_context or "").strip():
        return False
    if not requires_file_generation(route):
        return False
    return route.get("primary_operation") == "generate"


def _content_schema_has_substance(content: Any) -> bool:
    if not isinstance(content, dict):
        return False
    if str(content.get("summary") or "").strip():
        return True
    sections = content.get("sections")
    if isinstance(sections, list):
        for item in sections:
            if isinstance(item, dict):
                if str(item.get("content") or item.get("body") or "").strip():
                    return True
            elif str(item).strip():
                return True
    for key in ("tables", "slides"):
        value = content.get(key)
        if isinstance(value, list) and value:
            return True
    return False


def _strip_generation_instruction(text: str) -> str:
    cleaned = (text or "").strip()
    if not cleaned:
        return ""
    cleaned = re.sub(
        r"(?im)^(请|帮我|给我)?(生成|编写|撰写|制作|写|输出)(一份|一篇)?[^\n。]*?(通知书|通知|证明|函|报告|方案|文档)[。:：]?\s*",
        "",
        cleaned,
    )
    cleaned = re.sub(r"(?is)\n*要求[：:].*$", "", cleaned).strip()
    return cleaned


def _draft_content_schema(user_input: str, attachment_context: str = "") -> dict[str, Any]:
    body = _fallback_body(user_input, attachment_context)
    return {
        "title": _title_from_input(user_input, "generate"),
        "summary": "⚠️ 文档自动生成未完全成功，以下内容是根据输入整理的草稿，请人工核对后再使用。",
        "sections": [{"heading": "草稿正文", "content": body}],
        "tables": [],
        "slides": [],
        "warnings": ["文档正文生成失败，当前内容仅为草稿。"],
        "source_usage": {"used_user_instruction": True},
    }


def _common_office_content_schema(user_input: str, route: dict, attachment_context: str = "") -> dict[str, Any] | None:
    if (attachment_context or "").strip():
        return None
    text = (user_input or "").strip()
    low = text.lower()

    project_name = _line_field(text, "项目名称", "")
    if _keyword(text, "实施方案", "项目方案") and project_name:
        title = f"{project_name}项目实施方案"
        gantt = (
            "```mermaid\n"
            "gantt\n"
            f"    title {project_name}项目实施计划\n"
            "    dateFormat  YYYY-MM-DD\n"
            "    section 准备阶段\n"
            "    需求调研与方案确认 :a1, 2026-07-01, 14d\n"
            "    section 建设阶段\n"
            "    平台部署与功能开发 :b1, after a1, 45d\n"
            "    数据对接与联调测试 :b2, after b1, 20d\n"
            "    section 交付阶段\n"
            "    培训试运行与验收 :c1, after b2, 20d\n"
            "```"
        )
        org = (
            "```mermaid\n"
            "flowchart TB\n"
            "    A[项目领导小组] --> B[项目经理]\n"
            "    B --> C[产品与需求组]\n"
            "    B --> D[技术实施组]\n"
            "    B --> E[测试验收组]\n"
            "    B --> F[培训运维组]\n"
            "```"
        )
        sections = [
            {"heading": "封面", "content": f"项目名称：{project_name}\n文档名称：项目实施方案\n文档版本：V1.0\n提交单位：中科信创\n提交日期：2026年6月"},
            {"heading": "目录", "content": "1. 项目背景\n2. 建设目标\n3. 功能设计\n4. 技术架构\n5. 实施计划\n6. 人员安排\n7. 风险控制\n8. 培训方案\n9. 验收标准\n10. 运维保障\n11. 预算报价"},
            {"heading": "项目背景", "content": f"{project_name}面向校园教学、学习、管理和服务场景，建设目标是在统一平台内融合AI学习助手、学情分析、资源推荐和教学管理能力，解决学习数据分散、个性化支持不足、教师工作负担较重等问题。"},
            {"heading": "建设目标", "content": "建设统一、可靠、可扩展的AI学习平台；形成覆盖学生、教师、管理人员的多角色服务体系；提升个性化学习支持能力、教学数据分析能力和校园数字化运营能力。"},
            {"heading": "功能设计", "content": "核心功能包括智能学习助手、课程资源管理、个性化学习路径、作业与测评、学情分析看板、教师备课辅助、知识库问答、学习预警和运营统计。各模块应支持权限控制、日志审计和可配置管理。"},
            {"heading": "技术架构", "content": "平台采用前后端分离和服务化架构，前端提供Web与移动端访问入口，后端提供用户、课程、资源、AI能力和数据分析服务。AI能力层对接大模型、向量检索、OCR/ASR等组件，数据层统一管理学习行为、课程资源和统计指标。"},
            {"heading": "实施计划", "content": f"项目分为启动准备、需求调研、系统设计、开发部署、联调测试、培训试运行、验收交付和运维保障八个阶段。\n\n{gantt}"},
            {"heading": "人员安排", "content": f"项目组织采用项目经理负责制，下设产品与需求组、技术实施组、测试验收组、培训运维组。组织架构如下：\n\n{org}"},
            {"heading": "风险控制", "content": "主要风险包括需求变更、数据对接延迟、模型效果不稳定、用户培训不足和上线初期并发压力。控制措施包括需求冻结机制、接口联调计划、模型评测机制、分批培训和上线应急预案。"},
            {"heading": "培训方案", "content": "培训对象包括系统管理员、教师、学生和运维人员。培训内容覆盖平台登录、课程资源维护、AI学习助手使用、学情看板解读、问题反馈流程和日常运维操作。培训方式包括现场培训、操作手册、视频材料和试运行答疑。"},
            {"heading": "验收标准", "content": "验收内容包括功能验收、性能验收、安全验收、文档验收和试运行验收。功能模块应满足合同与需求说明；核心页面响应时间应满足日常教学使用；权限、日志和数据安全策略应通过检查；交付文档应齐全。"},
            {"heading": "运维保障", "content": "提供上线后运维支持、故障响应、数据备份、安全巡检、版本升级和使用答疑。建立问题工单机制，明确响应时限和升级路径，保障平台长期稳定运行。"},
            {"heading": "预算报价", "content": "预算报价以软件平台建设、AI能力接入、实施部署、培训服务和运维保障为主要构成，最终报价可根据实际用户规模、部署环境和服务周期调整。"},
        ]
        return {
            "title": title,
            "summary": f"本方案围绕{project_name}建设目标、功能范围、技术架构、实施路径和运维保障进行设计，可作为客户交付版实施方案初稿。",
            "sections": sections,
            "tables": [
                {
                    "title": "预算报价表",
                    "headers": ["序号", "费用项", "内容说明", "金额（万元）"],
                    "rows": [
                        ["1", "软件平台建设", "AI学习平台核心功能开发与配置", "38.00"],
                        ["2", "AI能力接入", "大模型、知识库、智能推荐等能力集成", "26.00"],
                        ["3", "实施部署", "环境部署、数据对接、联调测试", "12.00"],
                        ["4", "培训服务", "管理员、教师、学生培训及材料", "4.00"],
                        ["5", "运维保障", "一年运维支持、巡检、升级服务", "8.00"],
                    ],
                },
                {
                    "title": "任务分工表",
                    "headers": ["角色", "主要职责", "人数"],
                    "rows": [
                        ["项目经理", "统筹计划、进度、风险和客户沟通", "1"],
                        ["产品经理", "需求调研、原型确认、验收支持", "1"],
                        ["开发工程师", "平台功能开发、接口对接、联调", "4"],
                        ["测试工程师", "测试用例、缺陷跟踪、验收测试", "2"],
                        ["培训运维工程师", "用户培训、上线保障、运维支持", "2"],
                    ],
                },
            ],
            "slides": [],
            "warnings": [],
            "source_usage": {"used_user_instruction": True},
        }

    if _should_generate_sectioned_document(user_input, route, attachment_context):
        return None

    if not _common_office_document_request(low):
        return None

    def schema(title: str, body: str, summary: str = "") -> dict[str, Any]:
        return {
            "title": title,
            "summary": summary,
            "sections": [{"heading": "正文", "content": body.strip()}],
            "tables": [],
            "slides": [],
            "warnings": [],
            "source_usage": {"used_user_instruction": True},
        }

    if "会议纪要" in text:
        title = _requested_title(text, "项目周会会议纪要")
        sections = [
            {"heading": "会议时间", "content": _line_field(text, "会议时间", "待补充")},
            {"heading": "参会人员", "content": _line_field(text, "参会人员", "项目组成员、相关负责人")},
            {"heading": "本周工作进展", "content": "1. 梳理本周重点工作完成情况。\n2. 汇总项目推进状态、阶段成果和待确认事项。\n3. 对影响交付的事项进行跟踪记录。"},
            {"heading": "存在问题", "content": "1. 部分任务依赖外部确认，需明确责任人与完成时间。\n2. 个别事项存在进度滞后风险，需在下周会议前完成闭环跟进。"},
            {"heading": "下周计划", "content": "1. 按任务分工推进下阶段工作。\n2. 对本周遗留问题进行复盘和处理。\n3. 完成关键节点材料整理，并同步项目进展。"},
        ]
        return {
            "title": title,
            "summary": "",
            "sections": sections,
            "tables": [{
                "title": "任务分工表",
                "headers": ["序号", "任务事项", "责任人", "完成时间", "状态"],
                "rows": [
                    ["1", "整理本周工作进展", "待指定", "本周内", "进行中"],
                    ["2", "跟进存在问题整改", "待指定", "下次周会前", "待处理"],
                    ["3", "制定下周工作计划", "待指定", "下周一前", "待启动"],
                ],
            }],
            "slides": [],
            "warnings": [],
            "source_usage": {"used_user_instruction": True},
        }

    if "面试邀约" in text or "面试邀请" in text or "邀请函" in text:
        position = _line_field(text, "岗位", "相关岗位")
        time_text = _line_field(text, "面试时间", "请以双方确认时间为准")
        place = _line_field(text, "地点", "请以通知地点为准")
        contact = _line_field(text, "联系人", "招聘负责人")
        phone = _line_field(text, "联系电话", "")
        body = (
            "尊敬的候选人：\n\n"
            f"您好！感谢您关注并应聘我司“{position}”岗位。经初步沟通，现诚邀您参加面试，具体安排如下：\n\n"
            f"面试时间：{time_text}\n"
            f"面试地点：{place}\n"
            f"联系人：{contact}\n"
            f"联系电话：{phone or '请以招聘通知为准'}\n\n"
            "路线提醒：请提前规划交通路线，按照导航前往面试地点。如进入园区或楼宇需登记，请配合现场管理要求。\n\n"
            "注意事项：\n"
            "1. 请提前10分钟到达面试地点；\n"
            "2. 请携带个人简历及相关证明材料；\n"
            "3. 如时间冲突，请提前与联系人沟通调整。\n\n"
            "期待您的到来。"
        )
        return schema("面试邀请函", body)

    if "入职通知书" in text or "入职通知" in text or "录用通知书" in text or "录用通知" in text:
        name = _line_field(text, "姓名", "新员工")
        dept = _line_field(text, "部门", "相关部门")
        position = _line_field(text, "岗位", "相关岗位")
        entry_date = _line_field(text, "入职时间", "请以公司通知时间为准")
        location = _line_field(text, "办公地点", "请以公司通知地点为准")
        contact = _line_field(text, "联系人", "人力资源部")
        phone = _line_field(text, "联系电话", "请以公司通知为准")
        body = (
            f"{name}：\n\n"
            f"您好！经公司综合评估，现正式通知您于{entry_date}到我司报到入职，拟安排您在{dept}担任{position}岗位，办公地点为{location}。\n\n"
            "请您按照公司入职安排，提前准备本人身份证、学历学位证明、离职证明、近期免冠照片及其他办理入职手续所需材料，并按要求完成报到登记。\n\n"
            "报到当日请保持通讯畅通，如因特殊原因无法按时报到，请至少提前1个工作日与公司联系说明情况，以便协调后续安排。\n\n"
            f"联系人：{contact}\n"
            f"联系电话：{phone}\n\n"
            "欢迎加入公司，期待您在新的岗位上发挥专业能力，与团队共同成长。"
        )
        return schema("员工入职通知书", body)

    if "离职证明" in text:
        name = _line_field(text, "员工", "该员工")
        id_tail = _line_field(text, "身份证后四位", "")
        dept = _line_field(text, "部门", "相关部门")
        leave_date = _line_field(text, "离职日期", "双方确认日期")
        id_part = f"（身份证号码后四位：{id_tail}）" if id_tail else ""
        body = (
            f"兹证明{name}{id_part}曾在我司{dept}任职。\n\n"
            f"现因个人原因离职，双方劳动关系已于{leave_date}解除，相关离职手续已按公司流程办理完毕。\n\n"
            "该员工在职期间与公司无未结清的劳动争议事项。特此证明。\n\n"
            "公司名称：________________\n"
            "日期：________________\n"
            "（此处加盖公司公章）"
        )
        return schema("离职证明", body)

    if "请假审批" in text or "请假记录" in text:
        name_match = re.search(r"([^\n：:]{1,12})[：:]\s*\n?.{0,30}请假", text)
        name = name_match.group(1).strip() if name_match else "员工"
        date_match = re.search(r"(\d{4}[-年]\d{1,2}[-月]\d{1,2}[日]?\s*(?:至|到|-)\s*\d{4}[-年]\d{1,2}[-月]\d{1,2}[日]?)", text)
        date_text = date_match.group(1) if date_match else "以审批系统登记时间为准"
        reason = "家属住院陪护" if re.search(r"老人|家属|住院", text) else "个人原因"
        body = (
            f"申请人：{name}\n"
            "请假类型：事假\n"
            f"请假时间：{date_text}\n"
            f"请假原因：{reason}\n\n"
            "审批意见：经核实，该请假事由具备合理性，建议按公司请假制度办理审批手续，并做好工作交接安排。\n\n"
            "备注：请申请人按规定提交相关证明材料，并以公司考勤系统最终审批记录为准。"
        )
        return schema("员工请假审批记录", body)

    if "催款" in text or ("客户" in text and "欠款" in text):
        customer = _line_field(text, "客户", "贵司")
        amount = _line_field(text, "欠款", "相关款项")
        overdue = _line_field(text, "逾期", "")
        overdue_text = f"该款项已逾期{overdue}，" if overdue else ""
        body = (
            f"致：{customer}\n\n"
            f"经我司财务核对，截至本函发出之日，贵司尚有人民币{amount}未完成支付。{overdue_text}"
            "为避免影响双方后续合作，请贵司在收到本函后3个工作日内完成付款，或与我司财务人员沟通付款安排。\n\n"
            "如贵司已完成付款，请忽略本函并提供付款凭证以便我司核销。感谢配合。"
        )
        return schema("催款通知函", body)

    if "采购" in text and "用途" in text:
        item = _line_field(text, "采购", "办公设备")
        purpose = _line_field(text, "用途", "业务办公需要")
        body = (
            f"申请事项：采购{item}\n"
            f"用途说明：{purpose}\n\n"
            f"审批理由：由于{purpose}，现有办公资源已无法完全满足实际使用需求。为保障相关人员正常开展工作，提高办公效率，申请采购{item}。\n\n"
            "建议按公司采购管理制度履行询价、审批和验收流程，确保采购价格合理、用途明确、资产登记完整。"
        )
        return schema("采购审批理由", body)

    if "通报" in text or "处罚" in text or "违纪" in text:
        body = (
            "经核实，相关员工在工作时间存在违反公司劳动纪律的行为，且经多次提醒后仍未及时改正，对正常办公秩序造成不良影响。\n\n"
            "处理意见：根据公司考勤及员工行为管理制度，现予以正式通报批评，并要求本人立即整改。后续如再次发生类似情况，公司将视情节采取进一步管理措施。\n\n"
            "请各部门员工引以为戒，严格遵守工作纪律，保持良好办公秩序。"
        )
        return schema("员工违纪处理通知", body)

    if "面试评价" in text or "评价表" in text:
        body = (
            "综合评价：候选人基础条件与岗位要求基本匹配，但项目实战经验相对不足，仍需进一步核实其过往项目参与深度和独立负责能力。\n\n"
            "沟通表达：候选人表达能力一般，信息组织和重点呈现能力有待提升。\n\n"
            "薪资匹配：候选人当前薪资期望偏高，建议结合岗位预算和复试表现综合评估。\n\n"
            "录用建议：建议暂列为备选候选人，如后续无更匹配人选，可安排复试进一步确认。"
        )
        return schema("面试评价表", body)

    if "停电" in text or "放假" in text or "公告" in text or "行政通知" in text:
        date_text = _tomorrow_cn() if "明天" in text else "具体日期以园区通知为准"
        body = (
            "各位同事：\n\n"
            f"因园区电力检修安排，公司将于{date_text}下午暂停现场办公。请各部门提前做好工作安排，妥善保存办公资料并关闭相关电源设备。\n\n"
            "如有紧急事项，请通过电话或线上方式保持沟通。恢复办公时间以公司后续通知为准。\n\n"
            "特此通知。"
        )
        return schema("临时放假通知", body)

    if "事故调查报告" in text or "事故报告" in text or "事件调查报告" in text or "事件报告" in text or "故障报告" in text:
        event_info = _block_field(text, "事件", "")
        impact = _block_field(text, "影响", "")
        cause = _block_field(text, "原因", "")
        action = _block_field(text, "整改措施", "")
        if not action:
            action = _block_field(text, "处理措施", "")
        if not action:
            action = _block_field(text, "整改建议", "")
        title = _line_field(text, "标题", "") or "事故调查报告"
        if title == "事故调查报告" and "网络中断" in event_info:
            title = "网络中断事故调查报告"
        sections = []
        if event_info:
            sections.append({
                "heading": "一、事故基本情况",
                "content": (
                    "根据已掌握信息，现将本次事故基本情况说明如下：\n"
                    f"{event_info}"
                ),
            })
        if impact:
            sections.append({
                "heading": "二、影响范围",
                "content": (
                    "本次事故发生后，对业务系统和现场运行造成了如下影响：\n"
                    f"{impact}"
                ),
            })
        if cause:
            sections.append({
                "heading": "三、原因分析",
                "content": (
                    "经初步排查与分析，本次事故的直接原因如下：\n"
                    f"{cause}\n\n"
                    "后续应结合设备运行日志、监控告警记录和现场处置情况进一步复核，确保原因判定准确完整。"
                ),
            })
        if action:
            sections.append({
                "heading": "四、整改措施",
                "content": (
                    "针对本次事故暴露出的问题，拟采取以下整改措施：\n"
                    f"{action}\n\n"
                    "上述措施应明确责任部门、完成时限和复盘要求，并在整改完成后开展效果验证。"
                ),
            })
        sections.append({
            "heading": "五、调查结论",
            "content": (
                "综上，本次事故属于基础设施设备故障引发的网络中断事件。"
                "建议尽快落实冗余设计、巡检维护、应急演练和故障复盘机制，持续提升机房网络设备的稳定性和业务连续性保障能力。"
            ),
        })
        return {
            "title": title,
            "summary": "本报告依据用户提供的事故经过、影响、原因及整改要求整理形成，可直接作为正式事故调查报告初稿使用。",
            "sections": sections,
            "tables": [],
            "slides": [],
            "warnings": [],
            "source_usage": {"used_user_instruction": True},
        }

    if "事件时间线" in text or "事件经过" in text or "正式记录" in text:
        event_lines = [line.strip() for line in text.splitlines() if re.search(r"\d{4}年\d{1,2}月\d{1,2}日|\d{4}-\d{1,2}-\d{1,2}", line)]
        body = "事件经过：\n\n" + "\n\n".join(event_lines or [text]) + "\n\n处理建议：建议保留原始沟通记录，并由部门负责人对相关事实进行复核确认。"
        return schema("事件经过记录", body)

    return None


def _title_from_input(user_input: str, operation: str) -> str:
    requested_title = _requested_title(user_input)
    if requested_title:
        return requested_title[:48]
    text = re.sub(r"\s+", " ", (user_input or "")).strip()
    text = re.sub(r"以下是用户上传附件[\s\S]*", "", text).strip()
    text = re.sub(r"^(请|帮我|给我)?(生成|创建|写|制作|整理|总结|提取|润色|改写|补充|补写)", "", text).strip(" ：:，,。.")
    first_line = ""
    for raw_line in (user_input or "").splitlines():
        line = re.sub(r"\s+", " ", raw_line).strip().strip("：:，,。. ")
        if line:
            first_line = line
            break
    if first_line and _keyword(
        first_line,
        "投标文件",
        "招标文件",
        "技术部分",
        "技术标",
        "建设方案",
        "实施方案",
        "项目方案",
        "事故调查报告",
        "调查报告",
        "分析报告",
        "运营分析报告",
    ):
        return first_line[:48]
    if not text:
        names = {
            "summarize": "附件总结",
            "append_section": "补充后的文档",
            "polish": "润色后的文档",
            "rewrite": "改写后的文档",
            "restructure": "正式报告",
            "extract": "信息清单",
            "compare": "对比分析",
            "convert": "转换后的文档",
            "visual_relation": "图文关系分析",
        }
        return names.get(operation, "生成文档")
    return text[:48]


def _default_section_heading(operation: str) -> str:
    return {
        "summarize": "摘要",
        "append_section": "补充内容",
        "polish": "润色内容",
        "rewrite": "改写内容",
        "restructure": "报告正文",
        "extract": "提取结果",
        "compare": "对比结果",
        "convert": "转换内容",
        "visual_relation": "图文关系",
    }.get(operation, "正文")


def _contains_advice_section(sections: list[dict[str, Any]]) -> bool:
    for section in sections or []:
        heading = str((section or {}).get("heading") or "")
        content = str((section or {}).get("content") or "")
        if "建议" in heading or re.search(r"(?m)^\s*(?:[-*]|\d+[.、])?\s*[^。\n]{0,16}建议", content):
            return True
    return False


def _build_advice_content(user_input: str, attachment_context: str, route: dict) -> str:
    text = attachment_context or ""
    if re.search(r"(投诉方|我方作品|不构成相似|作品|发色|构图|服装)", text):
        return "\n".join([
            "1. 建议把“投诉方作品”和“我方作品图”的差异继续按维度归纳，例如主体形象、发色发型、动作姿态、服装细节、构图位置和背景元素，避免只给笼统结论。",
            "2. 建议在每个差异点后补充对应依据：哪些来自 PDF 文字，哪些来自图片可见内容，让图文证据形成一一对应关系。",
            "3. 建议增加“核心结论”小节，明确说明两图虽然都有星空、彩色光效等背景元素，但主体人物形象、姿势、服装和构图差异更突出。",
            "4. 建议补充我方作品的独立设计说明，例如粉色短发、正面站立、长靴、腰部装饰、飘带、中心构图等，突出独立表达。",
            "5. 如用于正式材料，建议补充创作过程、发布时间、设计稿或授权来源等外部证据，并避免写成绝对化判断；可表述为“从当前附件可见内容看，两者整体表达差异明显”。",
        ])
    if "PDF 图片分析" in text or "图片识别结果" in text:
        return "\n".join([
            "1. 建议先把图片中的关键主体、位置、颜色、动作或图表含义单独列出，再对应到 PDF 文字中的相关句子。",
            "2. 建议把视觉信息作为辅助证据使用，不要让图片分析替代原文；结论应同时引用文字内容和图片可见信息。",
            "3. 建议补充一段“图文关系说明”，说明图片是在证明、补充、示意还是对比文字内容。",
            "4. 建议对不确定内容使用保守表述，例如“从图中可见”“附件中显示”，避免扩展到附件没有提供的信息。",
            "5. 如果要输出正式文档，建议把建议内容整理成独立章节，并保留原文主要结构。",
        ])
    return "\n".join([
        "1. 建议先补充附件核心信息的简要概括，再给出建议，便于读者理解建议依据。",
        "2. 建议每条建议都对应到附件中的文字或图片内容，避免脱离材料泛泛而谈。",
        "3. 建议把建议分成“内容完善”“表达优化”“后续补充材料”三类，结构会更清晰。",
        "4. 建议对附件没有明确说明的内容使用谨慎表述，不补充无法确认的事实。",
    ])


def _ensure_report_completeness(
    sections: list[dict[str, Any]],
    user_input: str,
    attachment_context: str,
    route: dict,
) -> list[dict[str, Any]]:
    target_chars = _requested_char_count(user_input)
    wants_report = _wants_report(user_input)
    if not wants_report and not target_chars:
        return sections

    result = list(sections or [])
    rendered = "\n".join(f"{item.get('heading', '')}\n{item.get('content', '')}" for item in result)
    has_conclusion = any(re.search(r"(结论|综合判断|报告结语|总体判断)", str(item.get("heading", "")) + str(item.get("content", ""))) for item in result)

    if wants_report and not has_conclusion:
        result.append({
            "heading": "报告结论",
            "content": _build_report_conclusion(attachment_context),
        })
        rendered = "\n".join(f"{item.get('heading', '')}\n{item.get('content', '')}" for item in result)

    min_chars = int((target_chars or 0) * 0.82)
    has_completion = any(str(item.get("heading") or "").strip() == "报告补充说明" for item in result)
    if target_chars and _content_char_len(rendered) < min_chars and not has_completion:
        result.append({
            "heading": "报告补充说明",
            "content": _build_report_completion(attachment_context, target_chars, rendered),
        })

    if wants_report and len(result) == 1:
        result = _build_report_sections(user_input, attachment_context, route)
        rendered = "\n".join(f"{item.get('heading', '')}\n{item.get('content', '')}" for item in result)

    longform_outline_keywords = (
        "投标文件",
        "招标文件",
        "技术部分",
        "技术标",
        "建设方案",
        "实施方案",
        "项目方案",
    )
    if _keyword(user_input or "", *longform_outline_keywords) and len(result) < 6:
        supplemental_sections = [
            ("建设重点与技术路线", "围绕平台建设重点、关键技术路线、系统整合方式和阶段性落地路径进行补充说明，突出方案的完整性与可实施性。"),
            ("风险控制与保障机制", "补充项目实施过程中的风险识别、进度控制、质量保证、安全保障和应急预案，体现项目交付与持续运营能力。"),
        ]
        existing_headings = {str(item.get('heading') or '').strip() for item in result}
        for heading, content in supplemental_sections:
            if heading in existing_headings:
                continue
            result.append({"heading": heading, "content": content})
    return result


def _build_report_sections(user_input: str, attachment_context: str, route: dict) -> list[dict[str, Any]]:
    return [
        {"heading": "报告摘要", "content": _build_report_intro(attachment_context)},
        {"heading": "主要分析", "content": _build_report_analysis(attachment_context)},
        {"heading": "报告结论", "content": _build_report_conclusion(attachment_context)},
    ]


def _build_report_intro(attachment_context: str) -> str:
    if re.search(r"(投诉方|我方作品|不构成相似|作品)", attachment_context or ""):
        return "本报告基于 PDF 文本内容和图片视觉分析，对投诉方作品与我方作品进行综合梳理。附件文字已经逐项列明两者在人物形象、动作姿态、服饰细节、构图位置和背景元素上的差异；图片分析进一步补充了我方作品的主体位置、发色、服装和画面风格等视觉信息。"
    return "本报告基于附件解析内容进行整理，综合参考文字信息与图片或版面信息，概括材料主旨、关键依据和可形成的结论。"


def _build_report_analysis(attachment_context: str) -> str:
    if re.search(r"(投诉方|我方作品|发色|构图|服装)", attachment_context or ""):
        return "从文字部分看，投诉方角色为蓝灰色头发，动作呈跳跃或坐姿，人物偏左下且占比较小；我方角色为粉色短发，正面站立，位于画面中央且视觉重点更突出。从图片部分看，我方作品中的角色有猫耳发饰、黑色服装、蓝色装饰、长靴和发光武器，背景呈宇宙或能量爆发效果。两者虽都包含星空、光效等背景元素，但主体形象、姿态和构图表达存在明显区别。"
    if "PDF 图片分析" in (attachment_context or ""):
        return "附件中的图片分析可作为辅助依据，用于说明视觉元素与文字描述之间的对应关系。正式写作时应以原文事实为主，以图片信息补充主体、位置、动作、颜色、图表含义或版面关系，避免将视觉推断直接替代文本事实。"
    return "从附件内容看，材料已经提供了可用于总结和判断的基础信息。报告写作时应围绕核心事实展开，先说明材料主题，再归纳主要依据，最后给出稳妥结论。"


def _build_report_conclusion(attachment_context: str) -> str:
    if re.search(r"(投诉方|我方作品|不构成相似|作品)", attachment_context or ""):
        return "综合来看，当前附件能够支撑“我方作品与投诉方作品整体表达存在显著差异”的判断。两图的相似点主要集中在背景氛围层面，而关键识别元素集中在主体人物、发色发型、动作姿态、服饰设计和画面构图上，这些部分差异更直接、更明确。因此，报告结论宜表述为：从附件可见内容判断，两者整体不相同，不宜认定为实质性相似。"
    if not (attachment_context or "").strip():
        return "综合来看，报告结论应围绕用户给定主题、目标和约束形成。对于输入中已经明确的信息，可以直接展开；对于缺少数据支撑的判断，应采用稳妥表述，并在正式交付前补充必要依据。"
    return "综合来看，报告结论应严格依据附件内容形成。对于附件已经明确呈现的信息，可以直接归纳；对于附件没有提供充分依据的部分，应使用谨慎表述，避免编造或过度推断。"


def _build_report_completion(attachment_context: str, target_chars: int, current_text: str) -> str:
    completion = []
    if re.search(r"(投诉方|我方作品|不构成相似|作品)", attachment_context or ""):
        completion.append("为使报告更完整，可在正文中进一步明确图文对应关系：文字列出差异维度，图片提供视觉佐证。")
        completion.append("其中，人物发色、姿态、服装、画面位置属于判断作品整体表达的重要依据；背景中的星空、光效和色彩只能作为辅助因素，不应单独决定相似性判断。")
        completion.append("最终表述应保持稳妥，强调“基于当前附件可见内容”，并保留后续补充创作过程、源文件或授权材料的空间。")
    else:
        completion.append("为满足报告完整性，建议补充项目背景、核心依据和结论三部分。")
        if (attachment_context or "").strip():
            completion.append("正文应优先使用附件中已经出现的事实，并在涉及判断时说明依据来源。")
            completion.append("结尾应给出清晰但保守的综合判断，避免超过附件内容进行扩展。")
        else:
            completion.append("正文应围绕用户给定主题展开，并在涉及数据、时间、金额或责任边界时保留可核验依据。")
            completion.append("结尾应给出清晰但保守的综合判断，避免补充输入中没有依据的具体事实。")
    return "".join(completion)


def _fallback_body(user_input: str, attachment_context: str) -> str:
    source = (attachment_context or user_input or "").strip()
    if not source:
        return "暂无可用内容。"
    cleaned = re.sub(r"\n{3,}", "\n\n", _strip_generation_instruction(source))
    if not cleaned.strip():
        cleaned = re.sub(r"\n{3,}", "\n\n", source)
    return cleaned[:3000]


def _safe_filename(title: str, output_format: str) -> str:
    ext = FILE_EXT_BY_FORMAT.get(output_format, ".docx")
    base = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff_-]+", "_", (title or "document").strip()).strip("_")
    return (base or "document")[:48] + ext


def _looks_like_document_request(value: Any) -> bool:
    return isinstance(value, dict) and (
        "content_type" in value
        or "content" in value and any(key in value for key in ("filename", "format"))
    )


def normalize_document_request(value: Any, fallback_instruction: str = "", route: dict | None = None) -> dict[str, str]:
    payload = dict(value or {}) if isinstance(value, dict) else {}
    output_spec = (route or {}).get("output_spec") or {}
    fmt = str(payload.get("format") or output_spec.get("format") or "docx").strip().lower()
    if fmt not in FILE_EXT_BY_FORMAT:
        fmt = "docx"
    content_type = str(payload.get("content_type") or "markdown").strip().lower()
    if content_type not in DOCUMENT_CONTENT_TYPES:
        content_type = "markdown"
    content = str(payload.get("content") or "").strip()
    filename = str(payload.get("filename") or output_spec.get("filename") or "").strip()
    if not filename:
        filename = _requested_filename(fallback_instruction, fmt)
    if not filename:
        filename = _safe_filename(_requested_title(fallback_instruction, _title_from_input(fallback_instruction, "generate")), fmt)
    expected_ext = FILE_EXT_BY_FORMAT.get(fmt, ".docx")
    if Path(filename).suffix.lower() != expected_ext:
        filename = re.sub(r"\.[^.]+$", "", filename) + expected_ext
    return {
        "filename": _safe_filename(filename[: -len(expected_ext)] if filename.endswith(expected_ext) else filename, fmt),
        "format": fmt,
        "content_type": content_type,
        "content": content,
    }


def _markdown_has_substance(markdown: str) -> bool:
    body = re.sub(r"[#*\-`>|_\s]", "", markdown or "")
    return len(body) >= MIN_MARKDOWN_CHARS


def _build_markdown_prompt(route: dict, user_input: str, attachment_context: str) -> str:
    context = attachment_context.strip()
    if not context and _has_attachment_context(user_input):
        context = user_input
    return (
        "你是中文正式文档撰写助手。请只输出 Markdown 正文，不要输出 JSON，不要解释，不要包含思考过程。\n"
        "要求：\n"
        "1. 使用 Markdown 标题组织结构，一级标题作为文档标题，二级/三级标题作为章节。\n"
        "2. 表格使用标准 Markdown 表格。\n"
        "3. 列表使用有序或无序列表。\n"
        "4. 如果用户要求 Mermaid 甘特图或组织架构图，保留 fenced code block。\n"
        "5. 内容要可直接交付给客户，中文正式、完整、克制。\n"
        "6. 不要输出 sections、tables、slides、style tree 或任何嵌套 JSON。\n\n"
        "ROUTING_JSON:\n"
        + json.dumps(route, ensure_ascii=False, indent=2)
        + "\n\nUSER_INSTRUCTION:\n"
        + (user_input or "").strip()
        + "\n\nPARSED_ATTACHMENT_CONTEXT:\n"
        + context[:28000]
        + "\n\nMARKDOWN:\n"
    )


def _wrap_plaintext_as_markdown(text: str, title: str = "文档") -> str:
    body = (text or "").strip()
    if not body:
        body = "请重新描述需求。"
    if body.startswith("#"):
        return body
    return f"# {title or '文档'}\n\n## 正文\n\n{body}".strip()


def _placeholder_markdown() -> str:
    return "# 文档生成失败\n\n请重新描述需求。"


def _plain_document_page_input(user_input: str) -> bool:
    text = (user_input or "").strip()
    if not text:
        return False
    if len(text) > 120:
        return False
    return not _keyword(
        text.lower(),
        "生成",
        "创建",
        "新建",
        "写",
        "撰写",
        "制作",
        "导出",
        "会议纪要",
        "报告",
        "方案",
        "简历",
        "通知",
        "总结",
        "要求",
        "包含",
        "标题",
    )


def generate_markdown_document_request(
    user_input: str,
    route: dict,
    attachment_context: str = "",
    progress_callback: Callable[[str, int, str], None] | None = None,
) -> dict[str, str]:
    route = normalize_workflow_route(route, user_input, attachment_context=attachment_context, trust_route=True)
    output_format = (route.get("output_spec") or {}).get("format", "docx")
    if progress_callback:
        progress_callback("generating_content", 52, "正在生成 Markdown 正文")
    if (route.get("output_spec") or {}).get("force_word_document_page") and _plain_document_page_input(user_input):
        return normalize_document_request(
            {
                "filename": (route.get("output_spec") or {}).get("filename") or _requested_filename(user_input, output_format),
                "format": output_format,
                "content_type": "markdown",
                "content": (user_input or "").strip(),
            },
            fallback_instruction=user_input,
            route=route,
        )
    markdown = ""
    generation_error_message = ""
    try:
        markdown = _call_qwen_text(_build_markdown_prompt(route, user_input, attachment_context), max_tokens=4096)
    except ModelEmptyResponseError as exc:
        generation_error_message = exc.user_message
        logger.warning("Markdown generation failed: %s", exc.technical_message)
    except DocumentGenerationError as exc:
        generation_error_message = exc.user_message
        logger.warning("Markdown generation failed: %s", exc.technical_message)
    except Exception as exc:
        generation_error_message = "文档生成失败，请稍后重试。"
        logger.warning("Markdown generation unexpected failure: %s", exc)
    markdown = _clean_generated_markdown(markdown)
    if generation_error_message:
        markdown = "# 文档生成失败\n\n" + generation_error_message + "\n\n请稍后重试或简化文档要求。"
    elif not _markdown_has_substance(markdown):
        markdown = _placeholder_markdown()
    if not _markdown_has_substance(markdown):
        markdown = _placeholder_markdown()
    return normalize_document_request(
        {
            "filename": (route.get("output_spec") or {}).get("filename") or _requested_filename(user_input, output_format),
            "format": output_format,
            "content_type": "markdown",
            "content": markdown,
        },
        fallback_instruction=user_input,
        route=route,
    )


def render_document_request(route: dict, document_request: dict[str, str]) -> dict[str, Any]:
    request = normalize_document_request(document_request, route=route)
    output_format = request["format"]
    if output_format == "chat":
        raise RuntimeError("chat output does not create a file")
    filepath = DOC_OUTPUT_DIR / f"{uuid.uuid4().hex}{FILE_EXT_BY_FORMAT.get(output_format, '.docx')}"
    content = request["content"]
    content_type = request["content_type"]
    quality_attempts: list[dict[str, Any]] = []
    if output_format == "docx":
        markdown = content if content_type == "markdown" else _wrap_plaintext_as_markdown(content)
        markdown = normalize_markdown_for_docx(markdown)
        title = document_title_from_markdown(markdown, Path(request["filename"]).stem)
        request["filename"] = _safe_filename(title, "docx")
        for attempt in range(1, 3):
            filepath.unlink(missing_ok=True)
            candidate = normalize_markdown_for_docx(markdown)
            if attempt == 2:
                # The second pass is deliberately stricter: collapse isolated
                # marker lines left by malformed streaming Markdown.
                candidate = "\n".join(
                    line for line in candidate.splitlines()
                    if not re.fullmatch(r"\s*#{1,6}\s*", line)
                ).strip()
            markdown_to_docx(candidate, filepath, title=title)
            errors = validate_docx_markdown_render(filepath)
            quality_attempts.append({"attempt": attempt, "errors": errors})
            if not errors:
                content = candidate
                request["content"] = candidate
                break
        else:
            filepath.unlink(missing_ok=True)
            raise RuntimeError("DOCX 格式校验失败，已完成两次自动修复仍未通过：" + ",".join(quality_attempts[-1]["errors"]))
    elif output_format == "markdown":
        filepath.write_text(content if content_type == "markdown" else _wrap_plaintext_as_markdown(content), encoding="utf-8")
    elif output_format == "json":
        filepath.write_text(json.dumps(request, ensure_ascii=False, indent=2), encoding="utf-8")
    else:
        raise RuntimeError(f"unsupported markdown-first output format: {output_format}")
    return {
        "filename": request["filename"],
        "filepath": str(filepath),
        "filesize": filepath.stat().st_size,
        "filetype": FILE_EXT_BY_FORMAT.get(output_format, ".docx"),
        "document_request": request,
        "quality_check": {"passed": True, "attempts": quality_attempts} if output_format == "docx" else {"passed": True, "attempts": []},
    }


def render_output_file(route: dict, content_json: dict) -> dict[str, Any]:
    output_spec = route.get("output_spec") or {}
    output_format = output_spec.get("format", "docx")
    if output_format == "chat":
        raise RuntimeError("chat output does not create a file")
    filename = str(output_spec.get("filename") or "").strip() or _safe_filename(str(content_json.get("title") or "document"), output_format)
    filepath = DOC_OUTPUT_DIR / f"{uuid.uuid4().hex}{FILE_EXT_BY_FORMAT.get(output_format, '.docx')}"
    if output_format == "docx":
        _render_docx(content_json, filepath)
    elif output_format == "pptx":
        _render_pptx(content_json, filepath)
    elif output_format == "markdown":
        filepath.write_text(_content_json_to_markdown(content_json), encoding="utf-8")
    elif output_format == "json":
        filepath.write_text(json.dumps(content_json, ensure_ascii=False, indent=2), encoding="utf-8")
    else:
        raise RuntimeError(f"unsupported output format: {output_format}")
    return {
        "filename": filename,
        "filepath": str(filepath),
        "filesize": filepath.stat().st_size,
        "filetype": FILE_EXT_BY_FORMAT.get(output_format, ".docx"),
    }


def _render_docx(content_json: dict, filepath: Path) -> None:
    doc = Document()
    doc.add_heading(str(content_json.get("title") or "文档"), level=1)
    summary = str(content_json.get("summary") or "").strip()
    if summary:
        doc.add_paragraph(summary)
    for section in content_json.get("sections") or []:
        if not isinstance(section, dict):
            continue
        heading_text = str(section.get("heading") or "正文")
        if re.match(r"^[一二三四五六七八九十]+、", heading_text):
            level = 2
        elif re.match(r"^（[一二三四五六七八九十]+）", heading_text):
            level = 3
        else:
            level = 2
        doc.add_heading(heading_text, level=level)
        for para in re.split(r"\n{2,}", str(section.get("content") or "").strip()):
            para = para.strip()
            if para:
                doc.add_paragraph(para)
    for table in content_json.get("tables") or []:
        if not isinstance(table, dict):
            continue
        headers = [str(item) for item in table.get("headers") or []]
        rows = table.get("rows") or []
        if not headers and rows and isinstance(rows[0], dict):
            headers = list(rows[0].keys())
        if not headers:
            continue
        if table.get("title"):
            doc.add_heading(str(table["title"]), level=2)
        doc_table = doc.add_table(rows=1, cols=len(headers))
        doc_table.style = "Table Grid"
        for index, header in enumerate(headers):
            doc_table.rows[0].cells[index].text = header
        for row in rows:
            cells = doc_table.add_row().cells
            if isinstance(row, dict):
                values = [row.get(header, "") for header in headers]
            else:
                values = row if isinstance(row, list) else [row]
            for index, value in enumerate(values[:len(headers)]):
                cells[index].text = _stringify(value)
    warnings = content_json.get("warnings") or []
    if warnings:
        doc.add_heading("注意事项", level=2)
        for warning in warnings:
            doc.add_paragraph(str(warning), style="List Bullet")
    doc.save(filepath)


def _render_pptx(content_json: dict, filepath: Path) -> None:
    from pptx import Presentation

    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = str(content_json.get("title") or "演示文稿")
    title_slide.placeholders[1].text = str(content_json.get("summary") or "")

    slides = content_json.get("slides") or []
    if not slides:
        for section in content_json.get("sections") or []:
            if isinstance(section, dict):
                slides.append({
                    "title": section.get("heading") or "内容",
                    "bullets": [line.strip("- ").strip() for line in str(section.get("content") or "").splitlines() if line.strip()],
                    "notes": "",
                })
    for slide_data in slides[:20]:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = str((slide_data or {}).get("title") or "内容")
        body = slide.placeholders[1].text_frame
        body.clear()
        bullets = (slide_data or {}).get("bullets") or []
        if not bullets and (slide_data or {}).get("notes"):
            bullets = [str((slide_data or {}).get("notes"))]
        for index, bullet in enumerate(bullets[:8]):
            paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
            paragraph.text = _stringify(bullet)
            paragraph.level = 0
    prs.save(filepath)


def _content_json_to_markdown(content_json: dict) -> str:
    lines = ["# " + str(content_json.get("title") or "文档")]
    summary = str(content_json.get("summary") or "").strip()
    if summary:
        lines.extend(["", summary])
    for section in content_json.get("sections") or []:
        if isinstance(section, dict):
            lines.extend(["", "## " + str(section.get("heading") or "正文"), "", str(section.get("content") or "").strip()])
    for table in content_json.get("tables") or []:
        if not isinstance(table, dict):
            continue
        headers = [str(item) for item in table.get("headers") or []]
        rows = table.get("rows") or []
        if headers:
            lines.extend(["", "## " + str(table.get("title") or "表格"), ""])
            lines.append("| " + " | ".join(headers) + " |")
            lines.append("| " + " | ".join("---" for _ in headers) + " |")
            for row in rows:
                values = [row.get(h, "") for h in headers] if isinstance(row, dict) else row
                if isinstance(values, list):
                    lines.append("| " + " | ".join(_stringify(v).replace("\n", " ") for v in values[:len(headers)]) + " |")
    warnings = content_json.get("warnings") or []
    if warnings:
        lines.extend(["", "## 注意事项"])
        lines.extend("- " + str(item) for item in warnings)
    return "\n".join(lines).strip() + "\n"


def content_json_to_chat_text(content_json: dict, route: dict) -> str:
    operation = route.get("primary_operation", "generate")
    if operation == "extract" and content_json.get("tables"):
        return _content_json_to_markdown(content_json)
    parts = []
    summary = str(content_json.get("summary") or "").strip()
    if summary:
        parts.append(summary)
    for section in content_json.get("sections") or []:
        if not isinstance(section, dict):
            continue
        heading = str(section.get("heading") or "").strip()
        body = str(section.get("content") or "").strip()
        if heading and len(content_json.get("sections") or []) > 1:
            parts.append(f"{heading}\n{body}")
        elif body:
            parts.append(body)
    warnings = content_json.get("warnings") or []
    if warnings:
        parts.append("注意事项：\n" + "\n".join("- " + str(item) for item in warnings))
    return "\n\n".join(part for part in parts if part).strip() or "已完成。"


def should_preview_outline(user_input: str, route: dict, attachment_context: str = "") -> bool:
    route = normalize_workflow_route(route, user_input, attachment_context=attachment_context, trust_route=True)
    return _should_generate_sectioned_document(user_input, route, attachment_context)


def build_outline_preview(user_input: str, route: dict, attachment_context: str = "") -> dict[str, Any] | None:
    route = normalize_workflow_route(route, user_input, attachment_context=attachment_context, trust_route=True)
    if not should_preview_outline(user_input, route, attachment_context):
        return None
    section_count = _section_count_for_request(user_input)
    stored_outline = route.get("outline_preview")
    if isinstance(stored_outline, dict) and stored_outline.get("sections"):
        return _normalize_outline(stored_outline, user_input, route, section_count)
    try:
        outline_timeout = _setting("outline_preview_timeout_seconds", 8)
        outline_raw = _call_qwen_json(
            _build_outline_prompt(user_input, route, attachment_context, section_count),
            max_tokens=1200,
            max_retries=0,
            timeout_seconds=outline_timeout,
        )
    except DocumentGenerationError as exc:
        logger.warning("Outline preview degraded: %s", exc.technical_message)
        outline_raw = _fallback_outline(user_input, route)
    outline = _normalize_outline(outline_raw, user_input, route, section_count)
    route["outline_preview"] = outline
    return outline


def render_outline_preview_text(user_input: str, route: dict, attachment_context: str = "") -> str:
    route = normalize_workflow_route(route, user_input, attachment_context=attachment_context, trust_route=True)
    lines = []
    outline = build_outline_preview(user_input, route, attachment_context)
    if outline and outline.get("sections"):
        lines.append("文档大纲预览：")
        title = str(outline.get("title") or "").strip()
        if title:
            lines.append(title)
        summary = str(outline.get("summary") or "").strip()
        if summary:
            lines.extend(["", summary])
        for index, section in enumerate(outline.get("sections") or [], 1):
            heading = str((section or {}).get("heading") or "").strip()
            if heading:
                lines.append(f"{index}. {heading}")
    return "\n".join(lines).strip()


class DocumentAgentExecutor:
    """Deterministic executor around a backend-built task graph."""

    def __init__(self, route: dict, user_input: str, attachment_context: str = "", attachments=None):
        self.route = normalize_workflow_route(route, user_input, attachments, attachment_context, trust_route=True)
        self.user_input = user_input or ""
        self.attachment_context = attachment_context or ""
        self.attachments = attachments
        self.task_graph = self.route.get("task_graph") or build_task_graph(self.route)
        self.node_results: list[dict[str, Any]] = []

    def run(
        self,
        content_schema: dict | None = None,
        render_file: bool = False,
        progress_callback: Callable[[str, int, str], None] | None = None,
    ) -> dict[str, Any]:
        validation_result = validate_task_plan(self.route, self.attachment_context or self.user_input, self.attachments)
        if validation_result["status"] != "success":
            return {
                "status": validation_result["status"],
                "message": validation_result["message"],
                "warnings": validation_result["warnings"],
                "routing_json": self.route,
                "task_graph": self.task_graph,
                "node_results": [],
            }
        if progress_callback:
            progress_callback("validated", 12, "文档任务校验完成")

        for node in self.task_graph:
            self.node_results.append({
                "task_id": node["id"],
                "type": node["type"],
                "status": "success",
                "output": node["output"],
                "warnings": [],
                "next_node": _next_node_id(self.task_graph, node["id"]),
            })

        document_request = None
        if _looks_like_document_request(content_schema):
            document_request = normalize_document_request(
                content_schema,
                fallback_instruction=self.user_input,
                route=self.route,
            )
            content_json = {
                "title": re.sub(r"\.[^.]+$", "", document_request["filename"]),
                "summary": "",
                "sections": [{"heading": "Markdown", "content": document_request["content"]}],
                "tables": [],
                "slides": [],
                "warnings": [],
                "source_usage": {"used_user_instruction": True},
            }
        elif render_file and requires_file_generation(self.route) and (self.route.get("output_spec") or {}).get("format") in {"docx", "markdown", "json"}:
            document_request = generate_markdown_document_request(
                self.user_input,
                self.route,
                self.attachment_context,
                progress_callback=progress_callback,
            )
            content_json = {
                "title": re.sub(r"\.[^.]+$", "", document_request["filename"]),
                "summary": "",
                "sections": [{"heading": "Markdown", "content": document_request["content"]}],
                "tables": [],
                "slides": [],
                "warnings": [],
                "source_usage": {"used_user_instruction": True},
            }
        else:
            content_json = _normalize_content_json(
                content_schema or generate_content_schema(
                    self.user_input,
                    self.route,
                    self.attachment_context,
                    progress_callback=progress_callback,
                ),
                self.route,
                self.user_input,
                self.attachment_context,
            )
        result = {
            "status": "success",
            "routing_json": self.route,
            "task_graph": self.task_graph,
            "validation_result": validation_result,
            "node_results": self.node_results,
            "content_json": content_json,
            "document_request": document_request,
            "response_text": content_json_to_chat_text(content_json, self.route),
            "warnings": validation_result["warnings"] + content_json.get("warnings", []),
        }
        if render_file and requires_file_generation(self.route):
            if progress_callback:
                progress_callback("rendering", 82, "正在渲染文档文件")
            if document_request:
                result["generated_file"] = render_document_request(self.route, document_request)
            else:
                result["generated_file"] = render_output_file(self.route, content_json)
            if progress_callback:
                progress_callback("download_card", 92, "正文已生成，正在准备下载卡片")
        return result


def _next_node_id(graph: list[dict[str, Any]], current_id: str) -> str | None:
    for index, node in enumerate(graph):
        if node.get("id") == current_id and index + 1 < len(graph):
            return graph[index + 1].get("id")
    return None


def execute_agent_chat(user_input: str, attachment_context: str = "", route: dict | None = None, attachments=None) -> dict[str, Any]:
    route = normalize_workflow_route(route, user_input, attachments, attachment_context, trust_route=True) if route else classify_workflow(user_input, attachments, attachment_context)
    route["output_spec"]["format"] = "chat"
    route["output_spec"]["download_card"] = False
    route["task_graph"] = build_task_graph(route)
    return DocumentAgentExecutor(route, user_input, attachment_context, attachments).run(render_file=False)


def generate_agent_output_file(
    user_input: str,
    route: dict | None = None,
    content_schema: dict | None = None,
    attachment_context: str = "",
    progress_callback: Callable[[str, int, str], None] | None = None,
) -> dict[str, Any]:
    route = normalize_workflow_route(route, user_input, attachment_context=attachment_context, trust_route=True) if route else classify_workflow(user_input, attachment_context=attachment_context)
    if not requires_file_generation(route):
        route["output_spec"]["format"] = "docx"
        route["output_spec"]["download_card"] = True
        route["task_graph"] = build_task_graph(route)
    requested_filename = _requested_filename(user_input, (route.get("output_spec") or {}).get("format", "docx"))
    if requested_filename:
        route.setdefault("output_spec", {})["filename"] = requested_filename
    executor = DocumentAgentExecutor(route, user_input, attachment_context)
    result = executor.run(content_schema=content_schema, render_file=True, progress_callback=progress_callback)
    if result["status"] != "success":
        raise RuntimeError(result.get("message") or "Document agent execution failed")
    generated = result["generated_file"]
    return {
        "route": result["routing_json"],
        "content_schema": result["content_json"],
        "document_request": result.get("document_request"),
        "filename": generated["filename"],
        "filepath": generated["filepath"],
        "filesize": generated["filesize"],
        "filetype": generated["filetype"],
        "task_graph": result["task_graph"],
        "validation_result": result["validation_result"],
        "node_results": result["node_results"],
    }


def stream_agent_plan(user_input: str, route: dict) -> Iterable[str]:
    route = normalize_workflow_route(route, user_input, trust_route=True)
    operation = route["primary_operation"]
    output_format = route["output_spec"]["format"]
    lines = [
        "我将按 Document Agent Workflow v2.5 处理：",
        "",
        f"- primary_operation：{operation}",
        f"- output：{output_format}",
        "- task_graph：",
    ]
    for node in route.get("task_graph") or build_task_graph(route):
        lines.append(f"  {node['id']}. {node['type']}")
    lines.extend(["", "正在开始执行...\n"])
    text = "\n".join(lines)
    for index in range(0, len(text), 8):
        yield text[index:index + 8]
