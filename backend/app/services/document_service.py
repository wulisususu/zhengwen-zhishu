"""Document generation service: route intent, generate structured content, write DOCX."""
from __future__ import annotations

import json
import re
import uuid
from pathlib import Path
from typing import Any, Iterable

import requests
from docx import Document
from docxtpl import DocxTemplate

from app.config import get_config

DOC_OUTPUT_DIR = Path("/home/youyeetoo/backend/data/generated_docs")
TEMPLATE_DIR = Path("/home/youyeetoo/backend/app/templates/docx")
DOC_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)

VALID_DOC_TYPES = {"weekly_report", "resume", "checklist", "simple_report", "complex_report"}
VALID_ENGINES = {"docxtpl", "python-docx"}
TEMPLATE_BY_TYPE = {
    "weekly_report": "weekly_v1",
    "resume": "resume_v1",
    "checklist": "checklist_v1",
    "simple_report": "simple_v1",
    "complex_report": "complex_v1",
}
ENGINE_BY_TYPE = {
    "weekly_report": "docxtpl",
    "resume": "docxtpl",
    "checklist": "docxtpl",
    "simple_report": "docxtpl",
    "complex_report": "python-docx",
}

DOC_ROUTER_SYSTEM_PROMPT = """You are a document type and generation engine router.
Your only task is to classify the user's document request and choose the best generation engine.
Do not generate document content.

Available doc_type:
- weekly_report
- resume
- checklist
- simple_report
- complex_report

Available engine:
- docxtpl: weekly_report, resume, checklist, fixed-template documents, stable table structures
- python-docx: complex structures, unknown fields, custom layout, mixed content, non-template reports

Output JSON only, with no extra text:
{"doc_type":"","engine":"","template_id":""}

template_id rules:
weekly_report -> weekly_v1
resume -> resume_v1
checklist -> checklist_v1
simple_report -> simple_v1
complex_report -> complex_v1
"""

CONTENT_SYSTEM_PROMPT = """You are a professional structured document content generator.
Generate clear JSON data that can be used to fill a document template.
Do not mention Word, docxtpl, python-docx, layout, or implementation details.
Do not output Markdown or explanations. Output JSON only.

Common structure:
{"title":"","summary":"","sections":{}}

For weekly_report:
{"title":"","this_week_work":[],"issues":[],"next_week_plan":[]}

For resume:
{"name":"","profile":"","skills":[],"experience":[],"education":[]}

For checklist:
{"title":"","items":[{"task":"","status":""}]}

For simple_report:
{"title":"","points":[]}

For complex_report:
{"title":"","sections":[{"heading":"","content":""}]}
"""

EXAMPLES = {
    "weekly_report": {"title": "", "this_week_work": [], "issues": [], "next_week_plan": []},
    "resume": {"name": "", "profile": "", "skills": [], "experience": [], "education": []},
    "checklist": {"title": "", "items": [{"task": "", "status": ""}]},
    "simple_report": {"title": "", "points": []},
    "complex_report": {"title": "", "sections": [{"heading": "", "content": ""}]},
}

ZH = {
    "user_input": "\u7528\u6237\u8f93\u5165\uff1a",
    "doc_type": "\u6587\u6863\u7c7b\u578b\uff1a",
    "request": "\u7528\u6237\u9700\u6c42\uff1a",
    "json_only": "\u8bf7\u4e25\u683c\u8f93\u51fa\u7b26\u5408\u8be5\u6587\u6863\u7c7b\u578b\u7684 JSON\u3002\u53c2\u8003\u7ed3\u6784\uff1a\n",
    "document": "\u6587\u6863",
    "section": "\u7ae0\u8282",
}


def _cfg() -> dict:
    return get_config().get("llm", {}) or {}


def _setting(name: str, default: Any) -> Any:
    return _cfg().get(name, default)


def _parse_json_object(text: str) -> dict:
    text = (text or "").strip()
    text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s*```$", "", text)
    try:
        value = json.loads(text)
        return value if isinstance(value, dict) else {}
    except Exception:
        pass
    match = re.search(r"\{[\s\S]*\}", text)
    if not match:
        return {}
    try:
        value = json.loads(match.group(0))
        return value if isinstance(value, dict) else {}
    except Exception:
        return {}


def _ollama_settings() -> tuple[str, str, float]:
    return (
        str(_setting("ollama_base_url", "http://192.168.10.212:21434")).rstrip("/"),
        str(_setting("ollama_model", "qwen3.5-35b-a3b-q4-4k:latest")),
        float(_setting("ollama_timeout_seconds", 240)),
    )


def _call_prompt_model(prompt: str, timeout: float = 12) -> str:
    url = str(_setting("prompt_model_url", "http://127.0.0.1:8001")).rstrip("/")
    resp = requests.post(url + "/v1/completions", json={"prompt": prompt, "stream": False}, timeout=timeout)
    resp.raise_for_status()
    return str(resp.json().get("text", ""))


def _call_qwen_json(prompt: str, max_tokens: int = 2048) -> dict:
    url, model, timeout = _ollama_settings()
    payload = {
        "model": model,
        "prompt": prompt,
        "stream": False,
        "format": "json",
        "options": {"num_predict": max_tokens, "temperature": 0.2, "top_p": 0.8, "repeat_penalty": 1.05},
    }
    resp = requests.post(url + "/api/generate", json=payload, timeout=timeout)
    resp.raise_for_status()
    return _parse_json_object(str(resp.json().get("response", "")))


def _valid_plan_text(text: str) -> bool:
    if not text:
        return False
    lowered = text.lower()
    bad_markers = ("<think", "thinking process", "analyze the request", "drafting the content", "determine the content")
    if any(marker in lowered for marker in bad_markers):
        return False
    chinese_count = len(re.findall(r"[\u4e00-\u9fff]", text))
    ascii_letters = len(re.findall(r"[A-Za-z]", text))
    return chinese_count >= 12 and chinese_count >= ascii_letters * 0.6


def _clean_plan_text(text: str) -> str:
    text = text or ""
    text = re.sub(r"<think>[\s\S]*?</think>", "", text, flags=re.IGNORECASE)
    text = re.sub(r"</?think>", "", text, flags=re.IGNORECASE)
    text = re.sub(r"Thinking Process:[\s\S]*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"Analyze the Request[\s\S]*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"Drafting the Content[\s\S]*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"Determine the Content[\s\S]*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\*\*[^\n]*\*\*", "", text)
    lines = []
    for raw in text.replace("\r\n", "\n").splitlines():
        line = raw.strip().strip("` ")
        if not line:
            if lines and lines[-1] != "":
                lines.append("")
            continue
        if re.search(r"[\u4e00-\u9fff]", line):
            lines.append(line)
    return "\n".join(lines).strip()


def _yield_text_chunks(text: str, chunk_size: int = 10) -> Iterable[str]:
    for i in range(0, len(text), chunk_size):
        yield text[i:i + chunk_size]


def stream_document_plan(user_input: str, route: dict) -> Iterable[str]:
    # The visible planning phase must be stable and Chinese-only; Qwen content
    # generation happens in the async document task immediately after this.
    plan = fallback_plan_text(route).rstrip()
    yield from _yield_text_chunks(plan, chunk_size=6)


def fallback_plan_text(route: dict) -> str:
    modules = {
        "weekly_report": ["\u57fa\u672c\u4fe1\u606f\u6a21\u5757", "\u672c\u5468\u5de5\u4f5c\u603b\u7ed3", "\u5de5\u4f5c\u4efb\u52a1\u62c6\u89e3", "\u98ce\u9669\u4e0e\u95ee\u9898", "\u4e0b\u4e00\u6b65\u8ba1\u5212"],
        "resume": ["\u57fa\u672c\u4fe1\u606f", "\u4e2a\u4eba\u7b80\u4ecb", "\u6280\u80fd\u6e05\u5355", "\u9879\u76ee\u4e0e\u5de5\u4f5c\u7ecf\u5386", "\u6559\u80b2\u80cc\u666f"],
        "checklist": ["\u6e05\u5355\u6807\u9898", "\u4efb\u52a1\u6761\u76ee", "\u72b6\u6001\u6807\u8bb0"],
        "simple_report": ["\u62a5\u544a\u6807\u9898", "\u6838\u5fc3\u8981\u70b9", "\u603b\u7ed3"],
        "complex_report": ["\u80cc\u666f", "\u5206\u6790", "\u65b9\u6848", "\u98ce\u9669", "\u8ba1\u5212"],
    }.get(route.get("doc_type"), ["\u6587\u6863\u7ed3\u6784", "\u5185\u5bb9\u6574\u7406"])
    lines = ["\u6211\u5c06\u6309\u7167\u4ee5\u4e0b\u7ed3\u6784\u751f\u6210Word\u6587\u6863\uff1a", ""]
    lines.extend(f"{i}. {name}" for i, name in enumerate(modules, 1))
    lines.append("")
    lines.append("\u6b63\u5728\u5f00\u59cb\u751f\u6210...\n")
    return "\n".join(lines)


def is_document_request(content: str) -> bool:
    text = (content or "").lower()
    doc_words = ("word", "docx", "\u6587\u6863", "\u5468\u62a5", "\u7b80\u5386", "\u6e05\u5355", "\u68c0\u67e5\u8868", "\u62a5\u544a")
    action_words = ("\u751f\u6210", "\u521b\u5efa", "\u65b0\u5efa", "\u5199", "\u505a", "\u5236\u4f5c", "\u5bfc\u51fa", "\u6574\u7406")
    return any(word in text for word in doc_words) and any(word in text for word in action_words)


def _fallback_route(user_input: str) -> dict:
    text = (user_input or "").lower()
    if "\u5468\u62a5" in text:
        doc_type = "weekly_report"
    elif "\u7b80\u5386" in text or "resume" in text:
        doc_type = "resume"
    elif "\u6e05\u5355" in text or "checklist" in text or "\u68c0\u67e5\u8868" in text:
        doc_type = "checklist"
    elif "\u590d\u6742" in text or "\u65b9\u6848" in text or "\u8c03\u7814" in text or "\u5206\u6790" in text or "\u6df7\u5408" in text:
        doc_type = "complex_report"
    elif "\u62a5\u544a" in text or "word" in text or "docx" in text or "\u6587\u6863" in text:
        doc_type = "simple_report"
    else:
        doc_type = "complex_report"
    return {"doc_type": doc_type, "engine": ENGINE_BY_TYPE[doc_type], "template_id": TEMPLATE_BY_TYPE[doc_type]}


def normalize_route(route: dict | None, user_input: str = "", trust_route: bool = False) -> dict:
    route = dict(route or {})
    fallback = _fallback_route(user_input)
    if not trust_route and fallback["doc_type"] in {"weekly_report", "resume", "checklist", "complex_report"}:
        route = fallback
    doc_type = route.get("doc_type")
    if doc_type not in VALID_DOC_TYPES:
        route = fallback
        doc_type = route["doc_type"]
    engine = route.get("engine") if route.get("engine") in VALID_ENGINES else ENGINE_BY_TYPE[doc_type]
    template_id = route.get("template_id") or TEMPLATE_BY_TYPE[doc_type]
    expected_template = TEMPLATE_BY_TYPE[doc_type]
    if template_id != expected_template:
        template_id = expected_template
    return {"doc_type": doc_type, "engine": engine, "template_id": template_id}


def classify_document(user_input: str) -> dict:
    prompt = DOC_ROUTER_SYSTEM_PROMPT + "\n" + ZH["user_input"] + (user_input or "")
    try:
        route = _parse_json_object(_call_prompt_model(prompt, timeout=float(_setting("prompt_timeout_seconds", 12))))
    except Exception:
        route = {}
    return normalize_route(route, user_input, trust_route=False)


def generate_content_schema(user_input: str, doc_type: str) -> dict:
    example = json.dumps(EXAMPLES.get(doc_type, EXAMPLES["complex_report"]), ensure_ascii=False, indent=2)
    prompt = CONTENT_SYSTEM_PROMPT + "\n" + ZH["doc_type"] + doc_type + "\n" + ZH["request"] + (user_input or "") + "\n" + ZH["json_only"] + example
    schema = _call_qwen_json(prompt)
    if not schema:
        raise RuntimeError("Qwen did not return valid JSON content_schema")
    return schema


def ensure_default_templates() -> None:
    templates = {
        "weekly_v1": ["{{ title }}", "\u672c\u5468\u5de5\u4f5c\u603b\u7ed3", "\u5de5\u4f5c\u5185\u5bb9\uff1a", "{{ this_week_work }}", "\u95ee\u9898\uff1a", "{{ issues }}", "\u4e0b\u5468\u8ba1\u5212\uff1a", "{{ next_week_plan }}"],
        "resume_v1": ["{{ name }}", "\u4e2a\u4eba\u7b80\u4ecb", "{{ profile }}", "\u6280\u80fd", "{{ skills }}", "\u5de5\u4f5c\u7ecf\u5386", "{{ experience }}", "\u6559\u80b2\u7ecf\u5386", "{{ education }}"],
        "checklist_v1": ["{{ title }}", "\u6e05\u5355\u5185\u5bb9", "{{ items }}"],
        "simple_v1": ["{{ title }}", "\u62a5\u544a\u8981\u70b9", "{{ points }}"],
    }
    heading_words = {"\u4e2a\u4eba\u7b80\u4ecb", "\u6280\u80fd", "\u5de5\u4f5c\u7ecf\u5386", "\u6559\u80b2\u7ecf\u5386", "\u6e05\u5355\u5185\u5bb9", "\u62a5\u544a\u8981\u70b9", "\u672c\u5468\u5de5\u4f5c\u603b\u7ed3"}
    for template_id, paragraphs in templates.items():
        path = TEMPLATE_DIR / f"{template_id}.docx"
        if path.exists():
            continue
        doc = Document()
        for index, paragraph in enumerate(paragraphs):
            if index == 0:
                doc.add_heading(paragraph, level=1)
            elif paragraph.endswith("\uff1a") or paragraph in heading_words:
                doc.add_heading(paragraph, level=2)
            else:
                doc.add_paragraph(paragraph)
        doc.save(path)


def _format_item(value: Any) -> str:
    if isinstance(value, dict):
        parts = []
        for key, val in value.items():
            if isinstance(val, (list, dict)):
                val = _list_text(val)
            parts.append(f"{key}: {val}")
        return "\uff1b".join(parts)
    return str(value)


def _list_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, list):
        return "\n".join(f"- {_format_item(item)}" for item in value)
    if isinstance(value, dict):
        return "\n".join(f"{key}: {_format_item(val)}" for key, val in value.items())
    return str(value)


def _template_context(content_schema: dict) -> dict:
    context = dict(content_schema or {})
    for key, value in list(context.items()):
        if isinstance(value, (list, dict)):
            context[key] = _list_text(value)
    for key in ("this_week_work", "issues", "next_week_plan", "skills", "experience", "education", "items", "points"):
        context[key] = _list_text((content_schema or {}).get(key, context.get(key, "")))
    context.setdefault("title", context.get("name", ZH["document"]))
    return context


def render_with_template(template_id: str, content_schema: dict, output_path: Path) -> None:
    ensure_default_templates()
    template_path = TEMPLATE_DIR / f"{template_id}.docx"
    if not template_path.exists():
        raise RuntimeError(f"template not found: {template_id}")
    tpl = DocxTemplate(str(template_path))
    tpl.render(_template_context(content_schema))
    tpl.save(output_path)


def _add_dynamic_value(doc: Document, key: str, value: Any, level: int = 1) -> None:
    title = str(key).replace("_", " ")
    if isinstance(value, list):
        doc.add_heading(title, level=min(level, 3))
        for item in value:
            if isinstance(item, dict):
                for sub_key, sub_val in item.items():
                    _add_dynamic_value(doc, sub_key, sub_val, level + 1)
            else:
                doc.add_paragraph(str(item), style="List Bullet")
    elif isinstance(value, dict):
        doc.add_heading(title, level=min(level, 3))
        for sub_key, sub_val in value.items():
            _add_dynamic_value(doc, sub_key, sub_val, level + 1)
    elif key not in {"title"}:
        doc.add_heading(title, level=min(level, 3))
        doc.add_paragraph(str(value or ""))


def build_dynamic_doc(content_schema: dict, output_path: Path) -> None:
    doc = Document()
    title = str((content_schema or {}).get("title") or ZH["document"])
    doc.add_heading(title, level=1)
    sections = (content_schema or {}).get("sections")
    if isinstance(sections, list):
        for section in sections:
            if isinstance(section, dict):
                heading = section.get("heading") or section.get("title") or ZH["section"]
                doc.add_heading(str(heading), level=2)
                content = section.get("content", "")
                if isinstance(content, list):
                    for item in content:
                        doc.add_paragraph(_format_item(item), style="List Bullet")
                else:
                    doc.add_paragraph(str(content or ""))
            else:
                doc.add_paragraph(str(section), style="List Bullet")
    else:
        for key, value in (content_schema or {}).items():
            _add_dynamic_value(doc, key, value, level=2)
    doc.save(output_path)


def _safe_filename(title: str, doc_type: str) -> str:
    base = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff_-]+", "_", (title or doc_type).strip()).strip("_")
    if not base:
        base = doc_type
    return base[:48] + ".docx"


def generate_document_file(user_input: str, route: dict | None = None, content_schema: dict | None = None) -> dict:
    route = normalize_route(route, user_input, trust_route=True) if route else classify_document(user_input)
    if content_schema is None:
        content_schema = generate_content_schema(user_input, route["doc_type"])
    if not isinstance(content_schema, dict):
        raise RuntimeError("content_schema must be an object")
    title = str(content_schema.get("title") or content_schema.get("name") or route["doc_type"])
    filename = _safe_filename(title, route["doc_type"])
    output_path = DOC_OUTPUT_DIR / f"{uuid.uuid4().hex}.docx"
    if route["engine"] == "docxtpl":
        render_with_template(route["template_id"], content_schema, output_path)
    elif route["engine"] == "python-docx":
        build_dynamic_doc(content_schema, output_path)
    else:
        raise RuntimeError(f"unsupported engine: {route['engine']}")
    return {"route": route, "content_schema": content_schema, "filename": filename, "filepath": str(output_path), "filesize": output_path.stat().st_size}


# Document Agent Workflow v2.5 overrides. The legacy helpers above stay in the
# file for compatibility, while the exported functions below implement the
# final routing_json + validation + task_graph + executor flow.
from app.services.document_agent import (  # noqa: E402
    build_outline_preview,
    classify_workflow as _agent_classify_workflow,
    execute_agent_chat,
    generate_agent_output_file,
    generate_content_schema as _agent_generate_content_schema,
    is_document_agent_request,
    normalize_workflow_route,
    requires_file_generation,
    render_outline_preview_text,
    should_preview_outline,
    stream_agent_plan,
    validate_task_plan,
)


def normalize_route(
    route: dict | None,
    user_input: str = "",
    trust_route: bool = False,
    attachments=None,
    attachment_context: str = "",
    force_output_format: str | None = None,
) -> dict:
    return normalize_workflow_route(
        route,
        user_input=user_input,
        attachments=attachments,
        attachment_context=attachment_context,
        trust_route=trust_route,
        force_output_format=force_output_format,
    )


def classify_document(
    user_input: str,
    attachments=None,
    attachment_context: str = "",
    force_output_format: str | None = None,
) -> dict:
    return _agent_classify_workflow(
        user_input,
        attachments=attachments,
        attachment_context=attachment_context,
        force_output_format=force_output_format,
    )


def is_document_request(content: str, has_attachments: bool = False) -> bool:
    return is_document_agent_request(content, has_attachments=has_attachments)


def stream_document_plan(user_input: str, route: dict) -> Iterable[str]:
    yield from stream_agent_plan(user_input, route)


def generate_content_schema(user_input: str, doc_type: str | dict, attachment_context: str = "", progress_callback=None) -> dict:
    return _agent_generate_content_schema(
        user_input,
        doc_type,
        attachment_context=attachment_context,
        progress_callback=progress_callback,
    )


def generate_document_file(
    user_input: str,
    route: dict | None = None,
    content_schema: dict | None = None,
    attachment_context: str = "",
    progress_callback=None,
) -> dict:
    return generate_agent_output_file(
        user_input,
        route=route,
        content_schema=content_schema,
        attachment_context=attachment_context,
        progress_callback=progress_callback,
    )


def build_document_outline_preview(user_input: str, route: dict, attachment_context: str = "") -> dict | None:
    return build_outline_preview(user_input, route, attachment_context=attachment_context)


def render_document_preview_text(user_input: str, route: dict, attachment_context: str = "") -> str:
    return render_outline_preview_text(user_input, route, attachment_context=attachment_context)


def should_render_document_outline(user_input: str, route: dict, attachment_context: str = "") -> bool:
    return should_preview_outline(user_input, route, attachment_context=attachment_context)


def validate_document_route(route: dict, attachment_context: str = "", attachments=None) -> dict:
    return validate_task_plan(route, attachment_context=attachment_context, attachments=attachments)
