"""Async file parsing and attachment context helpers."""
from __future__ import annotations

import hashlib
import json
import logging
import os
import re
import subprocess
import threading
import uuid
import urllib.request
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import Any

from app.database import SessionLocal
from app.models import File, FileParseTask, ParsedFile, ParsedFileChunk

logger = logging.getLogger(__name__)

PARSE_OUTPUT_DIR = Path("/home/youyeetoo/InNet_box/backend/data/parsed_files")
PIPELINE_RUNNER = Path("/home/youyeetoo/rkllm_model_zoo/mqw-tools/pipeline/run.py")
MINICPM_PYTHON = Path("/home/youyeetoo/rkllm_model_zoo/MiniCPM-V-2_6-rkllm/.venv/bin/python3")
MINICPM_LIB_DIR = Path("/home/youyeetoo/rkllm_model_zoo/MiniCPM-V-2_6-rkllm/lib")
MINICPM_OCR_RUNNER = Path("/home/youyeetoo/InNet_box/backend/app/services/minicpm_ocr_runner.py")
LLM_06B_SERVICE_URL = os.environ.get("LLM_06B_SERVICE_URL", "http://127.0.0.1:8003").rstrip("/")
MINICPM_ENABLED = os.environ.get("MINICPM_ENABLED", "true").strip().lower() not in {"0", "false", "no", "off"}

PDF_EXTS = {".pdf"}
STANDARD_IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".bmp", ".webp", ".tif", ".tiff"}
IMAGE_EXTS = STANDARD_IMAGE_EXTS | {".gif", ".jp2", ".j2k", ".ppm", ".pgm", ".pbm", ".pnm", ".heic", ".heif"}
IMAGE_PARSE_VERSION = "minicpm_v2_6_visual_rapidocr_fallback_v3"
PDF_PARSE_VERSION = "pdf_text_image_required_v7"
PDF_VISUAL_CACHE_VERSION = "pdf_image_required_visual_v7"
WORD_EXTS = {".docx", ".doc"}
OFFICE_ZIP_EXTS = {".pptx", ".xlsx"}
TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".json", ".jsonl", ".csv", ".tsv",
    ".yaml", ".yml", ".xml", ".html", ".htm", ".log", ".ini",
    ".py", ".js", ".ts", ".vue", ".css", ".sql",
}
SUPPORTED_PARSE_EXTS = PDF_EXTS | IMAGE_EXTS | WORD_EXTS | OFFICE_ZIP_EXTS | TEXT_EXTS
MAX_CONTEXT_CHARS = 22000
CHUNK_SIZE = 4500
CHUNK_OVERLAP = 350
PDF_MAX_VISUAL_IMAGES_PER_PAGE = 3
PDF_MIN_VISUAL_IMAGE_PIXELS = 80 * 80

PARSE_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
VISUAL_CACHE_DIR = PARSE_OUTPUT_DIR / "visual_cache"
VISUAL_CACHE_DIR.mkdir(parents=True, exist_ok=True)

# MiniCPM/OCR must not run concurrently on RK3588. Keep all parse jobs serialized.
_executor = ThreadPoolExecutor(max_workers=1)
_vector_executor = ThreadPoolExecutor(max_workers=1)
_minicpm_lock = threading.Lock()

VISUAL_QUESTION_RE = re.compile(
    r"(图片|图像|画面|插图|配图|图文|图表|示意图|照片|截图|视觉|版面|布局|"
    r"image|picture|photo|figure|chart|visual|layout)",
    re.IGNORECASE,
)
RELATION_QUESTION_RE = re.compile(r"(关系|关联|对应|说明|表达|含义|讲什么|为什么|relation|related)", re.IGNORECASE)


def _json_dump(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False)


def _json_load(value: str | None, default):
    if not value:
        return default
    try:
        return json.loads(value)
    except Exception:
        return default


def _sha256(path: str) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def is_parse_supported(filename: str | None) -> bool:
    ext = Path(filename or "").suffix.lower()
    return ext in SUPPORTED_PARSE_EXTS


def should_auto_parse_on_upload(filename: str | None) -> bool:
    ext = Path(filename or "").suffix.lower()
    return ext in SUPPORTED_PARSE_EXTS and ext not in IMAGE_EXTS


def needs_visual_attachment_context(question: str) -> bool:
    text = question or ""
    return bool(VISUAL_QUESTION_RE.search(text) or (RELATION_QUESTION_RE.search(text) and "文字" in text))


def should_run_visual_analysis(question: str, text_context: str) -> bool:
    if not needs_visual_attachment_context(question):
        return False
    q = question or ""
    text_len = len(re.sub(r"\s+", "", text_context or ""))
    asks_relationship = bool(re.search(r"(图文|图片.*文字|文字.*图片|关系|关联|对应|说明)", q))
    asks_image_content = bool(re.search(r"(图里|图片里|画面里|描述图片|描述画面|图像内容|图中|视觉元素)", q))
    if asks_relationship and text_len >= 220:
        return False
    if asks_image_content:
        return True
    return text_len < 220


def _cache_compatible(parsed: ParsedFile | None, filename: str | None) -> bool:
    if not parsed:
        return False
    ext = Path(filename or "").suffix.lower()
    if ext in IMAGE_EXTS:
        metadata = _json_load(parsed.metadata_json, {})
        return metadata.get("parse_version") == IMAGE_PARSE_VERSION
    if ext == ".pdf":
        metadata = _json_load(parsed.metadata_json, {})
        if metadata.get("parse_version") != PDF_PARSE_VERSION:
            return False
        markdown = str(parsed.markdown or "")
        if "MiniCPM 图片分析失败：" in markdown or "视觉分析失败：" in markdown:
            return False
        image_pages = {int(page) for page in metadata.get("image_pages") or [] if str(page).isdigit()}
        if not image_pages:
            return True
        attempted_pages = {
            int(page) for page in metadata.get("minicpm_attempted_pages") or metadata.get("visual_pages") or []
            if str(page).isdigit()
        }
        return bool(metadata.get("minicpm_required")) and image_pages.issubset(attempted_pages)
    return True


def _public_task(task: FileParseTask, db=None) -> dict[str, Any]:
    file_info = None
    parsed_info = None
    if db is not None:
        file = db.query(File).filter(File.id == task.file_id, File.user_id == task.user_id).first()
        if file:
            file_info = {
                "id": file.id,
                "filename": file.filename,
                "size": file.filesize,
                "type": file.filetype,
            }
        if task.parsed_file_id:
            parsed = db.query(ParsedFile).filter(
                ParsedFile.id == task.parsed_file_id,
                ParsedFile.user_id == task.user_id,
            ).first()
            if parsed:
                parsed_info = _public_parsed(parsed, include_markdown=False)
    return {
        "task_id": task.task_id,
        "file_id": task.file_id,
        "status": task.status,
        "progress": task.progress,
        "stage": task.stage,
        "error": task.error,
        "file": file_info,
        "parsed_file": parsed_info,
        "created_at": task.created_at.isoformat() if task.created_at else None,
        "updated_at": task.updated_at.isoformat() if task.updated_at else None,
    }


def _public_parsed(parsed: ParsedFile, include_markdown: bool = False) -> dict[str, Any]:
    data = {
        "id": parsed.id,
        "file_id": parsed.file_id,
        "sha256": parsed.sha256,
        "mime_type": parsed.mime_type,
        "title": parsed.title,
        "page_count": parsed.page_count,
        "ocr_pages": _json_load(parsed.ocr_pages_json, []),
        "metadata": _json_load(parsed.metadata_json, {}),
        "created_at": parsed.created_at.isoformat() if parsed.created_at else None,
    }
    if include_markdown:
        data["markdown"] = parsed.markdown or ""
    return data


def get_parse_task(task_id: str, user_id: int | None = None, file_id: int | None = None) -> dict[str, Any] | None:
    db = SessionLocal()
    try:
        query = db.query(FileParseTask).filter(FileParseTask.task_id == task_id)
        if user_id is not None:
            query = query.filter(FileParseTask.user_id == user_id)
        if file_id is not None:
            query = query.filter(FileParseTask.file_id == file_id)
        task = query.first()
        return _public_task(task, db) if task else None
    finally:
        db.close()


def get_parsed_file(file_id: int, user_id: int, include_markdown: bool = False) -> dict[str, Any] | None:
    db = SessionLocal()
    try:
        parsed = db.query(ParsedFile).filter(
            ParsedFile.file_id == file_id,
            ParsedFile.user_id == user_id,
        ).order_by(ParsedFile.id.desc()).first()
        return _public_parsed(parsed, include_markdown=include_markdown) if parsed else None
    finally:
        db.close()


def create_file_parse_task(file_id: int, user_id: int, force: bool = False) -> dict[str, Any] | None:
    db = SessionLocal()
    try:
        file = db.query(File).filter(File.id == file_id, File.user_id == user_id).first()
        if not file:
            return None
        if not is_parse_supported(file.filename):
            return _create_terminal_task(db, file, user_id, "error", "unsupported", "该文件类型暂不支持解析")

        if not force:
            running = db.query(FileParseTask).filter(
                FileParseTask.file_id == file_id,
                FileParseTask.user_id == user_id,
                FileParseTask.status.in_(("queued", "extracting", "ocr", "chunking")),
            ).order_by(FileParseTask.created_at.desc()).first()
            if running:
                return _public_task(running, db)

            parsed = db.query(ParsedFile).filter(
                ParsedFile.file_id == file_id,
                ParsedFile.user_id == user_id,
            ).order_by(ParsedFile.id.desc()).first()
            if parsed and _cache_compatible(parsed, file.filename):
                done_task = db.query(FileParseTask).filter(
                    FileParseTask.file_id == file_id,
                    FileParseTask.user_id == user_id,
                    FileParseTask.status == "done",
                ).order_by(FileParseTask.created_at.desc()).first()
                if done_task:
                    return _public_task(done_task, db)
                task = FileParseTask(
                    task_id=uuid.uuid4().hex,
                    user_id=user_id,
                    file_id=file_id,
                    parsed_file_id=parsed.id,
                    status="done",
                    progress=100,
                    stage="cached",
                )
                db.add(task)
                db.commit()
                db.refresh(task)
                return _public_task(task, db)

        task = FileParseTask(
            task_id=uuid.uuid4().hex,
            user_id=user_id,
            file_id=file_id,
            status="queued",
            progress=3,
            stage="queued",
        )
        db.add(task)
        db.commit()
        db.refresh(task)
        public = _public_task(task, db)
    finally:
        db.close()

    _executor.submit(_run_parse_task, public["task_id"])
    return public


def ensure_parse_tasks(file_ids: list[int], user_id: int) -> list[dict[str, Any]]:
    tasks = []
    for file_id in file_ids:
        task = create_file_parse_task(file_id, user_id, force=False)
        if task:
            tasks.append(task)
    return tasks


def _create_terminal_task(db, file: File, user_id: int, status: str, stage: str, error: str) -> dict[str, Any]:
    task = FileParseTask(
        task_id=uuid.uuid4().hex,
        user_id=user_id,
        file_id=file.id,
        status=status,
        progress=100,
        stage=stage,
        error=error,
    )
    db.add(task)
    db.commit()
    db.refresh(task)
    return _public_task(task, db)


def _update(task_id: str, **kwargs) -> None:
    db = SessionLocal()
    try:
        task = db.query(FileParseTask).filter(FileParseTask.task_id == task_id).first()
        if not task:
            return
        for key, value in kwargs.items():
            setattr(task, key, value)
        db.add(task)
        db.commit()
    finally:
        db.close()


def _run_parse_task(task_id: str) -> None:
    db = SessionLocal()
    try:
        task = db.query(FileParseTask).filter(FileParseTask.task_id == task_id).first()
        if not task:
            return
        file = db.query(File).filter(File.id == task.file_id, File.user_id == task.user_id).first()
        if not file or not os.path.exists(file.filepath):
            raise RuntimeError("文件不存在")

        task.status = "extracting"
        task.progress = 10
        task.stage = "正在计算文件指纹"
        db.add(task)
        db.commit()

        digest = _sha256(file.filepath)
        cached = db.query(ParsedFile).filter(
            ParsedFile.user_id == task.user_id,
            ParsedFile.sha256 == digest,
        ).order_by(ParsedFile.id.desc()).first()
        if cached and not _cache_compatible(cached, file.filename):
            cached = None
        if cached and cached.file_id != file.id:
            parsed = _clone_parsed_file(db, cached, file, task.user_id)
            task.status = "done"
            task.progress = 100
            task.stage = "已复用解析缓存"
            task.parsed_file_id = parsed.id
            db.add(task)
            db.commit()
            return

        task.status = _extract_status_for_file(file.filename)
        task.progress = 35
        task.stage = _extract_stage_for_file(file.filename)
        db.add(task)
        db.commit()

        markdown, metadata, ocr_pages = _extract_file(file, task_id=task_id)
        markdown = clean_markdown(markdown)
        if not markdown.strip():
            raise RuntimeError("未提取到有效文本")

        task.status = "chunking"
        task.progress = 82
        task.stage = "正在切分文档内容"
        db.add(task)
        db.commit()

        parsed = ParsedFile(
            user_id=task.user_id,
            file_id=file.id,
            sha256=digest,
            mime_type=file.filetype or Path(file.filename).suffix.lower(),
            title=file.filename,
            markdown=markdown,
            metadata_json=_json_dump(metadata),
            page_count=int(metadata.get("page_count") or 0),
            ocr_pages_json=_json_dump(ocr_pages),
        )
        db.add(parsed)
        db.commit()
        db.refresh(parsed)

        chunks = make_chunks(markdown)
        for index, chunk in enumerate(chunks):
            rec = ParsedFileChunk(
                parsed_file_id=parsed.id,
                user_id=task.user_id,
                chunk_index=index,
                page_start=int(chunk.get("page_start") or 0),
                page_end=int(chunk.get("page_end") or 0),
                content=chunk["content"],
                summary=chunk["summary"],
                token_estimate=max(1, len(chunk["content"]) // 2),
            )
            db.add(rec)
        db.commit()

        task.status = "done"
        task.progress = 100
        task.stage = "解析完成"
        task.error = None
        task.parsed_file_id = parsed.id
        db.add(task)
        db.commit()
        _schedule_vectorization(parsed.id)
    except Exception as exc:
        db.rollback()
        _update(task_id, status="error", progress=100, stage="解析失败", error=str(exc))
    finally:
        db.close()


def _schedule_vectorization(parsed_file_id: int) -> None:
    try:
        _vector_executor.submit(_vectorize_parsed_file, parsed_file_id)
    except Exception as exc:
        logger.warning("Unable to schedule KB vectorization for parsed_file %d: %s", parsed_file_id, exc)


def _vectorize_parsed_file(parsed_file_id: int) -> None:
    try:
        from app.kb.kb_ingest_service import vectorize_parsed_file

        result = vectorize_parsed_file(parsed_file_id)
        logger.info("KB vectorize result for parsed_file %d: %s", parsed_file_id, result)
    except Exception as exc:
        logger.warning("KB vectorize failed for parsed_file %d: %s", parsed_file_id, exc, exc_info=True)


def _clone_parsed_file(db, source: ParsedFile, file: File, user_id: int) -> ParsedFile:
    parsed = ParsedFile(
        user_id=user_id,
        file_id=file.id,
        sha256=source.sha256,
        mime_type=source.mime_type,
        title=file.filename,
        markdown=source.markdown,
        metadata_json=source.metadata_json,
        page_count=source.page_count,
        ocr_pages_json=source.ocr_pages_json,
    )
    db.add(parsed)
    db.commit()
    db.refresh(parsed)
    source_chunks = db.query(ParsedFileChunk).filter(
        ParsedFileChunk.parsed_file_id == source.id,
        ParsedFileChunk.user_id == user_id,
    ).order_by(ParsedFileChunk.chunk_index).all()
    for chunk in source_chunks:
        db.add(ParsedFileChunk(
            parsed_file_id=parsed.id,
            user_id=user_id,
            chunk_index=chunk.chunk_index,
            page_start=chunk.page_start,
            page_end=chunk.page_end,
            content=chunk.content,
            summary=chunk.summary,
            token_estimate=chunk.token_estimate,
        ))
    db.commit()
    return parsed


def _extract_status_for_file(filename: str | None) -> str:
    ext = Path(filename or "").suffix.lower()
    if ext in IMAGE_EXTS:
        return "ocr"
    return "extracting"


def _extract_stage_for_file(filename: str | None) -> str:
    ext = Path(filename or "").suffix.lower()
    if ext in IMAGE_EXTS:
        return "正在调用 MiniCPM 视觉解析图片"
    if ext == ".pdf":
        return "正在解析 PDF 文本层；如发现图片将调用 MiniCPM"
    if ext in WORD_EXTS:
        return "正在解析 Word 文档"
    if ext in OFFICE_ZIP_EXTS:
        return "正在解析 Office 文档"
    if ext in TEXT_EXTS:
        return "正在读取文本内容"
    return "正在解析文件内容"


def _extract_file(file: File, task_id: str | None = None) -> tuple[str, dict[str, Any], list[int]]:
    ext = Path(file.filename or file.filepath).suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(file.filepath, task_id=task_id)
    if ext in IMAGE_EXTS:
        return _extract_image(file.filepath)
    if ext in TEXT_EXTS:
        text = _extract_text_file(file.filepath, ext)
        return text, {"source": "plain_text", "page_count": 1}, []
    if ext == ".docx":
        from docx import Document

        doc = Document(file.filepath)
        parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n\n".join(parts), {"source": "python-docx", "page_count": 1}, []
    if ext == ".doc":
        return _extract_legacy_doc(file.filepath)
    if ext in OFFICE_ZIP_EXTS:
        return _extract_office_zip(file.filepath, ext)

    text = _decode_text_best_effort(file.filepath)
    if text.strip():
        return text, {"source": "best_effort_text", "page_count": 1}, []
    raise RuntimeError("该文件类型暂不支持解析")


def _extract_text_file(path: str, ext: str) -> str:
    text = _decode_text_best_effort(path)
    if ext == ".json":
        try:
            value = json.loads(text)
            return json.dumps(value, ensure_ascii=False, indent=2)
        except Exception:
            return text
    if ext == ".jsonl":
        lines = []
        for raw in text.splitlines():
            raw = raw.strip()
            if not raw:
                continue
            try:
                lines.append(json.dumps(json.loads(raw), ensure_ascii=False))
            except Exception:
                lines.append(raw)
        return "\n".join(lines)
    return text


def _decode_text_best_effort(path: str) -> str:
    data = Path(path).read_bytes()
    for enc in ("utf-8", "utf-8-sig", "gb18030", "gbk", "latin-1"):
        try:
            text = data.decode(enc)
            if text.strip():
                return text
        except Exception:
            continue
    return data.decode("utf-8", errors="ignore")


def _extract_legacy_doc(path: str) -> tuple[str, dict[str, Any], list[int]]:
    cmd = ["strings", "-n", "4", path]
    try:
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=60)
        text = proc.stdout or ""
    except Exception:
        text = _decode_text_best_effort(path)
    return clean_markdown(text), {"source": "strings_doc_fallback", "page_count": 1}, []


def _extract_office_zip(path: str, ext: str) -> tuple[str, dict[str, Any], list[int]]:
    wanted_prefixes = {
        ".pptx": ("ppt/slides/", "ppt/notesSlides/"),
        ".xlsx": ("xl/sharedStrings.xml", "xl/worksheets/"),
    }.get(ext, ())
    parts = []
    try:
        with zipfile.ZipFile(path) as zf:
            for name in sorted(zf.namelist()):
                if not any(name.startswith(prefix) for prefix in wanted_prefixes):
                    continue
                if not name.endswith(".xml"):
                    continue
                raw = zf.read(name).decode("utf-8", errors="ignore")
                text = _strip_xml_text(raw)
                if text:
                    parts.append("## " + name + "\n\n" + text)
    except Exception as exc:
        raise RuntimeError("Office 文件解析失败: %s" % exc)
    return "\n\n".join(parts), {"source": "office_zip_xml", "page_count": 1}, []


def _strip_xml_text(raw: str) -> str:
    raw = re.sub(r"<[^>]+>", "\n", raw)
    raw = raw.replace("&lt;", "<").replace("&gt;", ">").replace("&amp;", "&")
    lines = [line.strip() for line in raw.splitlines()]
    return "\n".join(line for line in lines if line)


def _parse_task_stage(task_id: str | None, progress: int, stage: str, status: str | None = None) -> None:
    if not task_id:
        return
    data: dict[str, Any] = {"progress": progress, "stage": stage}
    if status:
        data["status"] = status
    _update(task_id, **data)


def _extract_pdf(pdf_path: str, task_id: str | None = None) -> tuple[str, dict[str, Any], list[int]]:
    _parse_task_stage(task_id, 35, "正在解析 PDF 文本层", status="extracting")
    try:
        info = subprocess.run(["pdfinfo", pdf_path], capture_output=True, text=True, timeout=30)
        page_match = re.search(r"^Pages:\s*(\d+)", info.stdout or "", re.MULTILINE)
        page_count = int(page_match.group(1)) if page_match else 0
    except Exception:
        page_count = 0
    try:
        text_run = subprocess.run(["pdftotext", "-layout", pdf_path, "-"], capture_output=True, text=True, timeout=60)
        extracted_text = (text_run.stdout or "").strip()
        stdout = (text_run.stdout or "")[-6000:]
        stderr = (text_run.stderr or "")[-6000:]
    except Exception as exc:
        extracted_text = ""
        stdout = ""
        stderr = str(exc)
    pages = [part.strip() for part in re.split(r"\f+", extracted_text) if part.strip()]
    if pages:
        page_count = page_count or len(pages)
        blocks = [f"# {Path(pdf_path).name}", "", f"- Total pages: {page_count}", ""]
        for index, text in enumerate(pages, start=1):
            blocks.extend([f"## Page {index}", "", text, ""])
        markdown = "\n".join(blocks).strip()
    else:
        markdown = f"# {Path(pdf_path).name}\n\n- Total pages: {page_count or 1}"
    ocr_pages: list[int] = []
    text_fallback = {"used": bool(extracted_text), "chars": len(extracted_text), "reason": "poppler_primary"}
    digest = _sha256(pdf_path)
    image_infos = _pdf_image_infos(pdf_path, page_count)
    if not pages or _pdf_text_looks_garbled(markdown):
        forced_pages = range(1, min(page_count or 1, 3) + 1)
        existing_pages = {int(info.get("page") or 0) for info in image_infos}
        for page in forced_pages:
            if page in existing_pages:
                continue
            image_infos.append({
                "page": page,
                "num": 0,
                "type": "rendered_page",
                "width": 0,
                "height": 0,
                "pixels": PDF_MIN_VISUAL_IMAGE_PIXELS,
                "raw": "forced_visual_due_to_garbled_text",
            })
    image_pages = _pdf_image_pages_from_infos(image_infos, page_count)
    if image_pages:
        _parse_task_stage(
            task_id,
            52,
            f"检测到 PDF 图片，正在调用 MiniCPM 分析 {len(image_pages)} 页",
            status="ocr",
        )
    visual_pages = _analyze_pdf_image_pages(pdf_path, digest, image_infos, task_id=task_id)
    if not extracted_text.strip() and not visual_pages:
        raise RuntimeError("扫描 PDF 未识别到可用文字；MiniCPM 与 CPU OCR 均未返回文本")
    if visual_pages:
        markdown = _inject_pdf_visual_markdown(markdown, visual_pages)
    metadata = {
        "source": "poppler_pdf_hybrid",
        "parse_version": PDF_PARSE_VERSION,
        "page_count": page_count,
        "text_fallback": text_fallback,
        "image_count": len(image_infos),
        "image_pages": image_pages,
        "image_infos": image_infos[:80],
        "minicpm_required": bool(image_pages),
        "minicpm_attempted_pages": sorted(int(page) for page in visual_pages.keys()),
        "visual_pages": sorted(int(page) for page in visual_pages.keys()),
        "stdout": stdout,
        "stderr": stderr,
    }
    return markdown, metadata, ocr_pages


def _markdown_body_chars(markdown: str) -> int:
    body = re.sub(r"(?m)^#.*$", "", markdown or "")
    body = re.sub(r"(?m)^- (?:Total pages|Text pages|OCR pages|Processing time):.*$", "", body)
    body = re.sub(r"(?m)^---$", "", body)
    body = re.sub(r"(?m)^## Page.*$", "", body)
    return len(re.sub(r"\s+", "", body))


def _pdf_text_looks_garbled(markdown: str) -> bool:
    body = re.sub(r"(?m)^#.*$", "", markdown or "")
    compact = re.sub(r"\s+", "", body)
    if len(compact) < 20:
        return False
    bad = compact.count("?") + compact.count("�")
    cjk = len(re.findall(r"[\u4e00-\u9fff]", compact))
    return bad >= max(8, int(len(compact) * 0.22)) and cjk < max(6, int(len(compact) * 0.08))


def _merge_pdftotext_fallback(pdf_path: str, markdown: str, page_count: int) -> tuple[str, dict[str, Any]]:
    fallback = {"used": False, "chars": 0, "reason": ""}
    try:
        proc = subprocess.run(
            ["pdftotext", "-layout", pdf_path, "-"],
            capture_output=True,
            text=True,
            timeout=30,
        )
    except Exception as exc:
        fallback["reason"] = str(exc)
        return markdown, fallback

    text = (proc.stdout or "").strip()
    fallback["chars"] = len(re.sub(r"\s+", "", text))
    current_chars = _markdown_body_chars(markdown)
    if proc.returncode != 0 or fallback["chars"] <= max(current_chars + 80, current_chars * 2):
        fallback["reason"] = "opendataloader_text_sufficient"
        return markdown, fallback

    pages = [p.strip() for p in re.split(r"\f+", text) if p.strip()]
    if not pages:
        fallback["reason"] = "pdftotext_empty"
        return markdown, fallback

    header = [
        f"# {Path(pdf_path).name}",
        "",
        f"- Total pages: {page_count or len(pages)}",
        "- Text source: pdftotext fallback merged with OpenDataLoader",
        "",
        "---",
    ]
    for index, page_text in enumerate(pages, start=1):
        header.extend(["", f"## Page {index} (source: pdftotext, {len(page_text)} chars)", "", page_text])
    fallback["used"] = True
    fallback["reason"] = f"pdftotext_more_complete:{current_chars}->{fallback['chars']}"
    return "\n".join(header).strip(), fallback


def _pdf_image_infos(pdf_path: str, page_count: int) -> list[dict[str, Any]]:
    try:
        proc = subprocess.run(
            ["pdfimages", "-list", pdf_path],
            capture_output=True,
            text=True,
            timeout=30,
        )
    except Exception:
        return []
    if proc.returncode != 0:
        return []

    infos: list[dict[str, Any]] = []
    for raw in (proc.stdout or "").splitlines():
        line = raw.strip()
        if not line or line.startswith("page ") or line.startswith("---"):
            continue
        parts = line.split()
        if len(parts) < 5:
            continue
        try:
            page = int(parts[0])
            number = int(parts[1])
            width = int(parts[3])
            height = int(parts[4])
        except Exception:
            continue
        if page <= 0 or (page_count and page > page_count):
            continue
        infos.append({
            "page": page,
            "num": number,
            "type": parts[2],
            "width": width,
            "height": height,
            "pixels": width * height,
            "raw": line,
        })

    return infos


def _pdf_image_pages_from_infos(image_infos: list[dict[str, Any]], page_count: int = 0) -> list[int]:
    pages: list[int] = []
    for info in image_infos:
        try:
            page = int(info.get("page") or 0)
        except Exception:
            continue
        if page > 0 and page not in pages:
            pages.append(page)
    if not pages:
        return []
    if page_count:
        pages = [page for page in pages if page <= page_count]
    return sorted(pages)


def _pdf_image_pages(pdf_path: str, page_count: int) -> list[int]:
    return _pdf_image_pages_from_infos(_pdf_image_infos(pdf_path, page_count), page_count)


def _pdf_required_visual_prompt(page: int) -> str:
    return (
        f"?? PDF ? {page} ?????????????????????????????????????????????\n"
        "- ?????\n"
        "- ???\n"
        "- ?????\n"
        "- ???\n"
        "- ???????\n"
        "???\n"
        "1. ??????????????????????\n"
        "2. ???????? 09:30 ????????????????????????\n"
        "3. ???????????????????????\n"
        "4. ???????????????????"
    )


def _pdf_embedded_image_prompt(page: int, index: int, total: int) -> str:
    return (
        f"??? PDF ? {page} ?????? {index}/{total} ??????"
        "????????????????????????????????????"
        "??????????????????????????????????????"
        "???????????????????????????????????????? 120 ????"
    )


def _extract_pdf_page_images(pdf_path: str, sha256: str, page: int, image_infos: list[dict[str, Any]]) -> list[str]:
    page_infos = [info for info in image_infos if int(info.get("page") or 0) == page]
    if not page_infos:
        return []

    image_dir = VISUAL_CACHE_DIR / f"{sha256}_{PDF_VISUAL_CACHE_VERSION}_page_{page}_images"
    image_dir.mkdir(parents=True, exist_ok=True)
    existing = sorted(p for p in image_dir.glob("img-*") if p.is_file() and p.stat().st_size > 0)
    if not existing:
        prefix = image_dir / "img"
        proc = subprocess.run(
            ["pdfimages", "-png", "-f", str(page), "-l", str(page), pdf_path, str(prefix)],
            capture_output=True,
            text=True,
            timeout=90,
        )
        if proc.returncode != 0:
            return []
        existing = sorted(p for p in image_dir.glob("img-*") if p.is_file() and p.stat().st_size > 0)

    ranked: list[tuple[int, int, str]] = []
    for path in existing:
        pixels = 0
        try:
            from PIL import Image

            with Image.open(path) as img:
                width, height = img.size
                pixels = width * height
        except Exception:
            pixels = path.stat().st_size
        if pixels < PDF_MIN_VISUAL_IMAGE_PIXELS:
            continue
        ranked.append((pixels, path.stat().st_size, str(path)))

    ranked.sort(reverse=True)
    return [path for _, _, path in ranked[:PDF_MAX_VISUAL_IMAGES_PER_PAGE]]


def _analyze_pdf_image_pages(
    pdf_path: str,
    sha256: str,
    image_infos: list[dict[str, Any]],
    task_id: str | None = None,
) -> dict[int, str]:
    results: dict[int, str] = {}
    image_pages = _pdf_image_pages_from_infos(image_infos)
    total_pages = max(1, len(image_pages))
    for page in image_pages:
        cached = _cached_pdf_page_visual_by_sha(sha256, page)
        if cached:
            results[page] = cached
            continue
        try:
            page_index = image_pages.index(page) + 1
            _parse_task_stage(
                task_id,
                min(78, 52 + int(24 * page_index / total_pages)),
                f"正在调用 MiniCPM 分析 PDF 第 {page} 页图片",
                status="ocr",
            )
            extracted_images = _extract_pdf_page_images(pdf_path, sha256, page, image_infos)
            parts: list[str] = []
            if extracted_images:
                total_images = len(extracted_images)
                for index, image_path in enumerate(extracted_images, start=1):
                    _parse_task_stage(
                        task_id,
                        min(78, 52 + int(24 * (page_index - 1) / total_pages)),
                        f"MiniCPM 正在分析第 {page} 页图片 {index}/{total_images}",
                        status="ocr",
                    )
                    try:
                        _, payload = _run_minicpm_image(
                            image_path,
                            _pdf_embedded_image_prompt(page, index, total_images),
                            max_tokens=96,
                        )
                        text = (payload.get("text") or "").strip()
                        if text:
                            parts.append(f"#### 第{index}张内嵌图片\n{text}")
                    except Exception as image_exc:
                        try:
                            text = _run_cpu_ocr(image_path)
                            parts.append(f"#### 第{index}张内嵌图片\nCPU OCR（MiniCPM 不可用）\n{text}")
                        except Exception as fallback_exc:
                            logger.warning("PDF page %s image %s OCR failed: MiniCPM=%s; CPU=%s", page, index, image_exc, fallback_exc)

            if not parts:
                image_path = _render_pdf_page(pdf_path, sha256, page)
                try:
                    _, payload = _run_minicpm_image(image_path, _pdf_required_visual_prompt(page), max_tokens=128)
                    text = (payload.get("text") or "").strip()
                except Exception as image_exc:
                    text = _run_cpu_ocr(image_path)
                    text = "CPU OCR（MiniCPM 不可用）\n" + text
                if text:
                    parts.append("#### 页面截图\n" + text)

            text = "\n\n".join(part for part in parts if part.strip()).strip()
            if text:
                _write_pdf_page_visual_cache_by_sha(sha256, page, text)
                results[page] = text
        except Exception as exc:
            logger.warning("PDF page %s visual extraction failed: %s", page, exc)
    return results


def _inject_pdf_visual_markdown(markdown: str, visual_pages: dict[int, str]) -> str:
    if not visual_pages:
        return markdown
    pattern = re.compile(r"(^## Page\s+(\d+).*?$)", flags=re.MULTILINE | re.IGNORECASE)
    matches = list(pattern.finditer(markdown or ""))
    if not matches:
        blocks = [markdown.rstrip(), "", "## PDF 图片分析（MiniCPM-V）"]
        for page, text in sorted(visual_pages.items()):
            blocks.extend(["", f"### 第{page}页图片分析", text.strip()])
        return "\n".join(block for block in blocks if block is not None).strip()

    parts = []
    last = 0
    injected = set()
    for index, match in enumerate(matches):
        page = int(match.group(2))
        start = match.start()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(markdown)
        parts.append(markdown[last:start])
        block = markdown[start:end].rstrip()
        if page in visual_pages and "PDF 图片分析（MiniCPM-V）" not in block:
            block += (
                "\n\n### PDF 图片分析（MiniCPM-V）\n"
                + visual_pages[page].strip()
                + "\n"
            )
            injected.add(page)
        parts.append(block)
        last = end
    parts.append(markdown[last:])

    for page, text in sorted(visual_pages.items()):
        if page not in injected:
            parts.extend(["", f"## Page {page} 图片分析（MiniCPM-V）", text.strip()])
    return "".join(parts).strip()


def _extract_image(image_path: str) -> tuple[str, dict[str, Any], list[int]]:
    try:
        runner_image_path, payload = _run_minicpm_image(image_path)
        text = (payload.get("text") or "").strip()
        if not text:
            raise RuntimeError("MiniCPM 未识别到图片内容")
    except Exception as vision_exc:
        text = _run_cpu_ocr(image_path)
        markdown = "## 图片识别结果\n\n" + text
        metadata = {
            "source": "rapidocr_onnxruntime_fallback",
            "parse_version": IMAGE_PARSE_VERSION,
            "page_count": 1,
            "vision_error": str(vision_exc),
            "ocr_engine": "RapidOCR",
        }
        return markdown, metadata, [1]
    markdown = "## 图片识别结果\n\n" + text
    metadata = {
        "source": "minicpm_v2_6_visual",
        "parse_version": IMAGE_PARSE_VERSION,
        "page_count": 1,
        "normalized_image": runner_image_path != image_path,
        "vision_time_s": payload.get("vision_s"),
        "llm_time_s": payload.get("llm_s"),
        "tokens": payload.get("tokens"),
    }
    return markdown, metadata, [1]


def _run_cpu_ocr(image_path: str) -> str:
    try:
        from rapidocr_onnxruntime import RapidOCR
    except Exception as exc:
        raise RuntimeError("RapidOCR 依赖不可用") from exc

    try:
        results, _ = RapidOCR()(image_path)
    except Exception as exc:
        raise RuntimeError("RapidOCR 识别失败: " + str(exc)) from exc

    texts = []
    for result in results or []:
        if not isinstance(result, (list, tuple)) or len(result) < 2:
            continue
        text = str(result[1]).strip()
        if text:
            texts.append(text)
    if not texts:
        raise RuntimeError("RapidOCR 未识别到图片文字")
    return "\n".join(texts)


def _run_minicpm_image(image_path: str, prompt: str | None = None, max_tokens: int | None = None) -> tuple[str, dict[str, Any]]:
    if not MINICPM_ENABLED:
        raise RuntimeError("MiniCPM visual parsing disabled by configuration")
    if not MINICPM_PYTHON.exists():
        raise RuntimeError("MiniCPM Python 环境不存在")
    if not MINICPM_OCR_RUNNER.exists():
        raise RuntimeError("MiniCPM OCR runner 不存在")
    runner_image_path = _prepare_image_for_minicpm(image_path)
    env = os.environ.copy()
    env["LD_LIBRARY_PATH"] = str(MINICPM_LIB_DIR) + ":" + env.get("LD_LIBRARY_PATH", "")
    env.setdefault("RKLLM_LOG_LEVEL", "1")
    cmd = [str(MINICPM_PYTHON), str(MINICPM_OCR_RUNNER), runner_image_path]
    if prompt:
        cmd.append(prompt)
    if max_tokens:
        cmd.append(str(max_tokens))
    service_error = ""
    with _minicpm_lock:
        _pause_prompt_model()
        try:
            # Cold-start: spawn subprocess (loads model, does OCR, releases NPU)
            proc = subprocess.run(
                cmd,
                capture_output=True,
                text=True,
                timeout=900,
                env=env,
            )
        except Exception as exc:
            service_error = str(exc)
            # Fallback: try persistent service if available
            try:
                return runner_image_path, _run_minicpm_service(runner_image_path, prompt=prompt, max_tokens=max_tokens)
            except Exception:
                raise RuntimeError(service_error) from exc
        finally:
            _resume_prompt_model()
    raw = (proc.stdout or "") + "\n" + (proc.stderr or "")
    payload = {}
    for line in raw.splitlines():
        if line.startswith("###RESULT###"):
            try:
                payload = json.loads(line[len("###RESULT###"):])
            except Exception:
                payload = {}
    if proc.returncode != 0 and not payload:
        raise RuntimeError((proc.stderr or proc.stdout or "图片 OCR 失败").strip()[-1000:])
    if payload.get("error"):
        raise RuntimeError(payload["error"])
    if service_error:
        payload["fallback_from_service_error"] = service_error[-300:]
    return runner_image_path, payload


def _run_minicpm_service(image_path: str, prompt: str | None = None, max_tokens: int | None = None) -> dict[str, Any]:
    body = {
        "image_path": image_path,
        "prompt": prompt or "",
        "max_tokens": max_tokens or 0,
    }
    request = urllib.request.Request(
        LLM_06B_SERVICE_URL + "/analyze",
        data=json.dumps(body, ensure_ascii=False).encode("utf-8"),
        method="POST",
        headers={"Content-Type": "application/json"},
    )
    with urllib.request.urlopen(request, timeout=320) as resp:
        payload = json.loads(resp.read().decode("utf-8", errors="replace") or "{}")
    if payload.get("error"):
        raise RuntimeError(payload["error"])
    if not (payload.get("text") or "").strip():
        raise RuntimeError("MiniCPM 常驻服务未返回有效文本")
    return payload


def _prepare_image_for_minicpm(image_path: str) -> str:
    ext = Path(image_path).suffix.lower()
    normalized_path = PARSE_OUTPUT_DIR / (uuid.uuid4().hex + ".png")

    try:
        from PIL import Image, ImageOps

        with Image.open(image_path) as img:
            try:
                img.seek(0)
            except EOFError:
                pass
            img = ImageOps.exif_transpose(img).convert("RGB")
            img.thumbnail((1600, 1600))
            img.save(normalized_path, format="PNG")
        return str(normalized_path)
    except Exception as pil_exc:
        if ext in STANDARD_IMAGE_EXTS:
            return image_path

        cmd = [
            "ffmpeg",
            "-y",
            "-hide_banner",
            "-loglevel",
            "error",
            "-i",
            image_path,
            "-frames:v",
            "1",
            str(normalized_path),
        ]
        try:
            proc = subprocess.run(cmd, capture_output=True, text=True, timeout=60)
            if proc.returncode == 0 and normalized_path.exists():
                return str(normalized_path)
            detail = (proc.stderr or proc.stdout or str(pil_exc)).strip()
        except Exception as ffmpeg_exc:
            detail = f"{pil_exc}; {ffmpeg_exc}"
        raise RuntimeError("图片格式无法转换为 MiniCPM 可识别格式: " + detail[-500:])


def _prompt_model_admin(path: str, timeout: float = 8) -> None:
    request = urllib.request.Request(
        "http://127.0.0.1:8004" + path,
        data=b"{}",
        method="POST",
        headers={"Content-Type": "application/json"},
    )
    try:
        with urllib.request.urlopen(request, timeout=timeout) as resp:
            resp.read()
    except Exception:
        pass


def _pause_prompt_model() -> None:
    _prompt_model_admin("/admin/pause", timeout=12)


def _resume_prompt_model() -> None:
    _prompt_model_admin("/admin/resume", timeout=5)


def _read_pdf_metadata_from_markdown(markdown: str) -> tuple[int, list[int]]:
    total_match = re.search(r"Total pages:\s*(\d+)", markdown, flags=re.IGNORECASE)
    page_count = int(total_match.group(1)) if total_match else 0
    ocr_pages = []
    for match in re.finditer(r"## Page\s+(\d+)\s+\(source:\s*ocr", markdown, flags=re.IGNORECASE):
        try:
            ocr_pages.append(int(match.group(1)))
        except Exception:
            pass
    if not page_count:
        pages = [int(m.group(1)) for m in re.finditer(r"## Page\s+(\d+)", markdown, flags=re.IGNORECASE)]
        page_count = max(pages) if pages else 1
    return page_count, ocr_pages


def _pdf_visual_prompt(page: int) -> str:
    return (
        f"这是 PDF 第 {page} 页的页面截图。请重点分析页面中的图片、图表、插图、照片、版面布局，"
        "并说明这些视觉元素和页面文字之间的关系。不要逐字抄整页文字；只提取理解图文关系所必需的关键词。"
        "如果页面只有文字、没有明显独立图片或图表，也请明确说明。请用中文直接回答。"
    )


def _visual_pages_from_chunks(chunks: list[ParsedFileChunk], page_count: int) -> list[int]:
    pages = []
    for chunk in chunks:
        start = int(chunk.page_start or 0)
        end = int(chunk.page_end or start or 0)
        if start <= 0:
            continue
        for page in range(start, max(start, end) + 1):
            if page not in pages:
                pages.append(page)
            if len(pages) >= 3:
                return pages
    if not pages:
        max_page = max(1, min(int(page_count or 1), 3))
        pages = list(range(1, max_page + 1))
    return pages[:3]


def _render_pdf_page(pdf_path: str, sha256: str, page: int) -> str:
    prefix = VISUAL_CACHE_DIR / f"{sha256}_page_{page}"
    output = prefix.with_suffix(".png")
    if output.exists() and output.stat().st_size > 0:
        return str(output)
    cmd = [
        "pdftoppm",
        "-f", str(page),
        "-l", str(page),
        "-singlefile",
        "-png",
        "-r", "144",
        pdf_path,
        str(prefix),
    ]
    proc = subprocess.run(cmd, capture_output=True, text=True, timeout=90)
    if proc.returncode != 0 or not output.exists():
        raise RuntimeError((proc.stderr or proc.stdout or "PDF页面渲染失败").strip()[-800:])
    return str(output)


def _cached_pdf_page_visual(parsed: ParsedFile, page: int) -> str | None:
    return _cached_pdf_page_visual_by_sha(parsed.sha256, page)


def _cached_pdf_page_visual_by_sha(sha256: str, page: int) -> str | None:
    cache_path = VISUAL_CACHE_DIR / f"{sha256}_page_{page}_{PDF_VISUAL_CACHE_VERSION}.json"
    if not cache_path.exists():
        return None
    try:
        payload = json.loads(cache_path.read_text(encoding="utf-8"))
        text = (payload.get("text") or "").strip()
        if text.startswith("MiniCPM 图片分析失败：") or text.startswith("视觉分析失败："):
            return None
        return text or None
    except Exception:
        return None


def _write_pdf_page_visual_cache(parsed: ParsedFile, page: int, text: str) -> None:
    _write_pdf_page_visual_cache_by_sha(parsed.sha256, page, text)


def _write_pdf_page_visual_cache_by_sha(sha256: str, page: int, text: str) -> None:
    cache_path = VISUAL_CACHE_DIR / f"{sha256}_page_{page}_{PDF_VISUAL_CACHE_VERSION}.json"
    payload = {
        "version": PDF_VISUAL_CACHE_VERSION,
        "sha256": sha256,
        "page": page,
        "text": text,
    }
    cache_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _pdf_visual_context(db, parsed: ParsedFile, question: str, selected_chunks: list[ParsedFileChunk]) -> str:
    if not needs_visual_attachment_context(question):
        return ""
    file = db.query(File).filter(File.id == parsed.file_id, File.user_id == parsed.user_id).first()
    if not file or Path(file.filename or file.filepath).suffix.lower() != ".pdf":
        return ""
    if not os.path.exists(file.filepath):
        return ""

    pages = _visual_pages_from_chunks(selected_chunks, parsed.page_count or 1)
    parts = []
    for page in pages:
        try:
            text = _cached_pdf_page_visual(parsed, page)
            if not text:
                image_path = _render_pdf_page(file.filepath, parsed.sha256, page)
                _, payload = _run_minicpm_image(image_path, _pdf_visual_prompt(page), max_tokens=180)
                text = (payload.get("text") or "").strip()
                if text:
                    _write_pdf_page_visual_cache(parsed, page, text)
            if text:
                parts.append(f"### 第{page}页视觉分析\n{text}")
        except Exception as exc:
            parts.append(f"### 第{page}页视觉分析\n视觉分析失败：{exc}")
    if not parts:
        return ""
    return "## PDF页面视觉分析（MiniCPM-V）\n" + "\n\n".join(parts)


def clean_markdown(text: str) -> str:
    text = (text or "").replace("\r\n", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text)
    lines = []
    seen = {}
    for raw in text.splitlines():
        line = raw.rstrip()
        key = re.sub(r"\s+", "", line)
        if key and len(key) > 12:
            seen[key] = seen.get(key, 0) + 1
            if seen[key] > 3:
                continue
        lines.append(line)
    return "\n".join(lines).strip()


def make_chunks(markdown: str) -> list[dict[str, Any]]:
    page_blocks = _split_page_blocks(markdown)
    if page_blocks:
        chunks = []
        buffer = ""
        start_page = page_blocks[0][0]
        end_page = start_page
        for page, content in page_blocks:
            if buffer and len(buffer) + len(content) > CHUNK_SIZE:
                chunks.append(_chunk_record(buffer, start_page, end_page))
                buffer = ""
                start_page = page
            buffer += ("\n\n" if buffer else "") + content
            end_page = page
        if buffer:
            chunks.append(_chunk_record(buffer, start_page, end_page))
        return chunks

    chunks = []
    pos = 0
    index = 0
    while pos < len(markdown):
        part = markdown[pos:pos + CHUNK_SIZE]
        chunks.append(_chunk_record(part, index + 1, index + 1))
        pos += max(1, CHUNK_SIZE - CHUNK_OVERLAP)
        index += 1
    return chunks


def _split_page_blocks(markdown: str) -> list[tuple[int, str]]:
    matches = list(re.finditer(r"^## Page\s+(\d+).*?$", markdown, flags=re.MULTILINE | re.IGNORECASE))
    blocks = []
    for i, match in enumerate(matches):
        page = int(match.group(1))
        start = match.start()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(markdown)
        content = markdown[start:end].strip()
        if content:
            blocks.append((page, content))
    return blocks


def _chunk_record(content: str, page_start: int, page_end: int) -> dict[str, Any]:
    content = content.strip()
    summary = re.sub(r"\s+", " ", content)[:260]
    return {"content": content, "summary": summary, "page_start": page_start, "page_end": page_end}


def _normalize_file_ids(value) -> list[int]:
    if value is None:
        return []
    if isinstance(value, (int, str)):
        value = [value]
    result = []
    if isinstance(value, list):
        for item in value:
            try:
                file_id = int(item)
            except Exception:
                continue
            if file_id > 0 and file_id not in result:
                result.append(file_id)
    return result


def get_attachment_bundle(user_id: int, file_ids, question: str, include_visual: bool = True) -> dict[str, Any]:
    ids = _normalize_file_ids(file_ids)
    if not ids:
        return {"file_ids": [], "tasks": [], "pending": [], "errors": [], "context": "", "ready_count": 0}

    db = SessionLocal()
    try:
        tasks = ensure_parse_tasks(ids, user_id)
        pending = [task for task in tasks if task["status"] not in ("done", "error")]
        errors = [task for task in tasks if task["status"] == "error"]

        files = db.query(File).filter(File.user_id == user_id, File.id.in_(ids)).all()
        file_by_id = {file.id: file for file in files}
        parsed_items = db.query(ParsedFile).filter(
            ParsedFile.user_id == user_id,
            ParsedFile.file_id.in_(ids),
        ).order_by(ParsedFile.id.desc()).all()
        by_file: dict[int, ParsedFile] = {}
        for parsed in parsed_items:
            file = file_by_id.get(parsed.file_id)
            if file and not _cache_compatible(parsed, file.filename):
                continue
            by_file.setdefault(parsed.file_id, parsed)

        context_parts = []
        for file_id in ids:
            parsed = by_file.get(file_id)
            if not parsed:
                continue
            context_parts.append(_context_for_parsed(db, parsed, question, include_visual=include_visual))

        context = "\n\n".join(part for part in context_parts if part).strip()
        if len(context) > MAX_CONTEXT_CHARS:
            context = context[:MAX_CONTEXT_CHARS] + "\n\n[附件内容过长，已截断到关键片段。]"
        return {
            "file_ids": ids,
            "tasks": tasks,
            "pending": pending,
            "errors": errors,
            "context": context,
            "ready_count": len(context_parts),
        }
    finally:
        db.close()


def _context_for_parsed(db, parsed: ParsedFile, question: str, include_visual: bool = True) -> str:
    chunks = db.query(ParsedFileChunk).filter(
        ParsedFileChunk.parsed_file_id == parsed.id,
        ParsedFileChunk.user_id == parsed.user_id,
    ).order_by(ParsedFileChunk.chunk_index).all()
    if not chunks:
        return parsed.markdown[:MAX_CONTEXT_CHARS]

    selected = _select_relevant_chunks(chunks, question)
    selected_text = "\n\n".join((chunk.content or "") for chunk in selected)
    run_visual = include_visual and should_run_visual_analysis(question, selected_text)
    visual_context = _pdf_visual_context(db, parsed, question, selected) if run_visual else ""
    lines = [
        f"## 附件：{parsed.title or ('file-' + str(parsed.file_id))}",
        f"- 页数：{parsed.page_count or '未知'}",
    ]
    if visual_context:
        lines.append("\n" + visual_context)
    for chunk in selected:
        page_hint = ""
        if chunk.page_start and chunk.page_end:
            page_hint = f"（第{chunk.page_start}-{chunk.page_end}页）"
        lines.append(f"\n### 相关内容{page_hint}\n{chunk.content.strip()}")
    return "\n".join(lines)


def _select_relevant_chunks(chunks: list[ParsedFileChunk], question: str) -> list[ParsedFileChunk]:
    if not chunks:
        return []
    keywords = set(re.findall(r"[A-Za-z0-9_\-/]{2,}|[\u4e00-\u9fff]{2,}", question or ""))
    if not keywords:
        return chunks[: min(4, len(chunks))]

    scored = []
    for chunk in chunks:
        text = chunk.content or ""
        score = sum(1 for kw in keywords if kw in text)
        score += min(len(text), 1200) / 6000
        scored.append((score, chunk.chunk_index, chunk))
    scored.sort(key=lambda item: (-item[0], item[1]))
    selected = [item[2] for item in scored[:4]]
    selected.sort(key=lambda chunk: chunk.chunk_index)
    return selected
