"""Render Markdown into DOCX with python-docx.

This renderer intentionally supports a small, predictable Markdown subset used
by generated business documents. It does not try to be a full Markdown engine.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


PAGE_BREAK_MARKERS = {"<<<PAGE_BREAK>>>", "[[PAGE_BREAK]]", "<!-- pagebreak -->", "<!-- page break -->"}


def markdown_to_docx(markdown: str, output_path: str | Path, title: str = "") -> None:
    doc = Document()
    _ensure_document_styles(doc)
    text = (markdown or "").replace("\r\n", "\n").replace("\r", "\n")
    # Normalize: ensure markdown headings start on their own line (fixes 35B output)
    text = re.sub(r"([^\n])(#{1,3}\s)", r"\1\n\n\2", text)
    # Strip leading natural-language intro before first # heading
    first_heading = re.search(r"^#{1,3}\s", text, re.MULTILINE)
    if first_heading and first_heading.start() > 0:
        text = text[first_heading.start():]
    lines = text.split("\n")
    index = 0
    emitted_title = False

    while index < len(lines):
        raw = lines[index]
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped:
            index += 1
            continue

        if stripped in PAGE_BREAK_MARKERS:
            _add_page_break(doc)
            index += 1
            continue

        if stripped == "---":
            doc.add_paragraph("─" * 24)
            index += 1
            continue

        if stripped.startswith("```"):
            block, index = _consume_code_block(lines, index)
            if block:
                doc.add_paragraph(block)
            continue

        if _is_table_start(lines, index):
            rows, index = _consume_table(lines, index)
            _add_table(doc, rows)
            continue

        heading = re.match(r"^(#{1,3})\s+(.+?)\s*$", stripped)
        if heading:
            level = len(heading.group(1))
            content = heading.group(2).strip()
            if level == 1:
                paragraph = doc.add_heading(content, level=1)
                emitted_title = True
            else:
                paragraph = doc.add_heading(content, level=level)
            _style_heading_runs(paragraph, level)
            if content == "目录":
                _add_toc(doc)
            index += 1
            continue

        ordered = re.match(r"^\d+[.)]\s+(.+)$", stripped)
        if ordered:
            _add_inline_paragraph(doc, ordered.group(1), style="List Number")
            index += 1
            continue

        unordered = re.match(r"^[-*+]\s+(.+)$", stripped)
        if unordered:
            _add_inline_paragraph(doc, unordered.group(1), style="List Bullet")
            index += 1
            continue

        if stripped.startswith(">"):
            quote = stripped.lstrip("> ").strip()
            _add_inline_paragraph(doc, quote)
            index += 1
            continue

        _add_inline_paragraph(doc, stripped)
        index += 1

    if title and not emitted_title:
        first = doc.paragraphs[0]._p if doc.paragraphs else None
        heading = doc.add_heading(title, level=1)
        _style_heading_runs(heading, 1)
        if first is not None:
            doc._body._element.remove(heading._p)
            doc._body._element.insert(0, heading._p)

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def _ensure_document_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(11)
    _set_style_east_asia(normal, "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, font_size in (("Heading 1", 18), ("Heading 2", 15), ("Heading 3", 13)):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(font_size)
        style.font.bold = True
        _set_style_east_asia(style, "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(11)
        _set_style_east_asia(style, "Microsoft YaHei")

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = "Consolas"
        code.font.size = Pt(9)
        _set_style_east_asia(code, "Microsoft YaHei")


def _set_style_east_asia(style, font_name: str) -> None:
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)


def _set_run_font(run, font_name: str = "Microsoft YaHei", size: int = 11) -> None:
    run.font.name = font_name
    run.font.size = Pt(size)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font_name)


def _style_heading_runs(paragraph, level: int) -> None:
    sizes = {1: 18, 2: 15, 3: 13}
    for run in paragraph.runs:
        _set_run_font(run, "Microsoft YaHei", sizes.get(level, 11))
        run.bold = True


def _add_toc(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'

    fld_separate = OxmlElement("w:fldChar")
    fld_separate.set(qn("w:fldCharType"), "separate")

    placeholder = OxmlElement("w:t")
    placeholder.text = "请在 Word 中更新目录"

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_separate)
    run._r.append(placeholder)
    run._r.append(fld_end)


def _add_page_break(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    run.add_break(WD_BREAK.PAGE)


def _consume_code_block(lines: list[str], start: int) -> tuple[str, int]:
    index = start + 1
    block: list[str] = []
    while index < len(lines):
        if lines[index].strip().startswith("```"):
            return "\n".join(block).strip("\n"), index + 1
        block.append(lines[index].rstrip())
        index += 1
    return "\n".join(block).strip("\n"), index


def _is_table_start(lines: list[str], index: int) -> bool:
    if index + 1 >= len(lines):
        return False
    current = lines[index].strip()
    next_line = lines[index + 1].strip()
    return _is_pipe_table_row(current) and (
        _is_table_separator(next_line) or _is_pipe_table_row(next_line)
    )


def _consume_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines):
        stripped = lines[index].strip()
        if not _is_pipe_table_row(stripped) and not _is_table_separator(stripped):
            break
        if _is_table_separator(stripped):
            index += 1
            continue
        rows.append(_split_table_row(stripped))
        index += 1
    return rows, index


def _is_pipe_table_row(line: str) -> bool:
    return line.startswith("|") and line.endswith("|") and line.count("|") >= 2


def _is_table_separator(line: str) -> bool:
    return bool(
        re.match(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?$", line.strip())
    )


def _split_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip("|").split("|")]


def _add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    for r_index, row in enumerate(rows):
        for c_index in range(columns):
            cell = table.cell(r_index, c_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            run = paragraph.add_run(row[c_index] if c_index < len(row) else "")
            _set_run_font(run)
            if r_index == 0:
                run.bold = True
                _shade_cell(cell, "D9EAF7")
    doc.add_paragraph()


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def _add_inline_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    paragraph = doc.add_paragraph(style=style)
    _add_inline_runs(paragraph, text)


def _add_inline_runs(paragraph, text: str) -> None:
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")
    pos = 0
    for match in pattern.finditer(text or ""):
        if match.start() > pos:
            _set_run_font(paragraph.add_run(text[pos:match.start()]))
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        _set_run_font(run)
        pos = match.end()
    if pos < len(text or ""):
        _set_run_font(paragraph.add_run(text[pos:]))
