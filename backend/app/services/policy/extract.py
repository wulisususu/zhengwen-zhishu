"""Text extraction for imported policy sources (HTML page, txt/md, docx).

Only the stdlib HTML parser is used. On gov.cn article pages the body lives in
a known container (``#UCAP-CONTENT`` / ``.pages_content`` / TRS editor div);
when none is found the whole page body is used and the result is flagged
``reliable=False`` so every clause gets ``needs_review``. Text is never
rewritten, only whitespace-normalised.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from html import unescape
from html.parser import HTMLParser

from app.services.policy import errors
from app.services.policy.errors import PolicyError

CONTENT_IDS = {"ucap-content"}
CONTENT_CLASSES = {"pages_content", "trs_editor", "trs_editor_view", "article-content", "view"}
BLOCK_TAGS = {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article", "table"}
SKIP_TAGS = {"script", "style", "noscript", "head"}
VOID_TAGS = {"br", "img", "hr", "meta", "link", "input"}
TEXT_EXTENSIONS = {".txt", ".md"}
HTML_EXTENSIONS = {".html", ".htm"}
DOCX_EXTENSIONS = {".docx"}


@dataclass
class ExtractedDocument:
    text: str
    title: str = ""
    meta: dict = field(default_factory=dict)
    reliable: bool = True


class _PageParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.title = ""
        self.meta: dict[str, str] = {}
        self._in_title = False
        self._skip_depth = 0
        self._stack: list[tuple[str, bool]] = []  # (tag, opens_content_container)
        self._content_depth = 0
        self.content_found = False
        self._content_parts: list[str] = []
        self._body_parts: list[str] = []

    def _is_container(self, attrs) -> bool:
        attr = {k.lower(): (v or "") for k, v in attrs}
        if attr.get("id", "").lower() in CONTENT_IDS:
            return True
        classes = {c.lower() for c in attr.get("class", "").split()}
        return bool(classes & CONTENT_CLASSES)

    def _newline(self):
        self._body_parts.append("\n")
        if self._content_depth:
            self._content_parts.append("\n")

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        if tag == "meta":
            attr = {k.lower(): (v or "") for k, v in attrs}
            name = attr.get("name") or attr.get("property")
            if name and attr.get("content"):
                self.meta[name] = attr["content"].strip()
            return
        if tag == "title":
            self._in_title = True
        if tag in SKIP_TAGS:
            self._skip_depth += 1
        if tag in BLOCK_TAGS:
            self._newline()
        if tag in VOID_TAGS:
            return
        opens = self._is_container(attrs)
        if opens:
            self._content_depth += 1
            self.content_found = True
        self._stack.append((tag, opens))

    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag == "title":
            self._in_title = False
        if tag in SKIP_TAGS and self._skip_depth:
            self._skip_depth -= 1
        if tag in BLOCK_TAGS:
            self._newline()
        # Pop to the matching tag (tolerates unclosed children).
        for i in range(len(self._stack) - 1, -1, -1):
            if self._stack[i][0] == tag:
                for _t, opens in self._stack[i:]:
                    if opens and self._content_depth:
                        self._content_depth -= 1
                del self._stack[i:]
                break

    def handle_data(self, data):
        if self._in_title:
            self.title += data
            return
        if self._skip_depth:
            return
        self._body_parts.append(data)
        if self._content_depth:
            self._content_parts.append(data)

    def text(self, parts):
        # Collapse only HTML layout whitespace; full-width spaces (　) are part of the original text.
        lines = [re.sub(r"[ \t ]+", " ", line).strip(" \t ") for line in "".join(parts).splitlines()]
        return "\n".join(line for line in lines if line)


def extract_html(raw: str) -> ExtractedDocument:
    parser = _PageParser()
    parser.feed(raw)
    parser.close()
    reliable = parser.content_found
    text = parser.text(parser._content_parts if reliable else parser._body_parts)
    title = parser.meta.get("ArticleTitle") or unescape(parser.title).strip()
    return ExtractedDocument(text=text, title=title, meta=parser.meta, reliable=reliable and bool(text))


def decode_bytes(data: bytes, declared: str | None = None) -> str:
    candidates = [c for c in (declared, "utf-8", "gb18030") if c]
    for enc in candidates:
        try:
            return data.decode(enc)
        except (LookupError, UnicodeDecodeError):
            continue
    return data.decode("utf-8", errors="replace")


def _html_charset(data: bytes) -> str | None:
    head = data[:4096].decode("ascii", errors="ignore")
    m = re.search(r"charset=[\"']?([A-Za-z0-9_\-]+)", head, re.IGNORECASE)
    return m.group(1) if m else None


def extract_file(filename: str, data: bytes) -> ExtractedDocument:
    name = (filename or "").lower()
    ext = name[name.rfind("."):] if "." in name else ""
    if ext in TEXT_EXTENSIONS:
        text = decode_bytes(data)
        lines = [line.strip() for line in text.splitlines()]
        return ExtractedDocument(text="\n".join(line for line in lines if line))
    if ext in HTML_EXTENSIONS:
        return extract_html(decode_bytes(data, _html_charset(data)))
    if ext in DOCX_EXTENSIONS:
        try:
            import docx  # python-docx, listed in requirements.txt
        except ImportError:
            raise PolicyError(errors.UNSUPPORTED_FILE, "服务器未安装 python-docx，暂不能解析 .docx，请上传 .txt 或 .html。") from None
        try:
            document = docx.Document(io.BytesIO(data))
        except Exception:
            raise PolicyError(errors.EXTRACT_FAILED, "无法读取该 .docx 文件。") from None
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        return ExtractedDocument(text="\n".join(paragraphs))
    raise PolicyError(errors.UNSUPPORTED_FILE, "仅支持 .txt / .md / .html / .docx；PDF 暂不支持，请先转换。")
