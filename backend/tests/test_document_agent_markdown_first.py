import sys
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


from app.services.document_agent import (  # noqa: E402
    ModelEmptyResponseError,
    generate_agent_output_file,
    normalize_document_request,
)


class DocumentAgentMarkdownFirstTests(unittest.TestCase):
    def test_normalizes_light_markdown_document_request(self):
        request = normalize_document_request(
            {
                "filename": "项目周会纪要.docx",
                "format": "docx",
                "content_type": "markdown",
                "content": "# 项目周会会议纪要\n\n## 会议时间\n\n2026年6月8日",
            },
            fallback_instruction="创建项目周会纪要",
        )

        self.assertEqual("项目周会纪要.docx", request["filename"])
        self.assertEqual("docx", request["format"])
        self.assertEqual("markdown", request["content_type"])
        self.assertIn("# 项目周会会议纪要", request["content"])

    def test_renders_markdown_payload_without_content_schema(self):
        generated = generate_agent_output_file(
            "创建项目周会纪要",
            route={"primary_operation": "generate", "output_spec": {"format": "docx"}},
            content_schema={
                "filename": "项目周会纪要.docx",
                "format": "docx",
                "content_type": "markdown",
                "content": "# 项目周会会议纪要\n\n## 会议时间\n\n2026年6月8日\n\n| 任务 | 负责人 |\n| --- | --- |\n| 测试 | 张三 |",
            },
        )

        self.assertEqual("项目周会纪要.docx", generated["filename"])
        self.assertEqual("markdown", generated["document_request"]["content_type"])
        doc = Document(generated["filepath"])
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("项目周会会议纪要", text)
        self.assertEqual(1, len(doc.tables))

    def test_empty_model_output_still_generates_placeholder_docx(self):
        with patch(
            "app.services.document_agent._call_qwen_text",
            side_effect=ModelEmptyResponseError("empty response"),
        ):
            generated = generate_agent_output_file(
                "创建一个 Word 文档，标题：项目周会会议纪要",
                route={"primary_operation": "generate", "output_spec": {"format": "docx"}},
            )

        self.assertEqual(".docx", generated["filetype"])
        doc = Document(generated["filepath"])
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("文档生成失败", text)
        self.assertIn("请重新描述需求", text)

    def test_strips_thinking_trace_from_markdown_before_docx(self):
        leaked = (
            "# AI 文档生成平台软件需求规格说明书\n\n"
            "## 项目概述\n\n"
            "本项目旨在建设 AI 文档生成平台。\n\n"
            "<think>\n"
            "Thinking Process:\n"
            "Analyze the Request:\n"
            "Role: Chinese formal document writing assistant.\n"
            "</think>\n"
        )
        with patch("app.services.document_agent._call_qwen_text", return_value=leaked):
            generated = generate_agent_output_file(
                "我要开发一个AI文档生成平台，生成Word需求规格说明书",
                route={"primary_operation": "generate", "output_spec": {"format": "docx"}},
            )

        doc = Document(generated["filepath"])
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("AI 文档生成平台软件需求规格说明书", text)
        self.assertNotIn("<think>", text)
        self.assertNotIn("Thinking Process", text)
        self.assertNotIn("Analyze the Request", text)
        self.assertNotIn("Role: Chinese formal document writing assistant", text)

    def test_strips_duplicate_second_markdown_document(self):
        duplicated = (
            "# AI面试助手产品需求文档 (PRD)\n\n"
            "## 1. 产品背景\n\n"
            "第一版产品背景正文。\n\n"
            "## 2. 用户痛点\n\n"
            "第一版用户痛点正文。\n\n"
            "## 3. 功能需求\n\n"
            "| 模块 | 说明 |\n"
            "| --- | --- |\n"
            "| AI面试 | 自动分析候选人回答 |\n\n"
            "## 4. 页面设计说明\n\n"
            "第一版页面设计正文。\n\n"
            "## 5. 风险分析\n\n"
            "第一版风险分析正文。\n\n"
            "---\n\n"
            "# AI面试助手产品需求文档 (PRD)\n\n"
            "## 目录\n\n"
            "- 产品背景\n"
            "- 用户痛点\n\n"
            "## 1. 产品背景\n\n"
            "第二版重复正文。\n"
        )
        with patch("app.services.document_agent._call_qwen_text", return_value=duplicated):
            generated = generate_agent_output_file(
                "生成一份产品需求文档。产品名称：AI面试助手",
                route={"primary_operation": "generate", "output_spec": {"format": "docx"}},
            )

        doc = Document(generated["filepath"])
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertEqual(1, text.count("AI面试助手产品需求文档"))
        self.assertIn("第一版风险分析正文", text)
        self.assertNotIn("第二版重复正文", text)


if __name__ == "__main__":
    unittest.main()
