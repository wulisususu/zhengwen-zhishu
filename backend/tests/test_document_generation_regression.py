import sys
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


from app.services.document_agent import (  # noqa: E402
    ModelEmptyResponseError,
    generate_agent_output_file,
)


class DocumentGenerationRegressionTests(unittest.TestCase):
    def _generate(self, prompt: str, markdown: str):
        with patch("app.services.document_agent._call_qwen_text", return_value=markdown):
            result = generate_agent_output_file(
                prompt,
                route={"primary_operation": "generate", "output_spec": {"format": "docx"}},
            )
        self.assertEqual(".docx", result["filetype"])
        self.assertEqual("markdown", result["document_request"]["content_type"])
        self.assertTrue(Path(result["filepath"]).exists())
        return result, Document(result["filepath"])

    def _text(self, doc: Document) -> str:
        return "\n".join(p.text for p in doc.paragraphs)

    def test_01_resume(self):
        _, doc = self._generate("生成简历 Word 文档", "# 张三简历\n\n## 个人信息\n\n姓名：张三\n\n## 工作经历\n\n- AI平台工程师")
        self.assertIn("张三简历", self._text(doc))

    def test_02_meeting_minutes(self):
        _, doc = self._generate(
            "生成项目周会会议纪要 Word 文档",
            "# 项目周会会议纪要\n\n## 会议时间\n\n2026年6月8日\n\n## 任务分工\n\n| 任务 | 责任人 |\n| --- | --- |\n| 联调 | 张三 |",
        )
        self.assertIn("项目周会会议纪要", self._text(doc))
        self.assertEqual(1, len(doc.tables))

    def test_03_project_plan(self):
        _, doc = self._generate(
            "生成智慧校园AI学习平台项目实施方案",
            "# 智慧校园AI学习平台项目实施方案\n\n## 项目背景\n\n建设智慧校园AI学习平台。\n\n## 实施计划\n\n```mermaid\ngantt\n    title 实施计划\n```",
        )
        self.assertIn("智慧校园AI学习平台项目实施方案", self._text(doc))

    def test_04_notice(self):
        _, doc = self._generate("生成通知 Word 文档", "# 关于召开项目推进会的通知\n\n## 通知事项\n\n请相关人员按时参会。")
        self.assertIn("关于召开项目推进会的通知", self._text(doc))

    def test_05_research_report(self):
        _, doc = self._generate("生成研究报告 Word 文档", "# 2026年中国人工智能行业发展研究报告\n\n## 研究背景\n\n行业进入深化应用阶段。\n\n## 结论\n\n保持稳健增长。")
        self.assertIn("研究报告", self._text(doc))

    def test_06_training_document(self):
        _, doc = self._generate("生成培训文档 Word 文档", "# AI学习平台培训手册\n\n## 培训对象\n\n- 管理员\n- 教师\n\n## 培训内容\n\n1. 登录系统\n2. 使用AI助手")
        self.assertIn("AI学习平台培训手册", self._text(doc))

    def test_07_table_document(self):
        _, doc = self._generate("生成带表格文档", "# 预算报价\n\n| 序号 | 费用项 | 金额 |\n| --- | --- | --- |\n| 1 | 软件平台 | 38.00 |")
        self.assertEqual(1, len(doc.tables))
        self.assertEqual("费用项", doc.tables[0].cell(0, 1).text)

    def test_08_page_break_document(self):
        _, doc = self._generate("生成带分页文档", "# 第一部分\n\n正文\n\n<<<PAGE_BREAK>>>\n\n# 第二部分\n\n正文")
        self.assertIn("第二部分", self._text(doc))

    def test_09_long_document(self):
        body = "\n\n".join(f"## 第{i}章\n\n这是第{i}章正文，围绕项目建设、实施、风险和保障展开。" for i in range(1, 12))
        _, doc = self._generate("生成超长文档", "# 长篇项目报告\n\n" + body)
        self.assertIn("第11章", self._text(doc))

    def test_10_empty_model_output(self):
        with patch(
            "app.services.document_agent._call_qwen_text",
            side_effect=ModelEmptyResponseError("empty response"),
        ):
            result = generate_agent_output_file(
                "生成 Word 文档",
                route={"primary_operation": "generate", "output_spec": {"format": "docx"}},
            )
        doc = Document(result["filepath"])
        self.assertIn("文档生成失败", self._text(doc))


if __name__ == "__main__":
    unittest.main()
