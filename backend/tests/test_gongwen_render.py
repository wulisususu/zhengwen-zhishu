"""公文 IR → DOCX 版式自检（无需模型、无需数据库）。

用法：
    python tests/test_gongwen_render.py [输出目录]
逐项校验 GB/T 9704-2012 的关键版式参数，任何一项不符即退出码 1。
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from docx.shared import Mm, Pt

from app.gongwen import renderer
from app.gongwen.checker import check
from app.gongwen.schema import BodyPara, GongwenDoc

SAMPLE_BODY = [
    BodyPara("为深入贯彻全市政务服务工作会议精神，切实解决群众反映强烈的窗口服务问题，现就有关事项通知如下。", 0),
    BodyPara("一、提高思想认识", 1),
    BodyPara("窗口是政府联系群众的直接通道，服务质量直接影响政府形象。各单位要把窗口作风建设作为年度重点任务，主要负责同志亲自抓。", 0),
    BodyPara("二、明确重点任务", 1),
    BodyPara("（一）规范服务行为。统一着装、挂牌上岗、使用文明用语，严禁工作时间从事与业务无关的事项。", 2),
    BodyPara("（二）压缩办理时限。对材料齐全、符合法定形式的事项，当场办结率不低于百分之八十。", 2),
    BodyPara("三、加强组织保障", 1),
    BodyPara("建立月度通报机制，对群众投诉集中的窗口进行专项督查，结果纳入年度考核。", 0),
    BodyPara("特此通知。", 0),
]


def build_sample(doc_type: str = "通知") -> GongwenDoc:
    return GongwenDoc(
        doc_type=doc_type,
        title=f"XX市政务服务管理局关于加强政务服务中心窗口作风建设的{doc_type}",
        recipients=["各县（市、区）政务服务管理局", "局属各单位"],
        issuer_org="XX市政务服务管理局",
        signer_org="XX市政务服务管理局",
        sign_date="2026年10月24日",
        doc_number="X政服发〔2026〕12号",
        urgency="",
        body=list(SAMPLE_BODY),
        cc=["市委办公室", "市政府办公室"],
        print_org="XX市政务服务管理局办公室",
        print_date="2026年10月25日",
    )


def _all_text(doc) -> str:
    return "\n".join(p.text for p in doc.paragraphs)


def check_layout(path: Path, doc_type: str) -> list[str]:
    """逐项比对版式参数，返回问题清单。"""
    problems: list[str] = []
    doc = Document(str(path))
    sec = doc.sections[0]

    def near(value, expect, tol=0.3) -> bool:
        return abs(float(value) - float(expect)) < tol

    if not near(sec.page_width.mm, 210):
        problems.append(f"页宽 {sec.page_width.mm:.1f}mm ≠ 210mm")
    if not near(sec.page_height.mm, 297):
        problems.append(f"页高 {sec.page_height.mm:.1f}mm ≠ 297mm")
    if not near(sec.top_margin.mm, 37):
        problems.append(f"上边距 {sec.top_margin.mm:.1f}mm ≠ 37mm")
    if not near(sec.bottom_margin.mm, 35):
        problems.append(f"下边距 {sec.bottom_margin.mm:.1f}mm ≠ 35mm")
    if not near(sec.left_margin.mm, 28):
        problems.append(f"左边距 {sec.left_margin.mm:.1f}mm ≠ 28mm")
    if not near(sec.right_margin.mm, 26):
        problems.append(f"右边距 {sec.right_margin.mm:.1f}mm ≠ 26mm")

    text = _all_text(doc)
    if "特此通知" not in text and doc_type == "通知":
        problems.append("正文缺少结尾语「特此通知」")

    # 标题：2 号小标宋、居中
    title_ok = False
    for p in doc.paragraphs:
        if p.text.strip().startswith("XX市政务服务管理局关于加强") or "的" + doc_type in p.text:
            for run in p.runs:
                if run.font.size and abs(run.font.size.pt - 22) < 0.5 and run.font.name == "方正小标宋简体":
                    title_ok = True
            break
    if not title_ok:
        problems.append("标题未使用 2 号（22pt）方正小标宋简体")

    # 正文：3 号仿宋 + 首行缩进 2 字符
    body_ok = False
    for p in doc.paragraphs:
        if p.text.startswith("为深入贯彻"):
            for run in p.runs:
                if run.font.size and abs(run.font.size.pt - 16) < 0.5 and run.font.name == "仿宋_GB2312":
                    body_ok = True
            if body_ok and p.paragraph_format.first_line_indent is None:
                problems.append("正文段落缺少首行缩进")
            break
    if not body_ok:
        problems.append("正文未使用 3 号（16pt）仿宋_GB2312")

    # 一级标题：黑体
    h1_ok = False
    for p in doc.paragraphs:
        if p.text.strip() == "一、提高思想认识":
            h1_ok = any(r.font.name == "黑体" for r in p.runs)
            break
    if not h1_ok:
        problems.append("一级标题未使用黑体")

    # 二级标题：楷体
    h2_ok = False
    for p in doc.paragraphs:
        if p.text.strip().startswith("（一）规范服务行为"):
            h2_ok = any(r.font.name == "楷体_GB2312" for r in p.runs)
            break
    if not h2_ok:
        problems.append("二级标题未使用楷体_GB2312")

    # 红色分隔线
    red_line = False
    for p in doc.paragraphs:
        borders = p._p.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr")
        for pBdr in borders:
            for child in pBdr:
                if child.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}color") in ("E60000", "FF0000"):
                    red_line = True
    if not red_line:
        problems.append("缺少红色分隔线")

    # 页码
    footer_xml = sec.footer.paragraphs[0]._p.xml if sec.footer.paragraphs else ""
    if "PAGE" not in footer_xml:
        problems.append("页脚缺少 PAGE 页码域")

    # 抄送
    if "抄送：" not in text:
        problems.append("版记缺少抄送机关")

    return problems


def main() -> int:
    out_dir = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(tempfile.mkdtemp())
    out_dir.mkdir(parents=True, exist_ok=True)

    failed = False
    for doc_type in ("通知", "请示", "报告", "函", "纪要"):
        gw = build_sample(doc_type)
        if doc_type == "请示":
            gw.recipients = ["XX市人民政府"]
            gw.signer_name = "张某某"
        issues = check(gw)
        path = renderer.render(gw, out_dir / f"{doc_type}.docx")
        problems = check_layout(path, doc_type)
        errors = [i for i in issues if i.level == "error"]

        size_kb = path.stat().st_size / 1024
        status = "OK " if not problems else "FAIL"
        print(f"[{status}] {doc_type:<3} {path.name:<20} {size_kb:6.1f} KB  校核 error={len(errors)}")
        for p in problems:
            failed = True
            print(f"        - {p}")

    print(f"\n输出目录：{out_dir}")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
