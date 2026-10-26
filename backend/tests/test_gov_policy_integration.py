"""Real JWT/SQLite contracts for policy provenance and durable draft citations."""
import pytest
from test_gov_basic_api import app_client, create, action
from app.services.policy import library as policy_library
from app.services.policy.library import ImportRequest, PolicyLibraryService
from app.services.policy.settings import PolicySettings
from app.services.policy.store import PolicyStore
from app.services.policy.fetcher import SafeFetcher


@pytest.fixture()
def policy_client(app_client, tmp_path, monkeypatch):
    client, tokens, factory = app_client
    settings = PolicySettings(db_path=tmp_path / "policy.db")
    library = PolicyLibraryService(settings, PolicyStore(settings.db_path), SafeFetcher(settings.fetch))
    monkeypatch.setattr(policy_library, "_library", library)
    from app.routers.gov_policy import router
    client.app.include_router(router, prefix="/api/v1")
    imported = library.import_policy(ImportRequest(filename="测试政策.txt", title="测试政策（仅测试）", data="第一条 政务数据应当分类管理。\n第二条 办理结果应当公开。".encode()))
    assert imported.get("error") is None
    policy = imported["policy"]
    hit = library.search("分类管理", statuses=["unverified"])["hits"][0]
    return client, tokens, factory, library, policy, hit["citation_id"]


def link(client, token, doc, ids):
    return client.post(f"/api/v1/gov/documents/{doc['id']}/citations", headers=token,
                       json={"version": doc["version"], "citationIds": ids})


def test_selected_citation_survives_save_reload_and_has_server_provenance(policy_client):
    client, tokens, _, library, policy, cid = policy_client
    doc = create(client, tokens["writer"])
    response = link(client, tokens["writer"], doc, [cid])
    assert response.status_code == 200, response.text
    saved = response.json()["data"]
    ref = saved["policyRefs"][0]
    assert ref["id"] == cid and ref["title"] == policy["title"]
    assert ref["sourceUrl"] == policy["url"] and ref["policyId"] == policy["id"]
    assert ref["version"] == 1 and ref["status"] == "unverified"
    path = f"/api/v1/gov/documents/{doc['id']}"
    edited = client.put(path, headers=tokens["writer"], json={"version": saved["version"], "body": "修改正文"})
    assert edited.status_code == 200
    assert client.get(path, headers=tokens["writer"]).json()["data"]["policyRefs"] == saved["policyRefs"]
    assert link(client, tokens["writer"], saved, []).status_code == 409
    latest = edited.json()["data"]
    assert link(client, tokens["writer"], latest, []).json()["data"]["policyRefs"] == []


def test_citation_permissions_locks_and_missing_ids(policy_client):
    client, tokens, _, _, _, cid = policy_client
    doc = create(client, tokens["writer"])
    assert link(client, tokens["other"], doc, [cid]).status_code == 404
    assert link(client, tokens["writer"], doc, ["clause:999999"]).status_code == 422
    assert link(client, tokens["writer"], doc, ["garbage"]).status_code == 422
    pending = action(client, tokens["writer"], doc, "submit").json()["data"]
    assert link(client, tokens["writer"], pending, [cid]).status_code == 409
    assert link(client, tokens["reviewer"], pending, [cid]).status_code == 403


def test_live_citation_status_updates_without_rewriting_saved_version(policy_client):
    client, tokens, _, library, policy, cid = policy_client
    doc = create(client, tokens["writer"])
    saved = link(client, tokens["writer"], doc, [cid]).json()["data"]
    library.verify(policy["id"], status="abolished", note="测试核验", fields={}, verifier="reviewer")
    loaded = client.get(f"/api/v1/gov/documents/{doc['id']}", headers=tokens["writer"]).json()["data"]
    assert loaded["policyRefs"][0]["status"] == "abolished"
    assert loaded["policyRefs"][0]["version"] == 1


def test_policy_browse_detail_and_stats_use_real_library(policy_client):
    client, tokens, _, library, policy, cid = policy_client
    assert client.get("/api/v1/gov/policies/stats").status_code == 401
    stats = client.get("/api/v1/gov/policies/stats", headers=tokens["writer"])
    assert stats.status_code == 200
    assert stats.json()["data"] == {"origin": "backend", "available": True, "total": 1, "current": 0, "abolished": 0, "unverified": 1}
    detail = client.get(f"/api/v1/gov/policies/{policy['id']}", headers=tokens["writer"])
    assert detail.status_code == 200
    data = detail.json()["data"]
    assert data["policy"]["url"] == policy["url"]
    assert len(data["clauses"]) == 2 and "分类管理" in data["clauses"][0]["content"]
    citations = client.post("/api/v1/gov/citations/batch", headers=tokens["writer"], json={"ids": [cid]}).json()["data"]["citations"]
    assert citations[0]["sourceUrl"] == policy["url"] and citations[0]["version"] == 1
    assert client.get("/api/v1/gov/policies/999999", headers=tokens["writer"]).json()["code"] != 0


def test_unverified_status_cannot_be_forged_and_import_is_admin_only(policy_client):
    client, tokens, _, _, _, cid = policy_client
    response = client.post("/api/v1/gov/policies/import", headers=tokens["writer"], files={"file": ("policy.txt", b"text")})
    assert response.status_code == 403
    response = client.post("/api/v1/gov/documents", headers=tokens["writer"], json={"title": "伪造", "policyRefs": [{"id": cid, "status": "current"}]})
    assert response.status_code == 422


def test_internal_docx_retains_selected_policy_provenance(policy_client):
    from io import BytesIO
    from docx import Document
    client, tokens, _, _, policy, cid = policy_client
    doc = create(client, tokens["writer"])
    selected = link(client, tokens["writer"], doc, [cid]).json()["data"]
    response = client.post(f"/api/v1/gov/documents/{doc['id']}/export", headers=tokens["writer"], json={"kind": "internal_draft"})
    assert response.status_code == 200
    text = "\n".join(p.text for p in Document(BytesIO(response.content)).paragraphs)
    assert policy["title"] in text and policy["url"] in text
    assert "未核验" in text and "第 1 版" in text


def test_docx_policy_tables_are_kept_in_document_order_and_flagged_for_review():
    from io import BytesIO
    from docx import Document
    from app.services.policy.extract import extract_file
    document = Document()
    document.add_paragraph("第一条 总则")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "补助上限"
    table.cell(0, 1).text = "每户十万元"
    document.add_paragraph("第二条 办理流程")
    buf = BytesIO(); document.save(buf)
    extracted = extract_file("policy.docx", buf.getvalue())
    assert extracted.text.index("总则") < extracted.text.index("每户十万元") < extracted.text.index("办理流程")
    assert extracted.reliable is False


def test_superseded_policy_warning_is_kept_in_internal_export(policy_client):
    from io import BytesIO
    from docx import Document
    from app.services.policy.clauses import segment
    client, tokens, _, library, policy, cid = policy_client
    library.verify(policy["id"], status="current", note="测试", fields={}, verifier="reviewer")
    content = "第一条 新版本政策。"
    library.store.save_import(url=policy["url"], meta={"title": policy["title"]}, content=content, clauses=segment(content), origin="file_import")
    doc = create(client, tokens["writer"])
    assert link(client, tokens["writer"], doc, [cid]).status_code == 200
    response = client.post(f"/api/v1/gov/documents/{doc['id']}/export", headers=tokens["writer"], json={"kind": "internal_draft"})
    text = "\n".join(p.text for p in Document(BytesIO(response.content)).paragraphs)
    assert "历史" in text and "替代" in text


@pytest.mark.parametrize("failure", ["missing", "store_error"])
def test_unresolved_reference_does_not_block_documents_or_repairs(policy_client, monkeypatch, failure):
    client, tokens, _, library, policy, cid = policy_client
    doc = create(client, tokens["writer"])
    selected = link(client, tokens["writer"], doc, [cid]).json()["data"]
    other = create(client, tokens["writer"], title="另一份正常文稿")
    monkeypatch.setattr(library, "citations", lambda ids: {"citations": [], "missing": ids,
                        "error": {"code": "STORE_ERROR"} if failure == "store_error" else None})
    response = client.get("/api/v1/gov/documents", headers=tokens["writer"])
    assert response.status_code == 200
    assert len(response.json()["data"]) == 2
    loaded = next(d for d in response.json()["data"] if d["id"] == doc["id"])
    assert loaded["policyRefs"][0]["unresolved"] is True
    assert loaded["policyRefs"][0]["status"] == "unverified"
    assert loaded["policyRefs"][0]["sourceUrl"] == policy["url"]
    saved = client.put(f"/api/v1/gov/documents/{doc['id']}", headers=tokens["writer"], json={"version": selected["version"], "body": "政策库故障时仍可保存"})
    assert saved.status_code == 200
    assert link(client, tokens["writer"], saved.json()["data"], []).json()["data"]["policyRefs"] == []
