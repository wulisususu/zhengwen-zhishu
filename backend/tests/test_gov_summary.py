"""提炼总结：材料切分、出处校验与 /gov/summaries 接口（真实 JWT + SQLite，只替换模型调用）。"""
import json
import shutil
import subprocess

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app import config
from app.database import Base, get_db
from app.deps import hash_password
from app.gov import summary
from app.gov.router import router as gov_router
from app.models import File, User
from app.routers import auth

PASSWORD_HASH = hash_password("local-test-password")

MEETING = "\n".join([
    "第12次主任办公会议纪要",
    "会议议定：同意启动政务数据安全自查工作。",
    "由数据安全处牵头，于2025年11月30日前完成自查报告。",
    "办公室负责会议通知的印发。",
])


@pytest.fixture()
def app_client(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "_config", {"jwt": {"secret": "isolated-gov-test-secret", "algorithm": "HS256", "expire_minutes": 60}})
    monkeypatch.setenv("GOV_DATA_DIR", str(tmp_path / "materials"))
    engine = create_engine(f"sqlite:///{tmp_path / 'test.db'}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)
    with factory() as db:
        for name in ["writer", "other"]:
            db.add(User(username=name, nickname=name, password_hash=PASSWORD_HASH, is_admin=False))
        db.commit()
    app = FastAPI()
    app.include_router(auth.router, prefix="/api/v1")
    app.include_router(gov_router, prefix="/api/v1")

    def database():
        with factory() as db:
            yield db
    app.dependency_overrides[get_db] = database
    with TestClient(app) as client:
        tokens = {}
        for name in ["writer", "other"]:
            result = client.post("/api/v1/auth/login", json={"username": name, "password": "local-test-password"})
            tokens[name] = {"Authorization": "Bearer " + result.json()["data"]["access_token"]}
        yield client, tokens
    engine.dispose()


def upload(client, headers, name="纪要.txt", text=MEETING):
    response = client.post("/api/v1/gov/materials", headers=headers, files={"file": (name, text.encode("utf-8"), "text/plain")})
    assert response.status_code == 200, response.text
    return response.json()["data"]["id"]


def fake_model(monkeypatch, reply):
    seen = {}

    def call(messages):
        seen["messages"] = messages
        return reply if isinstance(reply, str) else json.dumps(reply, ensure_ascii=False)
    monkeypatch.setattr(summary, "call_model", call)
    return seen


def material_file(tmp_path, name, content: bytes) -> File:
    path = tmp_path / name
    path.write_bytes(content)
    return File(id=7, user_id=1, filename=name, filepath=str(path), filetype=path.suffix)


# ---------- 切分 ----------

def test_text_material_is_split_into_numbered_paragraphs(tmp_path):
    segments = summary.segment([material_file(tmp_path, "a.txt", MEETING.encode("utf-8"))])
    assert [s["id"] for s in segments] == ["S1", "S2", "S3", "S4"]
    assert segments[2] == {"id": "S3", "fileId": "7", "fileName": "a.txt", "page": None, "paragraph": 3,
                           "text": "由数据安全处牵头，于2025年11月30日前完成自查报告。"}


def test_long_paragraph_is_split_at_sentence_ends(tmp_path):
    sentence = "各单位要按要求完成自查并报送结果。" * 40
    segments = summary.segment([material_file(tmp_path, "long.txt", sentence.encode("utf-8"))])
    assert len(segments) > 1
    assert all(len(s["text"]) <= summary.MAX_SEGMENT_CHARS for s in segments)
    assert "".join(s["text"] for s in segments) == sentence


def test_docx_paragraphs_and_table_rows_become_segments(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("一、总体要求")
    doc.add_paragraph("")
    doc.add_paragraph("各单位于11月底前完成自查。")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "数据安全处"
    table.rows[0].cells[1].text = "牵头"
    path = tmp_path / "plan.docx"
    doc.save(path)
    segments = summary.segment([File(id=3, user_id=1, filename="plan.docx", filepath=str(path), filetype=".docx")])
    assert [s["text"] for s in segments] == ["一、总体要求", "各单位于11月底前完成自查。", "数据安全处 | 牵头"]


@pytest.mark.skipif(not shutil.which("pdftotext"), reason="需要 poppler 的 pdftotext")
def test_pdf_segments_keep_page_numbers(tmp_path):
    pdf = tmp_path / "two.pdf"
    _write_two_page_pdf(pdf)
    segments = summary.segment([File(id=4, user_id=1, filename="two.pdf", filepath=str(pdf), filetype=".pdf")])
    assert [(s["page"], s["paragraph"], s["text"]) for s in segments] == [(1, 1, "Page one first line"), (2, 1, "Page two")]


def _write_two_page_pdf(path):
    """A minimal two-page text PDF written by hand, so the test needs no PDF library."""
    pages = ["Page one first line", "Page two"]
    objects = ["<< /Type /Catalog /Pages 2 0 R >>",
               "<< /Type /Pages /Kids [" + " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages))) + f"] /Count {len(pages)} >>"]
    for i, text in enumerate(pages):
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * i} 0 R "
                       f"/Resources << /Font << /F1 {3 + 2 * len(pages)} 0 R >> >> >>")
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    objects.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out, offsets = "%PDF-1.4\n", []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n" + "".join(f"{o:010d} 00000 n \n" for o in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    path.write_bytes(out.encode("latin-1"))
    assert subprocess.run(["pdftotext", str(path), "-"], capture_output=True).returncode == 0


def test_empty_or_unsupported_material_is_refused(tmp_path):
    with pytest.raises(summary.HTTPException) as empty:
        summary.segment([material_file(tmp_path, "blank.txt", b"\n \n")])
    assert empty.value.status_code == 422 and "没有提取到文字" in empty.value.detail
    with pytest.raises(summary.HTTPException) as image:
        summary.segment([material_file(tmp_path, "scan.png", b"\x89PNG")])
    assert image.value.status_code == 422


def test_material_over_limit_is_refused_not_truncated(tmp_path, monkeypatch):
    monkeypatch.setattr(summary, "MAX_CHARS", 20)
    with pytest.raises(summary.HTTPException) as error:
        summary.segment([material_file(tmp_path, "a.txt", MEETING.encode("utf-8"))])
    assert error.value.status_code == 422 and "未截断" in error.value.detail


# ---------- 出处校验 ----------

SEGMENTS = [
    {"id": "S1", "text": "会议议定：同意启动政务数据安全自查工作。"},
    {"id": "S2", "text": "由数据安全处牵头，于2025年11月30日前完成自查报告。"},
]


def test_items_without_valid_sources_are_dropped_with_reason():
    sections, dropped = summary.validate({"summary": [
        {"text": "启动自查工作", "sources": ["s1"]},
        {"text": "没有出处的结论", "sources": []},
        {"text": "引用了不存在的片段", "sources": ["S9"]},
    ]}, "material", SEGMENTS)
    assert sections[0]["items"] == [{"text": "启动自查工作", "sources": ["S1"], "notes": []}]
    assert [d["reason"] for d in dropped] == ["没有引用有效的原文片段"] * 2


def test_numbers_must_appear_in_cited_text():
    sections, dropped = summary.validate({"keyData": [
        {"text": "截止日期为2025年11月30日", "sources": ["S2"]},
        {"text": "共涉及 35 个单位", "sources": ["S2"]},
        {"text": "截止日期为２０２５年11月30日", "sources": ["S1"]},
    ]}, "material", SEGMENTS)
    key_data = next(s for s in sections if s["key"] == "keyData")
    assert [i["text"] for i in key_data["items"]] == ["截止日期为2025年11月30日"]
    assert [d["reason"] for d in dropped] == ["数字「35」未出现在所引原文中", "数字「2025」未出现在所引原文中"]


def test_meeting_owner_and_deadline_must_be_quoted_from_source():
    sections, dropped = summary.validate({"tasks": [
        {"task": "完成自查报告", "owner": "数据安全处", "deadline": "2025年11月30日前", "sources": ["S2"]},
        {"task": "启动自查", "owner": "信息中心", "deadline": "", "sources": ["S1"]},
    ]}, "meeting", SEGMENTS)
    tasks = next(s for s in sections if s["key"] == "tasks")["items"]
    assert tasks[0]["owner"] == "数据安全处" and tasks[0]["deadline"] == "2025年11月30日前" and tasks[0]["notes"] == []
    assert tasks[1]["owner"] == "" and "信息中心" in tasks[1]["notes"][0]
    assert dropped == []


def test_unparseable_model_output_is_an_error():
    assert summary.parse_model_json('```json\n{"summary": []}\n```') == {"summary": []}
    with pytest.raises(summary.HTTPException) as error:
        summary.parse_model_json("抱歉，我无法完成")
    assert error.value.status_code == 502


# ---------- 接口 ----------

def test_summary_api_returns_only_cited_sources(app_client, monkeypatch):
    client, tokens = app_client
    file_id = upload(client, tokens["writer"])
    seen = fake_model(monkeypatch, {
        "decisions": [{"text": "同意启动政务数据安全自查工作", "sources": ["S2"]}],
        "tasks": [{"task": "完成自查报告", "owner": "数据安全处", "deadline": "2025年11月30日前", "sources": ["S3"]},
                  {"task": "编造的任务", "owner": "", "deadline": "", "sources": ["S99"]}],
    })
    response = client.post("/api/v1/gov/summaries", headers=tokens["writer"], json={"fileIds": [int(file_id)], "mode": "meeting"})
    assert response.status_code == 200, response.text
    data = response.json()["data"]
    assert data["origin"] == "backend" and data["mode"] == "meeting"
    assert data["materials"] == [{"fileId": file_id, "fileName": "纪要.txt", "segments": 4}]
    assert [s["label"] for s in data["sections"]] == ["会议决议", "工作任务"]
    assert sorted(data["sources"]) == ["S2", "S3"]
    assert data["sources"]["S3"]["paragraph"] == 3 and data["sources"]["S3"]["fileId"] == file_id
    assert data["dropped"] == [{"section": "工作任务", "text": "编造的任务", "reason": "没有引用有效的原文片段"}]
    prompt = seen["messages"][1]["content"]
    assert "[S3]（第3段）由数据安全处牵头" in prompt and "纪要.txt" in prompt


def test_summary_api_checks_ownership_and_scope(app_client, monkeypatch):
    client, tokens = app_client
    fake_model(monkeypatch, {"summary": []})
    file_id = int(upload(client, tokens["writer"]))
    assert client.post("/api/v1/gov/summaries", json={"fileIds": [file_id]}).status_code == 401
    other = client.post("/api/v1/gov/summaries", headers=tokens["other"], json={"fileIds": [file_id]})
    assert other.status_code == 404
    second = int(upload(client, tokens["writer"], name="二.txt"))
    multi = client.post("/api/v1/gov/summaries", headers=tokens["writer"], json={"fileIds": [file_id, second]})
    assert multi.status_code == 422 and "多材料" in multi.json()["detail"]
    bad_mode = client.post("/api/v1/gov/summaries", headers=tokens["writer"], json={"fileIds": [file_id], "mode": "other"})
    assert bad_mode.status_code == 422


def test_summary_api_reports_model_failure(app_client, monkeypatch):
    client, tokens = app_client
    file_id = int(upload(client, tokens["writer"]))

    def broken(messages):
        raise RuntimeError("egress blocked")
    monkeypatch.setattr(summary, "call_model", broken)
    response = client.post("/api/v1/gov/summaries", headers=tokens["writer"], json={"fileIds": [file_id]})
    assert response.status_code == 503 and "未生成" in response.json()["detail"]
