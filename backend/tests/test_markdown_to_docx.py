import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


class MarkdownToDocxTests(unittest.TestCase):
    def _render(self, markdown: str) -> Document:
        from app.services.markdown_to_docx import markdown_to_docx

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "out.docx"
            markdown_to_docx(markdown, output)
            self.assertTrue(output.exists())
            return Document(output)

    def test_renders_core_markdown_blocks(self):
        doc = self._render(
            "# 项目周会会议纪要\n\n"
            "这是普通段落，包含**重点内容**和*补充说明*。\n\n"
            "## 本周工作进展\n\n"
            "1. 完成模型接入\n"
            "2. 完成文档生成联调\n\n"
            "- 风险跟踪\n"
            "- 任务复盘\n\n"
            "> 会议要求按计划推进。\n\n"
            "```text\n"
            "status=ok\n"
            "```\n\n"
            "---\n\n"
            "<<<PAGE_BREAK>>>\n\n"
            "### 下周计划\n\n"
            "| 任务 | 负责人 | 状态 |\n"
            "| --- | --- | --- |\n"
            "| 联调测试 | 张三 | 进行中 |\n"
        )

        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        self.assertIn("项目周会会议纪要", paragraphs)
        self.assertIn("本周工作进展", paragraphs)
        self.assertIn("下周计划", paragraphs)
        self.assertIn("会议要求按计划推进。", paragraphs)
        self.assertIn("status=ok", paragraphs)
        self.assertEqual(1, len(doc.tables))
        self.assertEqual("任务", doc.tables[0].cell(0, 0).text)
        self.assertEqual("张三", doc.tables[0].cell(1, 1).text)

    def test_renders_resume_meeting_plan_and_report_examples(self):
        examples = [
            "# 张三简历\n\n## 个人信息\n\n姓名：张三\n\n## 教育经历\n\n- 本科：计算机科学与技术",
            "# 项目周会会议纪要\n\n## 会议时间\n\n2026年6月8日\n\n## 任务分工\n\n| 任务 | 责任人 |\n| --- | --- |\n| 测试 | 李四 |",
            "# 智慧校园AI学习平台项目实施方案\n\n## 项目背景\n\n建设AI学习平台。\n\n## 实施计划\n\n```mermaid\ngantt\n    title 项目计划\n```",
            "# 2026年中国人工智能行业发展研究报告\n\n## 研究背景\n\n本报告分析行业发展态势。\n\n## 结论\n\n行业进入深化应用阶段。",
        ]

        for markdown in examples:
            doc = self._render(markdown)
            text = "\n".join(p.text for p in doc.paragraphs)
            first_title = markdown.splitlines()[0].lstrip("# ").strip()
            self.assertIn(first_title, text)

    def test_renders_business_docx_formatting_for_loose_tables_toc_and_fonts(self):
        from app.services.markdown_to_docx import markdown_to_docx

        markdown = (
            "# AI 文档生成平台软件需求规格说明书\n\n"
            "## 目录\n\n"
            "## 功能需求\n\n"
            "### 聊天功能\n\n"
            "| 需求编号 | 功能描述 | 优先级 |\n"
            "| FR-01 | 支持自然语言对话 | P0 |\n"
            "| FR-02 | 支持多轮上下文 | P0 |\n\n"
            "## 非功能需求\n\n"
            "平台应支持 RK3588 Ubuntu22.04 部署。\n"
        )

        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "srs.docx"
            markdown_to_docx(markdown, output)
            doc = Document(output)

            self.assertGreaterEqual(len(doc.tables), 1)
            self.assertEqual("需求编号", doc.tables[0].cell(0, 0).text)
            self.assertEqual("FR-01", doc.tables[0].cell(1, 0).text)

            with zipfile.ZipFile(output) as package:
                document_xml = package.read("word/document.xml").decode("utf-8", errors="replace")
                styles_xml = package.read("word/styles.xml").decode("utf-8", errors="replace")

        self.assertIn("TOC", document_xml)
        self.assertIn('w:eastAsia="Microsoft YaHei"', styles_xml)
        self.assertIn('<w:sz w:val="22"', styles_xml)


if __name__ == "__main__":
    unittest.main()
